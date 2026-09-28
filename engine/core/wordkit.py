"""Small python-docx helper for engine-generated Word documents: A4 page, header / footer stamp, headings, paragraphs,
bullets and grid tables with a repeating shaded header row."""
from __future__ import annotations

from pathlib import Path


class Doc:
    def __init__(self, base: Path | None, header: str, footer: str, landscape: bool = False, size: float = 9):
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Cm, Pt
        self.doc = Document(str(base)) if base and base.exists() else Document()
        for p in list(self.doc.paragraphs):
            p._element.getparent().remove(p._element)
        for t in list(self.doc.tables):
            t._element.getparent().remove(t._element)
        self.doc.styles["Normal"].font.size = Pt(size)
        sec = self.doc.sections[0]
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(1.8)
        sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)
        self.width_cm = (29.7 if landscape else 21.0) - 3.6
        for part, text in ((sec.header, header), (sec.footer, footer)):
            p = part.paragraphs[0] if part.paragraphs else part.add_paragraph()
            p.text = text
            p.runs[0].font.size = Pt(7)

    def title(self, text, size=18):
        from docx.shared import Pt
        for run in self.doc.add_heading(text, level=0).runs:
            run.font.size = Pt(size)

    def h(self, text, level=1):
        self.doc.add_heading(text, level=level).paragraph_format.keep_with_next = True

    def p(self, text, bold_lead: str | None = None):
        par = self.doc.add_paragraph()
        if bold_lead:
            par.add_run(bold_lead).bold = True
        par.add_run(text)
        return par

    def bullets(self, items):
        for t in items:
            self.doc.add_paragraph(t, style="List Bullet")

    def table(self, head, rows, widths_cm, size=7.5):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
        if sum(widths_cm) > self.width_cm + 0.05:
            f = self.width_cm / sum(widths_cm)
            widths_cm = [w * f for w in widths_cm]
        t = self.doc.add_table(rows=1 if head else 0, cols=len(widths_cm))
        t.style = "Table Grid"
        if head:
            for c, hd in zip(t.rows[0].cells, head):
                c.text = hd
                run = c.paragraphs[0].runs[0]
                run.bold, run.font.size, run.font.color.rgb = True, Pt(size), RGBColor(0xFF, 0xFF, 0xFF)
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), "1F3864")
                c._element.get_or_add_tcPr().append(shd)
            rh = OxmlElement("w:tblHeader")
            rh.set(qn("w:val"), "true")
            t.rows[0]._tr.get_or_add_trPr().append(rh)
        for r in rows:
            cells = t.add_row().cells
            for c, v in zip(cells, r):
                c.text = "" if v is None else str(v)
                for par in c.paragraphs:
                    par.paragraph_format.space_after = Pt(0)
                    for run in par.runs:
                        run.font.size = Pt(size)
            cs = OxmlElement("w:cantSplit")
            cs.set(qn("w:val"), "true")
            t.rows[-1]._tr.get_or_add_trPr().append(cs)
        t.autofit = False
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths_cm):
            gc.set(qn("w:w"), str(int(w * 567)))
        for row in t.rows:
            for cell, w in zip(row.cells, widths_cm):
                cell.width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def image(self, path: Path, width_cm: float, caption: str | None = None):
        from docx.shared import Cm
        self.doc.add_picture(str(path), width=Cm(min(width_cm, self.width_cm)))
        if caption:
            self.p(caption).runs[0].italic = True

    def save(self, path: Path) -> Path:
        self.doc.save(path)
        return path
