"""Codes, standards and regulations register (Appendix A.19 of the Owner contract) with the permits and authorities register:
Excel + Word + PDF.

Sources: reference records (laws, regulations, Grid Code, standards; precedence per ER-01.07), permit records and the
authority parties. Checks:
- every code cited in an Employer's Requirement is in the register (ER coverage);
- every law / regulation names the enforcing authority;
- agreement: a register entry is agreed by the Owner and by every party responsible for compliance (IEC for the power
  island), a permit by the Owner, the applicant and every supporting party;
- permits: legal basis is a law, regulation or Grid Code; apply_by + lead time <= needed_by; granted permits carry the date;
- every authority has at least one regulation or permit.
Options: pdf=no."""
from __future__ import annotations

import re
from datetime import date, timedelta

from ..core.runner import Context, Engine, pdf_from_office
from .tie_in_register import TieInRegister

CITE = re.compile(r"(ASME (?:PTC ?[\d.]+|B[\d.]+M?|Section [IVX]+|IX)|IEC(?:/IEEE)? \d{4,5}(?:-\d+)*|(?:TS )?EN (?:ISO )?\d{3,5}"
                  r"(?:-\d+)*|ISO \d{3,5}(?:-\d+)*|NFPA \d+|IEEE \d+|API \d+|TS \d+|VGB-S-[\d-]+|HEI\b|PED [\d/]+EU|TBDY \d{4}|"
                  r"Law \(?No\. ?\d{3,4}|No\. ?\d{4}|Grid Code)")
PRECEDENCE = {1: "Turkish laws and regulations", 2: "Grid Code", 3: "Contract", 4: "Codes and standards (Appendix A.19)",
              5: "Good industry practice / project libraries"}
LEGAL = ("law", "regulation", "grid_code")


def _norm(s):
    s = re.sub(r"\s+", " ", s.strip()).replace("BPVC ", "").replace("Section ", "")
    m = re.match(r"(?:Law \(?)?No\. ?(\d{3,4})", s)
    return f"Law {m.group(1)}" if m else s


def cited_codes(text):
    return sorted({_norm(m) for m in CITE.findall(text or "")})


def covered(cite, codes):
    """cite matches a register code exactly, as its part/prefix ('IEC 60034' ~ 'IEC 60034-1'), or via a dual logo."""
    for code in codes:
        alts = {_norm(code)}
        m = re.match(r"^([A-Z]+)/([A-Z]+) (.+)$", code)
        if m:
            alts |= {f"{m.group(1)} {m.group(3)}", f"{m.group(2)} {m.group(3)}"}
        for c in alts:
            if cite == c or c.startswith(cite + "-") or cite.startswith(c + "-") or c.startswith(cite + " ") \
                    or (cite == "HEI" and c.startswith("HEI")):
                return True
    return False


def check_registers(refs, permits, requirements, parties, owner):
    out = []

    def add(rid, chk, st, txt):
        out.append((rid, chk, st, txt))

    codes = [r.get("code") or "" for r in refs]
    by_id = {r["id"]: r for r in refs}
    for req in sorted(requirements, key=lambda r: r["id"]):
        for c in cited_codes(req.get("text", "") + " " + (req.get("acceptance") or "")):
            if covered(c, codes):
                add(req["id"], "ER coverage", "OK", f"cites {c}: in the register")
            else:
                add(req["id"], "ER coverage", "WARN", f"cites {c}: not in the register")
    for r in sorted(refs, key=lambda x: (x.get("precedence") or 9, x.get("code") or "")):
        rid = r["id"]
        if r.get("kind") in LEGAL and not r.get("authority"):
            add(rid, "authority", "WARN", f"{r.get('code')}: {r['kind']} without enforcing authority")
        need = set(r.get("responsible") or []) | ({owner} if owner else set())
        if r.get("status") in ("agreed",):
            miss = need - set(r.get("agreed_by") or [])
            add(rid, "agreement", "WARN" if miss else "OK",
                f"{r.get('code')}: agreed by {', '.join(r.get('agreed_by') or [])}"
                + (f"; missing {', '.join(sorted(miss))}" if miss else ""))
        elif r.get("status") != "superseded":
            add(rid, "agreement", "WARN", f"{r.get('code')}: status {r.get('status') or '?'}")
        if not r.get("kind"):
            add(rid, "completeness", "WARN", f"{r.get('code') or r.get('title')}: kind / precedence missing")
    for p in sorted(permits, key=lambda x: x["id"]):
        pid = p["id"]
        bad = [b for b in p.get("legal_basis", []) if by_id.get(b, {}).get("kind") not in LEGAL]
        add(pid, "legal basis", "WARN" if bad or not p.get("legal_basis") else "OK",
            ("not a law/regulation: " + ", ".join(bad)) if bad else
            ", ".join(by_id.get(b, {}).get("code", b) for b in p.get("legal_basis", [])))
        au = parties.get(p.get("authority"), {})
        if au.get("role") not in ("authority", "tso"):
            add(pid, "authority", "WARN", f"{p.get('authority')} is not an authority party")
        if p.get("apply_by") and p.get("lead_time") is not None and p.get("needed_by"):
            ready = date.fromisoformat(p["apply_by"]) + timedelta(days=p["lead_time"])
            fl = (date.fromisoformat(p["needed_by"]) - ready).days
            add(pid, "timing", "WARN" if fl < 0 else ("INFO" if fl < 14 else "OK"),
                f"apply {p['apply_by']} + {p['lead_time']} d = {ready.isoformat()}; needed {p['needed_by']}: float {fl} d")
        elif p.get("permit_status") not in ("granted",) and p.get("needed_by") and p.get("lead_time") is None:
            add(pid, "timing", "INFO", f"no lead time: procedural, needed from {p['needed_by']}")
        if p.get("permit_status") == "granted" and not p.get("obtained_date"):
            add(pid, "timing", "INFO", "granted - date of the permit not recorded")
        need = {p.get("applicant")} | set(p.get("support_by") or []) | ({owner} if owner else set())
        if p.get("status") == "agreed":
            miss = need - set(p.get("agreed_by") or [])
            add(pid, "agreement", "WARN" if miss else "OK", f"agreed by {', '.join(p.get('agreed_by') or [])}"
                + (f"; missing {', '.join(sorted(miss))}" if miss else ""))
        else:
            add(pid, "agreement", "WARN", f"status {p.get('status')}")
    used = {r.get("authority") for r in refs} | {p.get("authority") for p in permits}
    for a in sorted(x for x, v in parties.items() if v.get("role") == "authority"):
        if a not in used:
            add(a, "authority use", "INFO", f"{parties[a]['name']}: no regulation or permit")
    return out


class StandardsRegister(Engine):
    name = "standards_register"
    title = "Codes, standards and regulations register (Appendix A.19) with permits and authorities (Excel + Word + PDF)"
    version = "1.0.0"
    inputs = ["project", "reference", "permit", "party", "requirement", "clarification", "contract"]
    formats = ["xlsx", "docx", "pdf"]
    code_deps = ["engine/engines/tie_in_register.py", "templates/docx/datasheet_base.docx"]

    _table = TieInRegister._table

    def run(self, ctx: Context):
        s = ctx.store
        refs, permits = s.records("reference"), s.records("permit")
        if not refs:
            ctx.warnings.append("no reference records - nothing generated")
            return []
        parties = {p["id"]: p for p in s.records("party")}
        owner = next((p for p, v in parties.items() if v.get("role") == "owner"), None)
        checks = check_registers(refs, permits, s.records("requirement"), parties, owner)
        for rid, chk, st, txt in checks:
            if st == "WARN":
                ctx.warnings.append(f"{rid} {chk}: {txt}")
        stem = f"{ctx.project_record()['id']}_codes_standards_permits_register"
        files = [self._workbook(ctx, refs, permits, parties, checks, stem)]
        doc = self._document(ctx, refs, permits, parties, checks, stem)
        files.append(doc)
        if ctx.options.get("pdf", "yes") != "no":
            pdf = pdf_from_office(doc, ctx)
            if pdf:
                files.append(pdf)
        return files

    @staticmethod
    def _order(refs):
        kind = {"law": 0, "regulation": 1, "grid_code": 2}
        return sorted(refs, key=lambda r: (r.get("precedence") or 9, kind.get(r.get("kind"), 3), r.get("jurisdiction") != "TR",
                                           r.get("code") or ""))

    @staticmethod
    def _names(ids, parties):
        return ", ".join(parties.get(i, {}).get("short_name") or i for i in ids or [])

    def _workbook(self, ctx, refs, permits, parties, checks, stem):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        head = PatternFill("solid", fgColor="1F3864")
        fills = {"WARN": PatternFill("solid", fgColor="F8CBAD"), "INFO": PatternFill("solid", fgColor="FFF2CC"),
                 "OK": PatternFill("solid", fgColor="E2EFDA")}
        wb = Workbook()

        def sheet(ws, title, cols, rows, wrap=()):
            ws["A1"] = f"{ctx.project_record().get('name', '')} - {title}"
            ws["A1"].font = Font(bold=True, size=14)
            ws["A2"] = (f"engine {self.name} v{self.version} | generated {ctx.stamp['generated_at'][:19]}Z by "
                        f"{ctx.stamp['generated_by']} | data {ctx.stamp['inputs_hash']}")
            ws["A2"].font = Font(italic=True, size=8, color="595959")
            for c, (n, w) in enumerate(cols, 1):
                cell = ws.cell(4, c, n)
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                ws.column_dimensions[get_column_letter(c)].width = w
            for i, row in enumerate(rows, 5):
                for c, v in enumerate(row, 1):
                    v = "\n".join(v) if isinstance(v, list) else v
                    ws.cell(i, c, v).alignment = Alignment(wrap_text=c in wrap, vertical="top")
            ws.freeze_panes = "B5"
            ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}4"
            ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            return ws

        ws = wb.active
        ws.title = "Codes and standards"
        sheet(ws, "Codes, standards and regulations register (Appendix A.19)",
              [("ID", 16), ("Precedence", 10), ("Kind", 10), ("Jurisdiction", 10), ("Code", 26), ("Title", 55), ("Edition", 18),
               ("Issuer", 10), ("Disciplines", 16), ("Applies to", 45), ("Responsible", 16), ("Authority", 18),
               ("Requirements", 14), ("Deviation / clarification", 40), ("Status", 9), ("Agreed by", 18), ("Agreed in", 14)],
              [[r["id"], r.get("precedence"), r.get("kind"), r.get("jurisdiction"), r.get("code"), r["title"], r.get("edition"),
                r.get("issuer"), r.get("disciplines"), r.get("applies_to"), self._names(r.get("responsible"), parties),
                parties.get(r.get("authority"), {}).get("short_name") or r.get("authority"), r.get("requirement_refs"),
                r.get("deviation"), r.get("status"), self._names(r.get("agreed_by"), parties), r.get("agreed_refs")]
               for r in self._order(refs)], wrap=(6, 9, 10, 13, 14, 17))
        refcode = {r["id"]: r.get("code") for r in refs}
        wp = wb.create_sheet("Permits")
        sheet(wp, "Permits, licences and approvals register",
              [("ID", 8), ("Permit", 36), ("Authority", 16), ("Legal basis", 22), ("Phase", 13), ("Applicant", 12),
               ("Support by", 14), ("Scope", 45), ("Deliverables (who provides)", 50), ("Lead time d", 8), ("Apply by", 11),
               ("Needed by", 11), ("Gates", 26), ("Permit status", 12), ("Obtained", 11), ("Tie-ins", 12), ("Status", 9),
               ("Agreed by", 18), ("Agreed in", 14)],
              [[p["id"], p["title"], parties.get(p["authority"], {}).get("short_name") or p["authority"],
                [refcode.get(b, b) for b in p.get("legal_basis", [])], p.get("phase"),
                parties.get(p.get("applicant"), {}).get("short_name") or p.get("applicant"),
                self._names(p.get("support_by"), parties), p.get("scope"), p.get("deliverables"), p.get("lead_time"),
                p.get("apply_by"), p.get("needed_by"), p.get("gates"), p.get("permit_status"), p.get("obtained_date"),
                p.get("tie_ins"), p.get("status"), self._names(p.get("agreed_by"), parties), p.get("agreed_refs")]
               for p in sorted(permits, key=lambda x: (x.get("needed_by") or "9", x["id"]))], wrap=(2, 4, 8, 9, 13, 16))
        wa = wb.create_sheet("Authorities")
        sheet(wa, "Authorities", [("ID", 10), ("Authority", 70), ("Short name", 22), ("Regulations", 50), ("Permits", 40)],
              [[a, v["name"], v.get("short_name"), [r.get("code") for r in refs if r.get("authority") == a],
                [f"{p['id']} {p['title']}" for p in permits if p.get("authority") == a]]
               for a, v in sorted(parties.items()) if v.get("role") in ("authority", "tso")], wrap=(4, 5))
        wc = wb.create_sheet("Checks")
        sheet(wc, "Register checks", [("Item", 16), ("Check", 16), ("Result", 8), ("Detail", 120)], checks, wrap=(4,))
        for i, c in enumerate(checks, 5):
            wc.cell(i, 3).fill = fills[c[2]]
        out = ctx.out_dir / f"{stem}.xlsx"
        wb.save(out)
        return out

    def _document(self, ctx, refs, permits, parties, checks, stem):
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Cm, Pt
        base = ctx.template("docx", "datasheet_base.docx")
        doc = Document(str(base)) if base.exists() else Document()
        for p in list(doc.paragraphs):
            p._element.getparent().remove(p._element)
        for t in list(doc.tables):
            t._element.getparent().remove(t._element)
        self.doc = doc
        doc.styles["Normal"].font.size = Pt(9)
        sec = doc.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(1.5)
        sec.top_margin, sec.bottom_margin = Cm(1.4), Cm(1.4)
        proj = ctx.project_record()
        hdr = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
        hdr.text = f"{proj.get('name', '')} | {stem} | Appendix A.19 codes, standards and regulations; permits and authorities"
        hdr.runs[0].font.size = Pt(7)
        ftr = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        ftr.text = (f"engine {self.name} v{self.version} | data {ctx.stamp['inputs_hash']} | generated "
                    f"{ctx.stamp['generated_at'][:10]} - generated from the project database; do not edit")
        ftr.runs[0].font.size = Pt(7)
        tqs = sorted({t for r in refs + permits for t in (r.get("agreed_refs") or [])})
        doc.add_heading(f"{proj.get('name', '')}\nCodes, Standards and Regulations (Appendix A.19), Permits and Authorities",
                        level=0)
        agreed = sum(1 for r in refs if r.get("status") == "agreed")
        doc.add_paragraph(
            f"{len(refs)} register entries ({agreed} agreed), {len(permits)} permits "
            f"({sum(1 for p in permits if p.get('status') == 'agreed')} agreed), "
            f"{sum(1 for v in parties.values() if v.get('role') in ('authority', 'tso'))} authorities. "
            + (f"Agreed in {', '.join(tqs)}. " if tqs else "")
            + f"Checks: {sum(1 for c in checks if c[2] == 'WARN')} warning(s), {sum(1 for c in checks if c[2] == 'INFO')} note(s).")
        doc.add_heading("1. Order of precedence (ER-01.07)", level=1)
        for k, v in PRECEDENCE.items():
            doc.add_paragraph(v + (" - where codes conflict the more stringent requirement applies" if k == 4 else ""),
                              style="List Number")
        doc.add_paragraph("Edition: 'edition at Base Date' = the edition in force 28 days before the tender submission date; "
                          "later editions apply only by variation. Responsible = parties whose design or work must comply "
                          "(IEPC = Istanbul EPC, IEC = Imaginary Electric).")
        doc.add_heading("2. Laws, regulations and Grid Code", level=1)
        rows = [[r.get("code"), r["title"], r.get("edition") or "-",
                 parties.get(r.get("authority"), {}).get("short_name") or "-", r.get("applies_to") or "-",
                 self._names(r.get("responsible"), parties), r.get("status")]
                for r in self._order(refs) if r.get("kind") in LEGAL]
        self._table(["Code", "Title", "Edition", "Authority", "Applies to / obligations", "Responsible", "Status"], rows,
                    [3.2, 6.8, 2.2, 2.4, 7.5, 2.6, 1.8], size=7)
        doc.add_heading("3. Codes and standards", level=1)
        rows = [[r.get("code"), r["title"], r.get("edition") or "-", r.get("applies_to") or "-",
                 ", ".join(r.get("requirement_refs") or []) or "-", self._names(r.get("responsible"), parties),
                 r.get("deviation") or "-", r.get("status")]
                for r in self._order(refs) if r.get("kind") not in LEGAL]
        self._table(["Code", "Title", "Edition", "Applies to", "ER", "Responsible", "Deviation / clarification", "Status"], rows,
                    [3.0, 5.6, 2.0, 5.0, 2.0, 2.0, 5.3, 1.8], size=7)
        doc.add_heading("4. Permits, licences and approvals", level=1)
        refcode = {r["id"]: r.get("code") for r in refs}
        rows = [[p["id"], p["title"], parties.get(p["authority"], {}).get("short_name") or p["authority"],
                 ", ".join(refcode.get(b, b) for b in p.get("legal_basis", [])),
                 parties.get(p.get("applicant"), {}).get("short_name") or p.get("applicant"),
                 (p.get("deliverables") or "-"),
                 f"{p.get('apply_by') or '-'} / {p.get('lead_time') if p.get('lead_time') is not None else '-'} d",
                 f"{p.get('needed_by') or '-'}\n{p.get('gates') or ''}", (p.get("permit_status") or "").replace("_", " ")]
                for p in sorted(permits, key=lambda x: (x.get("needed_by") or "9", x["id"]))]
        self._table(["ID", "Permit", "Authority", "Legal basis", "Applicant", "Deliverables (who provides)",
                     "Apply by / lead", "Needed by / gates", "Permit status"], rows,
                    [1.7, 4.0, 2.1, 3.0, 1.8, 6.9, 2.4, 2.8, 2.0], size=7)
        doc.add_heading("5. Authorities", level=1)
        rows = [[a, v["name"], ", ".join(r.get("code") for r in refs if r.get("authority") == a) or "-",
                 ", ".join(p["id"] for p in permits if p.get("authority") == a) or "-"]
                for a, v in sorted(parties.items()) if v.get("role") in ("authority", "tso")]
        self._table(["ID", "Authority", "Regulations enforced", "Permits"], rows, [2.0, 9.0, 10.0, 5.7], size=7)
        doc.add_heading("6. Check results", level=1)
        rows = [list(c) for c in checks if c[2] != "OK"]
        if rows:
            self._table(["Item", "Check", "Result", "Detail"], rows, [2.6, 2.6, 1.4, 20.1], size=7)
        else:
            doc.add_paragraph(f"No warnings or notes ({len(checks)} checks passed; details in the Excel register).")
        out = ctx.out_dir / f"{stem}.docx"
        doc.save(out)
        return out


ENGINE = StandardsRegister()
