"""Master Document Register (MDR) .xlsx: documents with their latest issue, review status, open comments,
plan vs actual dates, lateness and earned progress; plus revision history, comment log and discipline summary.
Issue state is derived from document_revision records (never typed twice)."""
from __future__ import annotations

from datetime import date

from ..core import planning
from ..core.runner import Context, Engine

PURPOSE_ORDER = ["IFI", "IFR", "IFA", "IFD", "IFP", "IFC", "AB"]


def latest_revisions(store) -> dict[str, dict]:
    out = {}
    for r in store.records("document_revision"):
        cur = out.get(r["document"])
        if cur is None or (r["issue_date"], r["revision"]) > (cur["issue_date"], cur["revision"]):
            out[r["document"]] = r
    return out


def register_rows(store, data_date: date) -> list[dict]:
    prog = planning.document_progress(store, data_date)
    latest = latest_revisions(store)
    revs_by_doc: dict = {}
    for r in store.records("document_revision"):
        revs_by_doc.setdefault(r["document"], []).append(r)
    open_by_doc: dict = {}
    for c in store.records("review_comment"):
        if c.get("status", "open") in ("open", "responded"):
            doc = store.get("document_revision", c["revision"])["document"]
            open_by_doc[doc] = open_by_doc.get(doc, 0) + 1
    rows = []
    for d in store.records("document"):
        lr = latest.get(d["id"], {})
        ifc = min((r["issue_date"] for r in revs_by_doc.get(d["id"], []) if r["purpose"] in ("IFC", "AB")), default=None)
        planned_ifc = d.get("planned_ifc")
        late = bool(planned_ifc and not ifc and d.get("status") != "cancelled"
                    and date.fromisoformat(planned_ifc) < data_date)
        rows.append({"doc": d, "latest": lr, "actual_ifc": ifc, "late": late,
                     "open_comments": open_by_doc.get(d["id"], 0), "progress": prog.get(d["id"])})
    return rows


class DocumentRegister(Engine):
    name = "document_register"
    title = "Master Document Register (Excel): status, revisions, comments, progress, lateness"
    version = "1.0.0"
    inputs = ["project", "document", "document_revision", "review_comment", "progress_rule", "system", "wbs"]
    formats = ["xlsx"]

    def run(self, ctx: Context):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        s = ctx.store
        prj = ctx.project_record()
        dd = date.fromisoformat(prj["data_date"]) if prj.get("data_date") else date.today()
        rows = register_rows(s, dd)
        wb = Workbook()
        head = PatternFill("solid", fgColor="1F3864")
        late_fill = PatternFill("solid", fgColor="F8D7D7")

        def sheet(ws, title, cols, widths):
            ws["A1"] = f"{prj.get('name')} - {title}"
            ws["A1"].font = Font(bold=True, size=14)
            ws["A2"] = (f"data date {dd.isoformat()} | engine {self.name} v{self.version} | generated "
                        f"{ctx.stamp['generated_at'][:19]}Z by {ctx.stamp['generated_by']} | data {ctx.stamp['inputs_hash']}")
            ws["A2"].font = Font(italic=True, size=8, color="595959")
            for c, name in enumerate(cols, 1):
                cell = ws.cell(4, c, name)
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for c, wd in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(c)].width = wd
            ws.freeze_panes = "B5"
            ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}4"
            ws.print_title_rows = "4:4"
            ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True

        ws = wb.active
        ws.title = "MDR"
        cols = ["Document no.", "Title", "Discipline", "Type", "AVEVA class", "System", "WBS", "Responsible",
                "Work status", "Latest rev", "Purpose", "Issue date", "Review code", "Planned IFR", "Planned IFA",
                "Planned IFC", "Forecast IFC", "Actual IFC", "Late", "Open comments", "Progress %", "Weight [h]"]
        sheet(ws, "Master Document Register", cols,
              [20, 40, 12, 12, 18, 10, 10, 11, 11, 7, 8, 11, 8, 11, 11, 11, 11, 11, 6, 9, 9, 9])
        for i, r in enumerate(sorted(rows, key=lambda r: r["doc"]["id"]), 5):
            d, lr = r["doc"], r["latest"]
            vals = [d["id"], d["title"], d["discipline"], d["doc_type"], d.get("aveva_class"), d.get("system"),
                    d.get("wbs"), d.get("responsible"), d.get("status"), lr.get("revision"), lr.get("purpose"),
                    _d(lr.get("issue_date")), lr.get("review_code"), _d(d.get("planned_ifr")), _d(d.get("planned_ifa")),
                    _d(d.get("planned_ifc")), _d(d.get("forecast_ifc")), _d(r["actual_ifc"]),
                    "LATE" if r["late"] else "", r["open_comments"] or None,
                    None if r["progress"] is None else round(r["progress"], 1), d.get("weight")]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(i, c, v)
                if isinstance(v, date):
                    cell.number_format = "dd-mmm-yy"
                if r["late"]:
                    cell.fill = late_fill
            if r["late"]:
                ws.cell(i, 19).font = Font(bold=True, color="D03B3B")
        if not rows:
            ws.cell(5, 1, "No documents in the database yet.")

        ws = wb.create_sheet("Revisions")
        sheet(ws, "Revision history", ["Document no.", "Rev", "Purpose", "Issue date", "Description", "Prepared",
                                       "Checked", "Approved", "Transmittal", "Review code", "Review date"],
              [20, 6, 8, 11, 34, 9, 9, 9, 16, 8, 11])
        for i, r in enumerate(sorted(s.records("document_revision"), key=lambda r: (r["document"], r["issue_date"])), 5):
            for c, v in enumerate([r["document"], r["revision"], r["purpose"], _d(r["issue_date"]), r.get("description"),
                                   r.get("prepared_by"), r.get("checked_by"), r.get("approved_by"), r.get("delivery"),
                                   r.get("review_code"), _d(r.get("review_date"))], 1):
                cell = ws.cell(i, c, v)
                if isinstance(v, date):
                    cell.number_format = "dd-mmm-yy"

        ws = wb.create_sheet("Comments")
        sheet(ws, "Review comment log", ["Comment", "Document rev", "Originator", "Location", "Comment text",
                                         "Response", "Status", "Closed in"], [16, 18, 14, 14, 50, 50, 10, 16])
        for i, c in enumerate(s.records("review_comment"), 5):
            for k, v in enumerate([c["id"], c["revision"], c["originator"], c.get("location"), c["comment"],
                                   c.get("response"), c.get("status"), c.get("closed_in")], 1):
                ws.cell(i, k, v).alignment = Alignment(wrap_text=k in (5, 6), vertical="top")

        ws = wb.create_sheet("Summary")
        sheet(ws, "Summary by discipline", ["Discipline", "Documents", "Issued (any)", "IFC / AB", "Late",
                                           "Open comments", "Weighted progress %"], [16, 11, 12, 10, 8, 14, 18])
        by: dict = {}
        for r in rows:
            if r["doc"].get("status") == "cancelled":
                continue
            by.setdefault(r["doc"]["discipline"], []).append(r)
        for i, (disc, rs) in enumerate(sorted(by.items()), 5):
            w = [x["doc"].get("weight") or 1.0 for x in rs]
            ws.append([disc, len(rs), sum(1 for x in rs if x["latest"]), sum(1 for x in rs if x["actual_ifc"]),
                       sum(1 for x in rs if x["late"]), sum(x["open_comments"] for x in rs),
                       round(sum((x["progress"] or 0) * wi for x, wi in zip(rs, w)) / sum(w), 1)])
        out = ctx.out_dir / f"{prj['id']}_document_register.xlsx"
        wb.save(out)
        n_late = sum(1 for r in rows if r["late"])
        if n_late:
            ctx.warnings.append(f"{n_late} document(s) late against planned IFC")
        return [out]


def _d(v):
    return date.fromisoformat(v) if isinstance(v, str) and v else v


ENGINE = DocumentRegister()
