"""Equipment datasheets (.docx, + .pdf when LibreOffice is available), one per equipment tag.

Options:  ids=TAG1,TAG2   only these tags (default: all not deleted)
          pdf=no          skip PDF conversion
Layout/styles come from templates/docx/datasheet_base.docx (edit that file in Word to restyle).
"""
from __future__ import annotations

from ..core.runner import Context, Engine, pdf_from_office

SECTIONS = [
    ("General", ["description", "system", "equipment_type", "service", "quantity", "redundancy", "location", "status"]),
    ("Design data", ["rated_power", "voltage", "design_flow", "design_pressure", "design_temperature", "material"]),
    ("Supply", ["manufacturer", "model"]),
]


class EquipmentDatasheets(Engine):
    name = "equipment_datasheets"
    title = "Equipment datasheets (Word + PDF), one per tag"
    version = "1.1.1"
    inputs = ["project", "system", "equipment", "document", "source", "reference", "decision"]
    formats = ["docx", "pdf"]

    def run(self, ctx: Context):
        from docx import Document
        from docx.shared import Cm

        s = ctx.store
        fields = s.schema("equipment").fields
        prj = ctx.project_record()
        items = [r for r in s.records("equipment") if r.get("status") != "deleted"]
        if ctx.options.get("ids"):
            wanted = set(ctx.options["ids"].split(","))
            missing = wanted - {r["id"] for r in items}
            if missing:
                raise ValueError(f"unknown/deleted equipment tags: {', '.join(sorted(missing))}")
            items = [r for r in items if r["id"] in wanted]
        base = ctx.template("docx", "datasheet_base.docx")
        files = []
        for eq in items:
            doc = Document(str(base)) if base.exists() else Document()
            ds = s.get("document", eq["datasheet"]) if eq.get("datasheet") else None
            docno = ds["id"] if ds else f"{prj['id']}-DS-{eq['id']}"
            rev = ds.get("revision", "A") if ds else "-"
            for p in list(doc.paragraphs):   # template body placeholder text is discarded
                p._element.getparent().remove(p._element)
            for sec in doc.sections:
                _set_text(sec.header.paragraphs[0], f"{prj.get('name')}  |  {docno}  Rev {rev}")
                _set_text(sec.footer.paragraphs[0], f"Generated {ctx.stamp['generated_at'][:19]}Z by {ctx.stamp['generated_by']}"
                                                    f" | {self.name} v{self.version} | data {ctx.stamp['inputs_hash']}"
                                                    f" | git {ctx.stamp['git_commit']} | do not edit - regenerate")
            doc.add_heading(f"Equipment Datasheet - {eq['id']}", level=1)
            block = [("Project", f"{prj['id']} - {prj.get('name', '')}"), ("Document no.", docno),
                     ("Revision", rev), ("Tag", eq["id"]), ("Title", eq["description"]),
                     ("Equipment status", eq.get("status", "-"))]
            t = doc.add_table(rows=0, cols=2)
            t.style = "Table Grid"
            for k, v in block:
                c = t.add_row().cells
                c[0].text, c[1].text = k, str(v)
                c[0].paragraphs[0].runs[0].bold = True
            _widths(t, (Cm(4.5), Cm(12.5)))

            for title, keys in SECTIONS:
                rows = [(k, eq.get(k)) for k in keys if eq.get(k) is not None]
                if not rows:
                    continue
                doc.add_heading(title, level=2)
                t = doc.add_table(rows=0, cols=3)
                t.style = "Table Grid"
                for k, v in rows:
                    if k == "system" and s.get("system", v):
                        v = f"{v} - {s.get('system', v)['title']}"
                    c = t.add_row().cells
                    c[0].text = k.replace("_", " ").capitalize()
                    c[1].text = str(v)
                    c[2].text = fields[k].get("unit", "")
                _widths(t, (Cm(4.5), Cm(10), Cm(2.5)))
            doc.add_heading("Basis", level=2)
            refs = eq.get("basis_refs") or []
            if refs:
                for ref in refs:
                    ent, rid = ref.split(":", 1)
                    r = s.get(ent, rid) or {}
                    doc.add_paragraph(f"{rid}  {r.get('code') or ''} {r.get('title', '')}".strip(), style="List Bullet")
            else:
                doc.add_paragraph("No basis recorded - values are unverified.")
            if eq.get("remarks"):
                doc.add_heading("Remarks", level=2)
                doc.add_paragraph(eq["remarks"])

            out = ctx.out_dir / f"{docno}_datasheet.docx"
            doc.save(out)
            files.append(out)
            if ctx.options.get("pdf", "yes") != "no":
                pdf = pdf_from_office(out, ctx)
                if pdf:
                    files.append(pdf)
        if not items:
            ctx.warnings.append("no equipment records - nothing generated")
        return files


def _set_text(paragraph, text):
    """Replace text but keep the template's run formatting (font size/colour)."""
    runs = paragraph.runs
    if not runs:
        paragraph.text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


def _widths(table, widths):
    """Word honours cell widths (not column widths); fixed layout keeps them in LibreOffice too."""
    table.autofit = False
    for col, w in zip(table.columns, widths):   # tblGrid (LibreOffice, PDF export)
        col.width = w
    for row in table.rows:                       # per-cell widths (Word)
        for cell, w in zip(row.cells, widths):
            cell.width = w


ENGINE = EquipmentDatasheets()
