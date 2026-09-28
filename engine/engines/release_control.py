"""Document release control: in which order, at which purpose and revision each document is released, what it waits for,
and which process gates (procurement, Owner acceptance, construction, commissioning, operations) it opens (Excel).

Derived from the records (engine/core/release.py and engine/core/workflow.py):
- release rules: a document is issued at a purpose only when its inputs have the maturity the rules require
  (IFR / IFA: doc_type.input_maturity; IFC: EPC inputs IFC, supplier inputs accepted, own Owner / EPC acceptance);
- revisions: letters before IFC (A, B, ...), numbers from IFC (0, 1, ...); a revision records the input revisions it was
  based on - an input revised later makes the document CHECK REQUIRED;
- gates (gate_rule): the documents and statuses each process step needs, with the need date from the plan (MR procurement
  dates, CWP starts, key dates, milestones); gate float = need - planned date all requirements are met.
Sheets: Rules, Gates, Gate requirements, Release plan (every document: status now, next issue, blocking inputs, planned
dates, gates it feeds), Relations (every input link with the maturity required at IFR and at IFC), Impacts (documents and
gates affected when a document is revised), Gate rules.
WARN: gate planned after its need date, mandatory gate requirement without documents, document CHECK REQUIRED."""
from __future__ import annotations

from collections import defaultdict
from datetime import date

from ..core import release, workflow
from ..core.runner import Context, Engine

PROCESS_TITLE = {"procurement": "Procurement", "owner_acceptance": "Owner acceptance", "construction": "Construction",
                 "commissioning": "Commissioning / acceptance", "operations": "Operations"}


def analyse(s, res=None):
    """(res, states, gate results, per-document release rows, relation rows, impact rows, warnings)."""
    res = res or workflow.compute(s)
    st = release.states(s)
    types = {t["id"]: t for t in s.records("doc_type")}
    docs = {d["id"]: d for d in s.records("document")}
    gs = release.gates(s, res, st)
    feeds = defaultdict(list)                       # document -> [(gate, scope, status, need, gate float)]
    for g in gs:
        for i, status in g.members:
            feeds[i].append((g.gate, g.scope_key, status, g.need, g.float))
    seq = release.release_sequence(s, res)
    wave = {k: w for w, k in seq}
    warn = []
    for g in gs:
        if g.float is not None and g.float < 0:
            warn.append(f"gate {g.gate} {g.scope_key}: planned {res.d(g.ready)} after the need date {res.d(g.need)} "
                        f"({g.float} wd)")
        for req in g.empty:
            warn.append(f"gate {g.gate} {g.scope_key}: requirement '{req}' selects no document")
    rows = []
    for _, k in seq:
        d, t, me = docs[k], res.docs[k], st[k]
        tp = types.get(d.get("type_code"), {})
        nxt = release.next_purpose(d, me, types)
        why = release.can_issue(s, k, nxt, st, docs, types) if nxt else []
        if me.suspect:
            warn.append(f"{k}: CHECK REQUIRED - inputs revised after its revision {me.rev} "
                        f"({', '.join(i for i, _ in me.suspect)})")
        f = feeds.get(k, [])
        first = min(f, key=lambda x: (x[3] is None, x[3] if x[3] is not None else 0)) if f else None
        rows.append({
            "id": k, "title": d.get("title"), "description": d.get("description"), "originator": d.get("originator", "EPC"), "discipline": d.get("discipline"),
            "type": d.get("type_code"), "review": release.review_class(d, types), "wave": wave.get(k),
            "status": release.NAME[me.level], "rev": me.rev, "code": me.code, "issued": me.issued,
            "action": release.next_action(d, me, types), "next": nxt,
            "next_rev": release.next_rev(me, nxt) if nxt else None,
            "can": ("yes" if not why else "no") if nxt else "-", "why": why,
            "suspect": [f"{i} rev {r.rsplit('_', 1)[-1]}" for i, r in me.suspect],
            "start": t.start, "ifr": t.ifr, "ifa": t.ifa, "ifc": t.ifc,
            "accepted": workflow.level_date(t, tp, "ACCEPTED"), "driver": t.driver,
            "gates": [f"{g} {sc} {stt}" for g, sc, stt, _, _ in f],
            "gate_need": first[3] if first else None, "gate_first": f"{first[0]} {first[1]}" if first else None,
        })
    rel = []
    for k in res.order:
        d = docs[k]
        for i in d.get("inputs") or []:
            if i not in res.docs:
                continue
            x = docs[i]
            ifr_lvl = release.required_input_level(d, "IFR", x, types)
            ifc_lvl = release.required_input_level(d, "IFC", x, types)
            ti = res.docs[i]
            ready_ifr = workflow.level_date(ti, types.get(x.get("type_code"), {}), release.NAME[ifr_lvl])
            rel.append([k, d.get("title"), i, x.get("title"), x.get("originator", "EPC"), release.NAME[ifr_lvl],
                        release.NAME[ifc_lvl], release.NAME[st[i].level], res.d(ready_ifr), res.d(res.docs[k].start),
                        None if ready_ifr is None else res.docs[k].start - ready_ifr])
    # impacts: documents downstream of each document (transitive) and the gates they feed
    succ = defaultdict(list)
    for k in res.order:
        for i in docs[k].get("inputs") or []:
            if i in res.docs:
                succ[i].append(k)
    down: dict[str, set] = {}
    for k in reversed(res.order):
        acc = set()
        for x in succ[k]:
            acc.add(x)
            acc |= down[x]
        down[k] = acc
    imp = []
    for k in res.order:
        dn = down[k]
        gset = {f"{g} {sc}" for x in dn | {k} for g, sc, _, _, _ in feeds.get(x, [])}
        if dn or gset:
            imp.append([k, docs[k].get("title"), len(succ[k]), len(dn),
                        len({docs[x].get("discipline") for x in dn}), len(gset),
                        ", ".join(sorted({g.split()[0] for g in gset}))])
    imp.sort(key=lambda r: (-r[3], r[0]))
    return res, st, gs, rows, rel, imp, warn


class ReleaseControl(Engine):
    name = "release_control"
    title = "Document release control: release order and revisions, prerequisites, process gates, impacts (Excel)"
    version = "1.1.0"
    inputs = ["project", "document", "document_revision", "doc_type", "gate_rule", "system", "mr", "cwp", "ewp", "cwa",
              "activity", "wbs", "milestone", "eng_resource", "kks_key"]
    formats = ["xlsx"]
    code_deps = ["engine/core/release.py", "engine/core/workflow.py", "engine/core/planning.py"]

    def run(self, ctx: Context):
        s = ctx.store
        if not s.records("document"):
            ctx.warnings.append("no document records - nothing generated")
            return []
        res, st, gs, rows, rel, imp, warn = analyse(s)
        ctx.warnings.extend(warn)
        return [self._workbook(ctx, res, gs, rows, rel, imp, warn, f"{ctx.project_record()['id']}_release_control")]

    def _workbook(self, ctx, res, gs, rows, rel, imp, warn, stem):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        s = ctx.store
        head = PatternFill("solid", fgColor="1F3864")
        red, amber, green = (PatternFill("solid", fgColor=c) for c in ("F8CBAD", "FFF2CC", "E2EFDA"))
        d = res.d
        wb = Workbook()

        def sheet(ws, title, cols, data, wrap=(), note=None):
            ws["A1"] = f"{ctx.project_record().get('name', '')} - {title}"
            ws["A1"].font = Font(bold=True, size=14)
            ws["A2"] = (f"engine {self.name} v{self.version} | generated {ctx.stamp['generated_at'][:19]}Z by "
                        f"{ctx.stamp['generated_by']} | data {ctx.stamp['inputs_hash']} | "
                        + (note or "planned dates from the levelled workflow network and the CPM (working days)"))
            ws["A2"].font = Font(italic=True, size=8, color="595959")
            for c, (n, w) in enumerate(cols, 1):
                cell = ws.cell(4, c, n)
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                ws.column_dimensions[get_column_letter(c)].width = w
            for i, row in enumerate(data, 5):
                for c, v in enumerate(row, 1):
                    v = "\n".join(map(str, v)) if isinstance(v, (list, tuple, set)) else v
                    cell = ws.cell(i, c, v)
                    cell.alignment = Alignment(wrap_text=c in wrap, vertical="top")
                    if isinstance(v, date):
                        cell.number_format = "yyyy-mm-dd"
            ws.freeze_panes = "B5"
            ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(4, 4 + len(data))}"
            ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            return ws

        # Rules (read me)
        ws = wb.active
        ws.title = "Rules"
        late = [g for g in gs if g.float is not None and g.float < 0]
        by_proc = defaultdict(lambda: [0, 0, 0])
        for g in gs:
            by_proc[g.process][0] += 1
            by_proc[g.process][1] += g.ok_now
            by_proc[g.process][2] += bool(g.float is not None and g.float < 0)
        text = [
            ["How to use", "Before issuing a document run 'python -m engine doc status <document>' (next purpose, revision "
                           "and what blocks it); issue with 'python -m engine doc issue <document> --purpose IFR|IFA|IFC|AB "
                           "--date ... --reason ...' (refused while blocked; --override records the justification); record the Owner / "
                           "EPC review with 'python -m engine doc review <document> --rev <revision> --code 1|2|3|4 --date ... --reason ...'. Gates: "
                           "'python -m engine gate [--process ...] [--scope ...] [--detail]'."],
            ["Status ladder", "NONE < IFR (first issue, also IFI / IFD / IFP) < IFA < ACCEPTED (review code 1 or 2 on the "
                              "latest issue; information and internal EPC documents are accepted when issued) < IFC "
                              "(issued for construction; final / certified for supplier documents) < AB (as built)"],
            ["Release rules", "IFR / IFA: every input at least at the maturity its type requires (doc_type.input_maturity "
                              "IFR or IFC); IFA needs a reviewed IFR. IFC: every EPC input IFC, every supplier input "
                              "ACCEPTED, the document itself ACCEPTED by the Owner (approval / review class) or by the EPC "
                              "(supplier documents), no code 3. AB: the document is IFC. An input in CHECK REQUIRED blocks."],
            ["Revisions", "Letters before IFC (A, B, C ...), numbers from IFC (0, 1, 2 ...). Every revision records the "
                          "input revisions it was based on; an input revised later makes the document CHECK REQUIRED "
                          "(re-issue, or confirm with a new revision of the same purpose)."],
            ["Release order", "Sheet 'Release plan' is in network order: a document comes after all its inputs (wave = "
                              "longest input chain). Within the plan the dates come from the levelled engineering timeline; "
                              "priorities follow the gate need dates (PO, CWP start, key dates, milestones)."],
            ["Gates", "A process step (enquiry, PO, fabrication, FAT, shipment, CWP / IWP release, mechanical completion, "
                      "start of commissioning, handover, Taking-Over, as-built) starts only when the documents its gate "
                      "requires have the required status. Gate float = need date - planned date all requirements are met."],
        ] + [[PROCESS_TITLE.get(p, p), f"{n} gate instance(s): {o} open now, {lt} planned late"]
             for p, (n, o, lt) in sorted(by_proc.items())] + [
            ["Findings", f"{len(warn)} WARN line(s): {len(late)} gate(s) planned late, "
                         f"{sum(1 for g in gs if g.empty)} gate(s) with a requirement selecting no document, "
                         f"{sum(1 for r in rows if r['suspect'])} document(s) CHECK REQUIRED"]] + [["WARN", w] for w in warn]
        sheet(ws, "Document release control - rules and summary", [("Item", 27), ("Rule / result", 150)], text, wrap=(2,),
              note="release rules: engine/core/release.py; gates: gate_rule records")
        # Gates
        gdesc = {g["id"]: g.get("description") for g in s.records("gate_rule")}
        data = []
        for g in sorted(gs, key=lambda g: (g.float if g.float is not None else 10 ** 6, g.gate, g.scope_key)):
            data.append([g.gate, PROCESS_TITLE.get(g.process, g.process), g.scope_key, g.title, gdesc.get(g.gate),
                         "OPEN" if g.ok_now else "blocked", d(g.need), d(g.ready), g.float,
                         len({i for i, _ in g.members}), sum(len(r[2]) for r in g.rows), g.empty])
        ws = sheet(wb.create_sheet("Gates"), "Process gates (sorted by float)",
                   [("Gate", 13), ("Process", 18), ("Scope", 14), ("Title", 44), ("Description", 40), ("Now", 8), ("Need", 11),
                    ("Planned ready", 11), ("Float wd", 8), ("Documents", 11), ("Not met now", 11),
                    ("Requirement without documents", 34)], data, wrap=(4, 5, 12))
        for i, row in enumerate(data, 5):
            fl = row[8]
            ws.cell(i, 9).fill = red if fl is not None and fl < 0 else (amber if fl is not None and fl < 10 else green)
        # Gate requirements
        data = []
        for g in sorted(gs, key=lambda g: (g.gate, g.scope_key)):
            for req, n, missing, when, drv in g.rows:
                data.append([g.gate, g.scope_key, req, n, len(missing), d(when), drv, d(g.need),
                             None if (when is None or g.need is None) else g.need - when])
            for req in g.empty:
                data.append([g.gate, g.scope_key, req, 0, 0, None, "NO DOCUMENT SELECTED", d(g.need), None])
        sheet(wb.create_sheet("Gate requirements"), "Gate requirements per gate instance",
              [("Gate", 13), ("Scope", 14), ("Requirement", 40), ("Documents", 11), ("Not met now", 11),
               ("Planned (last)", 11), ("Driving document", 30), ("Need", 11), ("Float wd", 8)], data)
        # Release plan
        data = [[r["wave"], r["id"], r["title"], r["description"], r["originator"], r["discipline"], r["type"], r["review"], r["status"],
                 r["rev"], r["code"], r["action"], r["next"], r["next_rev"], r["can"], r["why"][:6], r["suspect"],
                 d(r["start"]), d(r["ifr"]), d(r["ifa"]), d(r["accepted"]), d(r["ifc"]), r["driver"],
                 r["gate_first"], d(r["gate_need"]), r["gates"][:8]] for r in rows]
        ws = sheet(wb.create_sheet("Release plan"), "Release plan: every document in release order (network waves)",
                   [("Wave", 6), ("Document", 30), ("Title", 48), ("Description", 36), ("Orig.", 7), ("Discipline", 12), ("Type", 6),
                    ("Review", 10), ("Status now", 10), ("Rev", 5), ("Code", 5), ("Next action", 30), ("Next purpose", 8),
                    ("Next rev", 7), ("Can issue now", 7), ("Blocked by (first 6)", 60), ("Check required (inputs revised)", 30),
                    ("Start", 11), ("IFR", 11), ("IFA", 11), ("Accepted", 11), ("IFC", 11), ("Driver", 30),
                    ("First gate", 22), ("Gate need", 11), ("Gates fed (first 8)", 40)],
                   data, wrap=(3, 4, 16, 17, 26))
        for i, r in enumerate(rows, 5):
            ws.cell(i, 15).fill = green if r["can"] == "yes" else (red if r["can"] == "no" else amber)
            if r["suspect"]:
                ws.cell(i, 17).fill = red
        # Relations
        sheet(wb.create_sheet("Relations"), "Document relations: input links with the maturity required",
              [("Document", 30), ("Title", 44), ("Input", 30), ("Input title", 44), ("Input orig.", 8),
               ("Needed for IFR", 9), ("Needed for IFC", 9), ("Input status now", 9), ("Input planned at IFR level", 11),
               ("Document start", 11), ("Slack wd", 8)], rel, wrap=(2, 4))
        # Impacts
        sheet(wb.create_sheet("Impacts"), "Impact of a revision: documents downstream (transitive) and gates affected",
              [("Document", 30), ("Title", 50), ("Direct users", 8), ("All downstream", 9), ("Disciplines", 9),
               ("Gate instances", 9), ("Gates", 60)], imp, wrap=(2, 7))
        # Gate rules
        rules = sorted(s.records("gate_rule"), key=lambda g: g["id"])
        sheet(wb.create_sheet("Gate rules"), "Gate rules (gate_rule records)",
              [("Gate", 13), ("Process", 18), ("Scope", 8), ("Title", 44), ("Description", 40), ("Requires", 60), ("Need", 18),
               ("Offset wd", 8), ("Filter", 40), ("Basis", 30), ("Status", 9)],
              [[g["id"], PROCESS_TITLE.get(g["process"], g["process"]), g["scope"], g["title"], g.get("description"),
                g["requires"], g["need"],
                g.get("offset_days"), g.get("package_types") or g.get("categories") or g.get("parties"),
                g.get("basis_refs"), g.get("status")] for g in rules], wrap=(4, 5, 6, 9, 10))
        out = ctx.out_dir / f"{stem}.xlsx"
        wb.save(out)
        return out


ENGINE = ReleaseControl()
