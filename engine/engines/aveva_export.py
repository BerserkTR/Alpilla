"""AVEVA tag export (.xlsx): classed tags with their AVEVA class attributes, one sheet per AVEVA class.

For loading into AVEVA Engineering / E3D through their Excel import (map columns once in the AVEVA
import configuration; row 1 = attribute name, row 2 = AVEVA attribute ID, row 3 = unit).
Only records that carry an `aveva_class` are exported; the rest are listed on the 'Unclassified' sheet.
"""
from __future__ import annotations

import re

from ..core.runner import Context, Engine

ENTITIES = ["system", "equipment", "instrument", "line", "document"]


class AvevaExport(Engine):
    name = "aveva_export"
    title = "AVEVA Engineering / E3D tag import workbook (one sheet per AVEVA class)"
    version = "1.0.0"
    inputs = ENTITIES
    formats = ["xlsx"]

    def run(self, ctx: Context):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        s, lib = ctx.store, ctx.store.lib
        if lib is None:
            raise RuntimeError("AVEVA class library not built")
        entities = [e for e in ENTITIES if e in s.schemas]
        by_class: dict[str, list[tuple[str, dict]]] = {}
        unclassified = []
        for e in entities:
            for r in s.records(e):
                if r.get("aveva_class"):
                    by_class.setdefault(lib.resolve(r["aveva_class"])["id"], []).append((e, r))
                else:
                    unclassified.append((e, r["id"], r.get("description") or r.get("title", "")))

        wb = Workbook()
        idx = wb.active
        idx.title = "Index"
        bold, head = Font(bold=True), PatternFill("solid", fgColor="D9E1F2")
        idx.append([f"{ctx.project_record().get('name')} - AVEVA tag export"])
        idx.append([f"generated {ctx.stamp['generated_at'][:19]}Z by {ctx.stamp['generated_by']} | {self.name} "
                    f"v{self.version} | data {ctx.stamp['inputs_hash']} | class library source sha256 "
                    f"{_lib_sha(ctx)[:16]}"])
        idx.append([])
        idx.append(["Sheet", "AVEVA class", "AVEVA class ID", "Tags"])
        for c in idx[4]:
            c.font, c.fill = bold, head
        used_names = set()
        for cid in sorted(by_class, key=lambda i: lib.classes[i]["label"]):
            cls = lib.classes[cid]
            sheet = _sheet_name(cls["label"], used_names)
            ws = wb.create_sheet(sheet)
            idx.append([sheet, cls["label"], cid, len(by_class[cid])])
            eff = lib.effective(cls)
            names = sorted({k for _, r in by_class[cid] for k in (r.get("aveva_attrs") or {})})
            head1 = ["Name", "Entity", "Class", "Description"]
            head2 = ["", "", cid, ""]
            head3 = ["", "", "", ""]
            cols = []
            for n in names:
                a = eff[n]
                units = sorted({(r.get("aveva_attrs") or {}).get(n, {}).get("unit") for _, r in by_class[cid]
                                if isinstance((r.get("aveva_attrs") or {}).get(n), dict)} - {None})
                if a["type"] == "quantity" and len(units) > 1:
                    for u in units:  # mixed units are never silently converted: one column per unit
                        head1.append(a["label"]); head2.append(a["id"]); head3.append(u); cols.append((n, u))
                else:
                    head1.append(a["label"]); head2.append(a["id"]); head3.append(units[0] if units else "")
                    cols.append((n, units[0] if units else None))
            for row in (head1, head2, head3):
                ws.append(row)
            for c in ws[1]:
                c.font, c.fill = bold, head
            for c in list(ws[2]) + list(ws[3]):
                c.font = Font(italic=True, size=8, color="595959")
            for e, r in sorted(by_class[cid], key=lambda x: x[1]["id"]):
                attrs = r.get("aveva_attrs") or {}
                row = [r["id"], e, cls["label"], r.get("description") or r.get("title", "")]
                for n, u in cols:
                    v = attrs.get(n)
                    if isinstance(v, dict):
                        v = v["value"] if v.get("unit") == u else None
                    row.append(v)
                ws.append(row)
            ws.freeze_panes = "B4"
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = max(12, min(40, len(str(col[0].value or "")) + 2))
        if unclassified:
            ws = wb.create_sheet("Unclassified")
            ws.append(["Name", "Entity", "Description", "Action"])
            for c in ws[1]:
                c.font, c.fill = bold, head
            for e, rid, d in unclassified:
                ws.append([rid, e, d, "set aveva_class (python -m engine lib find <text>)"])
            ctx.warnings.append(f"{len(unclassified)} record(s) without aveva_class not exported (see 'Unclassified')")
        for col, w in zip("ABCD", (26, 34, 26, 8)):
            idx.column_dimensions[col].width = w
        out = ctx.out_dir / f"{ctx.project_record()['id']}_aveva_tag_export.xlsx"
        wb.save(out)
        return [out]


def _lib_sha(ctx) -> str:
    import json
    m = ctx.project.database / "classlib" / "manifest.json"
    return json.loads(m.read_text(encoding="utf-8"))["source_sha256"] if m.exists() else "n/a"


def _sheet_name(label: str, used: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", label)[:28]
    name, n = base, 2
    while name.lower() in used:
        name, n = f"{base[:25]}~{n}", n + 1
    used.add(name.lower())
    return name


ENGINE = AvevaExport()
