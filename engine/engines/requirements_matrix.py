"""Requirements traceability matrix (.xlsx): every contract requirement with its verification method,
acceptance criterion, guarantee, and the database records that trace to it (design parameters, equipment,
anything else citing `requirement:<id>` in basis_refs). Coverage is computed, never typed.
Sheets: RTM, Guarantees, Scope split, Coverage."""
from __future__ import annotations

from ..core.runner import Context, Engine


class RequirementsMatrix(Engine):
    name = "requirements_matrix"
    title = "Requirements traceability matrix (Excel): requirements -> guarantees, design data, equipment"
    version = "1.0.0"
    inputs = ["contract", "requirement", "guarantee", "scope_item", "party", "design_parameter", "equipment", "line"]
    formats = ["xlsx"]

    def run(self, ctx: Context):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        s = ctx.store
        reqs = sorted(s.records("requirement"), key=lambda r: r["id"])
        if not reqs:
            ctx.warnings.append("no requirement records")
            return []
        traces: dict[str, list[str]] = {}
        for ent, sch in s.schemas.items():
            if "basis_refs" not in sch.fields:
                continue
            for r in s.records(ent):
                for ref in r.get("basis_refs", []):
                    if ref.startswith("requirement:"):
                        traces.setdefault(ref.split(":", 1)[1], []).append(f"{ent}/{r['id']}")
        guar = {g["id"]: g for g in s.records("guarantee")}
        head, fill = PatternFill("solid", fgColor="1F3864"), PatternFill("solid", fgColor="F8D7D7")
        wb = Workbook()

        def sheet(ws, title, cols, widths):
            ws["A1"] = f"{ctx.project_record().get('name', '')} - {title}"
            ws["A1"].font = Font(bold=True, size=14)
            ws["A2"] = (f"engine {self.name} v{self.version} | generated {ctx.stamp['generated_at'][:19]}Z by "
                        f"{ctx.stamp['generated_by']} | data {ctx.stamp['inputs_hash']}")
            ws["A2"].font = Font(italic=True, size=8, color="595959")
            for c, n in enumerate(cols, 1):
                cell = ws.cell(4, c, n)
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for c, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(c)].width = w
            ws.freeze_panes = "B5"
            ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}4"
            ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.print_title_rows = "4:4"

        ws = wb.active
        ws.title = "RTM"
        sheet(ws, "Requirements Traceability Matrix",
              ["ID", "Section", "Title", "Requirement", "Discipline", "Verification", "Acceptance", "Guarantee",
               "Guaranteed value", "Traced by (database records)", "Traced", "Status"],
              [11, 26, 26, 70, 11, 12, 28, 9, 16, 34, 8, 12])
        for i, r in enumerate(reqs, 5):
            g = guar.get(r.get("guarantee"))
            gv = f"{'>=' if g['direction'] == 'min' else '<='} {g['guaranteed_value']:g} {g['unit']}" if g else ""
            tr = traces.get(r["id"], [])
            row = [r["id"], r["section"], r["title"], r["text"], r["discipline"], r["verification"], r.get("acceptance"),
                   r.get("guarantee"), gv, "\n".join(tr), "yes" if tr else "no", r.get("status")]
            for c, v in enumerate(row, 1):
                cell = ws.cell(i, c, v)
                cell.alignment = Alignment(wrap_text=c in (2, 3, 4, 7, 10), vertical="top")
            if not tr:
                ws.cell(i, 11).fill = fill

        ws = wb.create_sheet("Guarantees")
        sheet(ws, "Guarantees", ["ID", "Parameter", "Category", "Value", "Unit", "Direction", "Minimum acceptance",
                                 "LD rate", "LD unit", "Remedy", "Requirements"], [8, 38, 13, 10, 12, 9, 12, 12, 26, 40, 22])
        for i, g in enumerate(sorted(guar.values(), key=lambda g: g["id"]), 5):
            linked = ", ".join(r["id"] for r in reqs if r.get("guarantee") == g["id"])
            for c, v in enumerate([g["id"], g["parameter"], g["category"], g["guaranteed_value"], g["unit"], g["direction"],
                                   g.get("minimum_acceptance"), g.get("ld_rate"), g.get("ld_unit"), g.get("remedy"), linked], 1):
                ws.cell(i, c, v).alignment = Alignment(wrap_text=c in (2, 10), vertical="top")

        ws = wb.create_sheet("Scope split")
        sheet(ws, "Scope split and terminal points", ["ID", "Area", "Item", "Design", "Supply", "Install", "Commission",
                                                      "Terminal point"], [8, 16, 60, 9, 9, 9, 11, 40])
        for i, x in enumerate(sorted(s.records("scope_item"), key=lambda x: x["id"]), 5):
            for c, v in enumerate([x["id"], x["area"], x["item"], x.get("design"), x.get("supply"), x.get("install"),
                                   x.get("commission"), x.get("terminal_point")], 1):
                ws.cell(i, c, v).alignment = Alignment(wrap_text=c in (3, 8), vertical="top")

        ws = wb.create_sheet("Coverage")
        sheet(ws, "Coverage by section", ["Section", "Requirements", "With guarantee", "Traced to data", "Traced %",
                                          "test", "inspection", "analysis", "review", "demonstration"],
              [42, 13, 14, 14, 10, 8, 10, 9, 8, 13])
        sections = sorted({r["section"] for r in reqs})
        for i, sec in enumerate(sections, 5):
            rs = [r for r in reqs if r["section"] == sec]
            n_tr = sum(1 for r in rs if traces.get(r["id"]))
            vals = [sec, len(rs), sum(1 for r in rs if r.get("guarantee")), n_tr, round(100 * n_tr / len(rs), 1)]
            vals += [sum(1 for r in rs if r["verification"] == v) for v in ("test", "inspection", "analysis", "review", "demonstration")]
            for c, v in enumerate(vals, 1):
                ws.cell(i, c, v)
        total_tr = sum(1 for r in reqs if traces.get(r["id"]))
        ws.cell(5 + len(sections), 1, "Total").font = Font(bold=True)
        ws.cell(5 + len(sections), 2, len(reqs)).font = Font(bold=True)
        ws.cell(5 + len(sections), 4, total_tr).font = Font(bold=True)
        ws.cell(5 + len(sections), 5, round(100 * total_tr / len(reqs), 1)).font = Font(bold=True)
        out = ctx.out_dir / f"{ctx.project_record()['id']}_requirements_matrix.xlsx"
        wb.save(out)
        ctx.warnings.append(f"{total_tr} of {len(reqs)} requirements traced to database records "
                            f"({len(reqs) - total_tr} open - trace them as design data is added)")
        return [out]


ENGINE = RequirementsMatrix()
