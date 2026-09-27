"""Clarification (technical query) register (.xlsx): all TQs with rounds, answers, outcomes, price effects and
the records each answer changed; summary by round and by outcome."""
from __future__ import annotations

from ..core.runner import Context, Engine


class ClarificationRegister(Engine):
    name = "clarification_register"
    title = "Technical query / clarification register (Excel) with rounds, outcomes and price effects"
    version = "1.0.0"
    inputs = ["clarification", "party", "contract"]
    formats = ["xlsx"]

    def run(self, ctx: Context):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        s = ctx.store
        tqs = sorted(s.records("clarification"), key=lambda q: q["id"])
        if not tqs:
            ctx.warnings.append("no clarification records")
            return []
        parties = {p["id"]: p for p in s.records("party")}
        head = PatternFill("solid", fgColor="1F3864")
        wb = Workbook()
        ws = wb.active
        ws.title = "Register"
        ws["A1"] = f"{ctx.project_record().get('name', '')} - Technical Query Register"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = (f"engine {self.name} v{self.version} | generated {ctx.stamp['generated_at'][:19]}Z by {ctx.stamp['generated_by']} "
                    f"| data {ctx.stamp['inputs_hash']}")
        ws["A2"].font = Font(italic=True, size=8, color="595959")
        cols = ["TQ", "Round", "Follows", "Raised by", "Raised", "Discipline", "Subject", "Question", "Contractor proposal",
                "References", "Owner's answer", "Answered", "Status", "Outcome", "Impact", "Price effect (EUR)", "Records changed"]
        widths = [8, 7, 8, 18, 11, 12, 32, 60, 45, 30, 60, 11, 9, 20, 13, 13, 34]
        for c, (n, w) in enumerate(zip(cols, widths), 1):
            cell = ws.cell(4, c, n)
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = w
        for i, q in enumerate(tqs, 5):
            row = [q["id"], q["round"], q.get("follows"), parties.get(q["raised_by"], {}).get("short_name", q["raised_by"]),
                   q["raised_date"], q["discipline"], q["subject"], q["question"], q.get("proposal"),
                   "\n".join(q.get("references", [])), q.get("response"), q.get("response_date"), q.get("status"),
                   (q.get("outcome") or "").replace("_", " "), q.get("impact"), q.get("cost_impact"),
                   "\n".join(q.get("changed_records", []))]
            for c, v in enumerate(row, 1):
                ws.cell(i, c, v).alignment = Alignment(wrap_text=c in (7, 8, 9, 10, 11, 17), vertical="top")
        ws.freeze_panes = "B5"
        ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}4"
        ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        sm = wb.create_sheet("Summary")
        sm.append(["By round", "Raised", "Closed", "Open", "Price effect (EUR)"])
        for r in sorted({q["round"] for q in tqs}):
            qs = [q for q in tqs if q["round"] == r]
            sm.append([f"Round {r}", len(qs), sum(1 for q in qs if q.get("status") == "closed"),
                       sum(1 for q in qs if q.get("status") not in ("closed", "superseded")), sum(q.get("cost_impact") or 0 for q in qs)])
        sm.append(["Total", len(tqs), sum(1 for q in tqs if q.get("status") == "closed"),
                   sum(1 for q in tqs if q.get("status") not in ("closed", "superseded")), sum(q.get("cost_impact") or 0 for q in tqs)])
        sm.append([])
        sm.append(["By outcome", "Count"])
        for o in sorted({q.get("outcome") or "none" for q in tqs}):
            sm.append([o.replace("_", " "), sum(1 for q in tqs if (q.get("outcome") or "none") == o)])
        for row in sm.iter_rows():
            if row[0].value in ("By round", "By outcome", "Total"):
                for c in row:
                    c.font = Font(bold=True)
        for col, w in zip("ABCDE", (32, 10, 10, 10, 20)):
            sm.column_dimensions[col].width = w
        out = ctx.out_dir / f"{ctx.project_record()['id']}_clarification_register.xlsx"
        wb.save(out)
        return [out]


ENGINE = ClarificationRegister()
