"""Project schedule from activity records: CPM dates/float (engine/core/planning.py), outputs
  <prj>_schedule.xlsx   activity table + weekly Gantt + critical path sheet
  <prj>_schedule.html   interactive Gantt (hover tooltips), light/dark
  <prj>_schedule.xml    Microsoft Project XML (MSPDI) - opens in MS Project, imports into Primavera P6
Dates are computed, never stored in the database.
"""
from __future__ import annotations

import html
from datetime import date, timedelta
from xml.sax.saxutils import escape

from ..core import planning
from ..core.runner import Context, Engine

REL_MSP = {"FF": 0, "FS": 1, "SF": 2, "SS": 3}


def wbs_tree(store, acts=None):
    """WBS nodes in display order: siblings by earliest early start of their subtree, then code."""
    nodes = {w["id"]: w for w in store.records("wbs")}
    kids: dict = {}
    for w in nodes.values():
        kids.setdefault(w.get("parent"), []).append(w["id"])
    first: dict = {}
    for a in acts or []:
        k = a.wbs
        while k:
            first[k] = min(first.get(k, a.es), a.es)
            k = nodes[k].get("parent")
    order = []

    def walk(pid, depth):
        for k in sorted(kids.get(pid, []), key=lambda k: (first.get(k, 10 ** 9), k)):
            order.append((k, depth))
            walk(k, depth + 1)
    walk(None, 0)
    return nodes, order


def rows(acts, cal, store):
    """Display rows: WBS summaries followed by their activities (sorted by early start)."""
    nodes, order = wbs_tree(store, acts)
    by_wbs: dict = {}
    for a in acts:
        by_wbs.setdefault(a.wbs, []).append(a)

    def span(wid):
        sub = [x for x in by_wbs.get(wid, [])]
        for k, _ in order:
            if _descends(nodes, k, wid) and k != wid:
                sub += by_wbs.get(k, [])
        return sub
    out = []
    for wid, depth in order:
        sub = span(wid)
        if not sub:
            continue
        out.append({"kind": "wbs", "id": wid, "title": nodes[wid]["title"], "depth": depth,
                    "start": min(cal.date(a.es) for a in sub),
                    "finish": max(cal.finish_date(a.ef, a.es) for a in sub)})
        for a in sorted(by_wbs.get(wid, []), key=lambda x: (x.es, x.id)):
            out.append({"kind": "act", "a": a, "depth": depth + 1, "start": cal.date(a.es),
                        "finish": cal.finish_date(a.ef, a.es)})
    return out


def _descends(nodes, k, anc):
    while k:
        if k == anc:
            return True
        k = nodes[k].get("parent")
    return False


class Schedule(Engine):
    name = "schedule"
    title = "CPM schedule: Excel Gantt, interactive HTML Gantt, MS Project XML (P6 import)"
    version = "1.0.0"
    inputs = ["project", "wbs", "activity", "document", "document_revision", "progress_rule"]
    formats = ["xlsx", "html", "xml"]

    def run(self, ctx: Context):
        s = ctx.store
        if not s.records("activity"):
            ctx.warnings.append("no activity records - nothing scheduled")
            return []
        acts, cal, dd = planning.compute(s)
        prj = ctx.project_record()
        rs = rows(acts, cal, s)
        base = ctx.out_dir / f"{prj['id']}_schedule"
        files = [self.xlsx(ctx, prj, acts, rs, cal, dd, base.with_suffix(".xlsx")),
                 self.html(ctx, prj, acts, rs, cal, dd, base.with_suffix(".html")),
                 self.mspdi(ctx, prj, acts, cal, dd, s, base.with_suffix(".xml"))]
        late = [a.id for a in acts if a.tf < 0]
        if late:
            ctx.warnings.append(f"negative float (constraint missed) on: {', '.join(late[:10])}")
        return files

    # ------------------------------------------------------------------ xlsx
    def xlsx(self, ctx, prj, acts, rs, cal, dd, out):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Schedule"
        ws["A1"] = f"{prj.get('name')} - Project Schedule (data date {dd.isoformat()})"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = (f"CPM by engine {self.name} v{self.version} | generated {ctx.stamp['generated_at'][:19]}Z by "
                    f"{ctx.stamp['generated_by']} | data {ctx.stamp['inputs_hash']} | {cal.dpw}-day week, working days")
        ws["A2"].font = Font(italic=True, size=8, color="595959")
        cols = ["WBS / ID", "Title", "Type", "Dur [wd]", "Start", "Finish", "Late start", "Late finish",
                "Total float [wd]", "Critical", "% complete", "Actual start", "Actual finish", "Predecessors"]
        h = 4
        start = min(r["start"] for r in rs)
        finish = max(r["finish"] for r in rs)
        week0 = start - timedelta(days=start.weekday())
        weeks = []
        w = week0
        while w <= finish:
            weeks.append(w)
            w += timedelta(days=7)
        head = PatternFill("solid", fgColor="1F3864")
        for c, name in enumerate(cols + [x.strftime("%d-%b-%y") for x in weeks], 1):
            cell = ws.cell(h, c, name)
            cell.font = Font(bold=True, color="FFFFFF", size=9 if c > len(cols) else 10)
            cell.fill = head
            cell.alignment = Alignment(text_rotation=90 if c > len(cols) else 0, wrap_text=True, vertical="center")
        fills = {"crit": PatternFill("solid", fgColor="D03B3B"), "norm": PatternFill("solid", fgColor="2A78D6"),
                 "done": PatternFill("solid", fgColor="C3C2B7"), "wbs": PatternFill("solid", fgColor="898781"),
                 "band": PatternFill("solid", fgColor="E7ECF5"), "dd": PatternFill("solid", fgColor="FAB219")}
        row = h + 1
        for r in rs:
            if r["kind"] == "wbs":
                vals = ["  " * r["depth"] + r["id"], r["title"], "WBS", None, r["start"], r["finish"]]
                for c, v in enumerate(vals, 1):
                    ws.cell(row, c, v).font = Font(bold=True)
                for c in range(1, len(cols) + 1):
                    ws.cell(row, c).fill = fills["band"]
                kind = "wbs"
            else:
                a = r["a"]
                vals = ["  " * r["depth"] + a.id, a.title, a.type.replace("_", " "), a.dur, r["start"], r["finish"],
                        None if a.status == "completed" else cal.date(a.ls),
                        None if a.status == "completed" else cal.finish_date(a.lf, a.ls),
                        "done" if a.status == "completed" else a.tf, "CRIT" if a.critical else "",
                        round(a.pct, 1), a.rec.get("actual_start"), a.rec.get("actual_finish"),
                        ", ".join(a.rec.get("predecessors", []))]
                for c, v in enumerate(vals, 1):
                    ws.cell(row, c, v)
                if a.critical:
                    ws.cell(row, 10).font = Font(bold=True, color="D03B3B")
                kind = "done" if a.status == "completed" else "crit" if a.critical else "norm"
            for i, wk in enumerate(weeks):
                if r["start"] <= wk + timedelta(days=6) and r["finish"] >= wk:
                    cell = ws.cell(row, len(cols) + 1 + i)
                    cell.fill = fills[kind]
                    if r["kind"] == "act" and r["a"].type != "task":
                        cell.value, cell.fill = "◆", PatternFill(fill_type=None)
                elif wk <= dd <= wk + timedelta(days=6):
                    ws.cell(row, len(cols) + 1 + i).fill = fills["dd"]
            row += 1
        for c, wd in enumerate([18, 38, 11, 8, 11, 11, 11, 11, 9, 8, 9, 11, 11, 22], 1):
            ws.column_dimensions[get_column_letter(c)].width = wd
        for i in range(len(weeks)):
            ws.column_dimensions[get_column_letter(len(cols) + 1 + i)].width = 3
        for rr in ws.iter_rows(min_row=h + 1, max_row=row, min_col=5, max_col=13):
            for cell in rr:
                if isinstance(cell.value, date):
                    cell.number_format = "dd-mmm-yy"
        ws.freeze_panes = ws.cell(h + 1, 3)
        ws.row_dimensions[h].height = 60
        ws.cell(row + 1, 1, "Legend: red = critical (CRIT), blue = non-critical, grey = completed, "
                            "yellow column = data date, ◆ = milestone").font = Font(italic=True, size=9)
        ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        cp = wb.create_sheet("Critical path")
        cp.append(["ID", "Title", "Start", "Finish", "Total float [wd]"])
        for c in cp[1]:
            c.font = Font(bold=True)
        for a in sorted((a for a in acts if a.critical), key=lambda a: (a.es, a.id)):
            cp.append([a.id, a.title, cal.date(a.es), cal.finish_date(a.ef, a.es), a.tf])
        for col, wd in zip("ABCDE", (16, 44, 12, 12, 14)):
            cp.column_dimensions[col].width = wd
        wb.save(out)
        return out

    # ------------------------------------------------------------------ html
    def html(self, ctx, prj, acts, rs, cal, dd, out):
        import jinja2
        t0 = min(r["start"] for r in rs) - timedelta(days=7)
        t1 = max(r["finish"] for r in rs) + timedelta(days=14)
        span = (t1 - t0).days
        pxd = max(1.0, min(14.0, 1100.0 / span))
        LW, RH, TOP = 430, 24, 44
        W = LW + int(span * pxd) + 20
        H = TOP + RH * len(rs) + 10
        X = lambda d: LW + (d - t0).days * pxd
        e = html.escape
        svg = [f'<svg class="gantt" width="{W}" height="{H}" viewBox="0 0 {W} {H}" role="img" '
               f'aria-label="Gantt chart of {len(acts)} activities">']
        # month grid + labels
        m = date(t0.year, t0.month, 1)
        while m <= t1:
            if m >= t0:
                x = X(m)
                svg.append(f'<line x1="{x:.1f}" y1="{TOP - 16}" x2="{x:.1f}" y2="{H}" stroke="var(--grid)" stroke-width="1"/>')
                svg.append(f'<text x="{x + 3:.1f}" y="{TOP - 20}" fill="var(--muted)" font-size="11">{m.strftime("%b %y")}</text>')
            m = date(m.year + (m.month == 12), m.month % 12 + 1, 1)
        svg.append(f'<text x="4" y="{TOP - 20}" fill="var(--text-secondary)" font-size="11" font-weight="600">ID / Activity</text>')
        svg.append(f'<text x="{LW - 118}" y="{TOP - 20}" fill="var(--text-secondary)" font-size="11" font-weight="600">Start</text>')
        svg.append(f'<text x="{LW - 58}" y="{TOP - 20}" fill="var(--text-secondary)" font-size="11" font-weight="600">Float</text>')
        svg.append(f'<line x1="0" y1="{TOP - 6}" x2="{W}" y2="{TOP - 6}" stroke="var(--axis)"/>')
        for i, r in enumerate(rs):
            y = TOP + i * RH
            cy = y + RH / 2
            if r["kind"] == "wbs":
                svg.append(f'<text x="{4 + 12 * r["depth"]}" y="{cy + 4}" fill="var(--text-primary)" font-size="12" '
                           f'font-weight="600">{e(r["id"])}  {e(r["title"][:34])}</text>')
                x0, x1 = X(r["start"]), X(r["finish"] + timedelta(days=1))
                svg.append(f'<rect x="{x0:.1f}" y="{cy - 2}" width="{max(2, x1 - x0):.1f}" height="4" rx="2" fill="var(--muted)" '
                           f'data-tip="{e(r["id"])} {e(r["title"])}&#10;{r["start"]} to {r["finish"]}"/>')
                continue
            a = r["a"]
            label = f'{a.id}  {a.title}'
            svg.append(f'<text x="{4 + 12 * r["depth"]}" y="{cy + 4}" fill="var(--text-primary)" font-size="12">'
                       f'{e(label[:44 - r["depth"] * 2])}</text>')
            svg.append(f'<text x="{LW - 118}" y="{cy + 4}" fill="var(--text-secondary)" font-size="11">'
                       f'{r["start"].strftime("%d-%b-%y")}</text>')
            fl = "done" if a.status == "completed" else "CRIT" if a.critical else str(a.tf)
            style = 'fill="var(--critical)" font-weight="600"' if a.critical else 'fill="var(--text-secondary)"'
            svg.append(f'<text x="{LW - 58}" y="{cy + 4}" font-size="11" {style}>{fl}</text>')
            tip = (f"{a.id} {a.title}&#10;{r['start']} to {r['finish']} ({a.dur} wd)&#10;"
                   f"total float {a.tf} wd{' - CRITICAL' if a.critical else ''}&#10;{a.pct:.0f}% complete ({a.status.replace('_', ' ')})")
            color = "var(--done)" if a.status == "completed" else "var(--critical)" if a.critical else "var(--series-1)"
            if a.type != "task":
                x = X(r["start"])
                svg.append(f'<path d="M{x:.1f},{cy - 7} l7,7 l-7,7 l-7,-7 z" fill="{color}" stroke="var(--surface-1)" '
                           f'stroke-width="2" data-tip="{e(tip)}"/>')
            else:
                x0, x1 = X(r["start"]), X(r["finish"] + timedelta(days=1))
                wbar = max(3.0, x1 - x0 - 1)
                svg.append(f'<rect x="{x0:.1f}" y="{cy - 6}" width="{wbar:.1f}" height="12" rx="3" fill="{color}" '
                           f'data-tip="{e(tip)}"/>')
                if 0 < a.pct < 100:
                    svg.append(f'<rect x="{x0:.1f}" y="{cy + 3}" width="{wbar * a.pct / 100:.1f}" height="3" rx="1.5" '
                               f'fill="var(--text-primary)" opacity="0.55" pointer-events="none"/>')
        xd = X(dd)
        svg.append(f'<line x1="{xd:.1f}" y1="{TOP - 16}" x2="{xd:.1f}" y2="{H}" stroke="var(--text-secondary)" '
                   f'stroke-width="1.5" stroke-dasharray="4 3"/>')
        svg.append(f'<text x="{xd + 4:.1f}" y="{TOP - 8}" fill="var(--text-secondary)" font-size="11">data date {dd.strftime("%d-%b-%y")}</text>')
        svg.append("</svg>")
        crit = [a for a in acts if a.critical]
        finish = max(cal.finish_date(a.ef, a.es) for a in acts)
        env = jinja2.Environment(loader=jinja2.FileSystemLoader(ctx.project.templates), autoescape=False)
        text = env.get_template("html/schedule.html.j2").render(
            prj=prj, stamp=ctx.stamp, svg="\n".join(svg), dd=dd, finish=finish, n=len(acts), ncrit=len(crit),
            late=[a for a in acts if a.tf < 0], engine=self)
        out.write_text(text, encoding="utf-8")
        return out

    # ------------------------------------------------------------------ MSPDI
    def mspdi(self, ctx, prj, acts, cal, dd, store, out):
        nodes, order = wbs_tree(store, acts)
        byid = {a.id: a for a in acts}
        uid = {}
        tasks = []
        n = 0

        def dt(d, t):
            return f"{d.isoformat()}T{t}"
        for wid, depth in order:
            n += 1
            uid["WBS:" + wid] = n
            sub = [a for a in acts if _descends(nodes, a.wbs, wid)]
            if not sub:
                n -= 1
                continue
            tasks.append(("wbs", wid, depth + 1, n, sub))
            for a in sorted((x for x in acts if x.wbs == wid), key=lambda x: (x.es, x.id)):
                n += 1
                uid[a.id] = n
                tasks.append(("act", a.id, depth + 2, n, None))
        x = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
             '<Project xmlns="http://schemas.microsoft.com/project">',
             f"<SaveVersion>14</SaveVersion><Name>{escape(prj['id'])}_schedule</Name><Title>{escape(prj.get('name', ''))}</Title>",
             f"<ScheduleFromStart>1</ScheduleFromStart><StartDate>{dt(cal.start, '08:00:00')}</StartDate>",
             f"<StatusDate>{dt(dd, '17:00:00')}</StatusDate><CalendarUID>1</CalendarUID>",
             "<DefaultStartTime>08:00:00</DefaultStartTime><DefaultFinishTime>17:00:00</DefaultFinishTime>",
             f"<MinutesPerDay>480</MinutesPerDay><MinutesPerWeek>{480 * cal.dpw}</MinutesPerWeek><DaysPerMonth>{4 * cal.dpw}</DaysPerMonth>",
             "<DurationFormat>7</DurationFormat><WorkFormat>2</WorkFormat>",
             "<Calendars><Calendar><UID>1</UID><Name>Project calendar</Name><IsBaseCalendar>1</IsBaseCalendar><WeekDays>"]
        for msday in range(1, 8):                       # 1 = Sunday ... 7 = Saturday
            pyday = (msday + 5) % 7                     # python: Monday = 0
            working = pyday < cal.dpw
            x.append(f"<WeekDay><DayType>{msday}</DayType><DayWorking>{int(working)}</DayWorking>"
                     + ("<WorkingTimes><WorkingTime><FromTime>08:00:00</FromTime><ToTime>12:00:00</ToTime></WorkingTime>"
                        "<WorkingTime><FromTime>13:00:00</FromTime><ToTime>17:00:00</ToTime></WorkingTime></WorkingTimes>"
                        if working else "") + "</WeekDay>")
        for h in sorted(cal.holidays):
            x.append(f"<WeekDay><DayType>0</DayType><DayWorking>0</DayWorking><TimePeriod><FromDate>{dt(h, '00:00:00')}"
                     f"</FromDate><ToDate>{dt(h, '23:59:00')}</ToDate></TimePeriod></WeekDay>")
        x.append("</WeekDays></Calendar></Calendars><Tasks>")
        x.append(f"<Task><UID>0</UID><ID>0</ID><Name>{escape(prj.get('name', prj['id']))}</Name><OutlineLevel>0</OutlineLevel>"
                 "<Summary>1</Summary></Task>")
        for i, (kind, key, level, u, sub) in enumerate(tasks, 1):
            if kind == "wbs":
                st = min(cal.date(a.es) for a in sub)
                fi = max(cal.finish_date(a.ef, a.es) for a in sub)
                x.append(f"<Task><UID>{u}</UID><ID>{i}</ID><Name>{escape(nodes[key]['title'])}</Name><WBS>{escape(key)}</WBS>"
                         f"<OutlineLevel>{level}</OutlineLevel><Summary>1</Summary><Start>{dt(st, '08:00:00')}</Start>"
                         f"<Finish>{dt(fi, '17:00:00')}</Finish></Task>")
                continue
            a = byid[key]
            st, fi = cal.date(a.es), cal.finish_date(a.ef, a.es)
            ms = a.type != "task"
            parts = [f"<UID>{u}</UID><ID>{i}</ID><Name>{escape(a.id + ' ' + a.title)}</Name><WBS>{escape(a.wbs)}</WBS>",
                     f"<OutlineLevel>{level}</OutlineLevel><Summary>0</Summary><Milestone>{int(ms)}</Milestone>",
                     f"<Start>{dt(st, '08:00:00')}</Start><Finish>{dt(fi, '08:00:00' if ms else '17:00:00')}</Finish>",
                     f"<Duration>PT{a.dur * 8}H0M0S</Duration><DurationFormat>7</DurationFormat>",
                     f"<PercentComplete>{int(round(a.pct))}</PercentComplete><Critical>{int(a.critical)}</Critical>"]
            r = a.rec
            if r.get("constraint") and r.get("constraint_date"):
                ctype = 4 if r["constraint"] == "start_no_earlier_than" else 7
                ctime = "08:00:00" if ctype == 4 else "17:00:00"
                parts.append(f"<ConstraintType>{ctype}</ConstraintType><ConstraintDate>{r['constraint_date']}T{ctime}</ConstraintDate>")
            if r.get("actual_start"):
                parts.append(f"<ActualStart>{r['actual_start']}T08:00:00</ActualStart>")
            if r.get("actual_finish"):
                parts.append(f"<ActualFinish>{r['actual_finish']}T17:00:00</ActualFinish>")
            for pid, rel, lag in a.preds:
                parts.append(f"<PredecessorLink><PredecessorUID>{uid[pid]}</PredecessorUID><Type>{REL_MSP[rel]}</Type>"
                             f"<CrossProject>0</CrossProject><LinkLag>{lag * 4800}</LinkLag><LagFormat>7</LagFormat></PredecessorLink>")
            x.append("<Task>" + "".join(parts) + "</Task>")
        x.append("</Tasks></Project>")
        out.write_text("\n".join(x) + "\n", encoding="utf-8")
        return out


ENGINE = Schedule()
