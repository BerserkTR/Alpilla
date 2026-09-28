"""Tie-in / terminal point register (Excel + Word + PDF) with the interface checks.

Every tie_in record (TP-: Owner/Contractor terminal points, IF-: interfaces between consortium members) is shown as a
data card (scope break, connection, isolation, protection, design and operating envelope, flow, electrical data,
quality, dates, agreement). The engine checks:
- completeness per category (a piped point needs size, connection, design and operating p/T, flow, isolation, quality);
- the HMB envelope over all non-superseded hmb_case records: the process stream crossing the point (hmb_stream) must stay
  inside the agreed operating range, below the design values and below the flow capacity; electrical points with an
  hmb_metric check the GT/ST/gross/net power against their rating (MW, or MVA at power_factor);
- the chain to the upstream point: minimum upstream pressure less chain_dp covers this point's minimum pressure, flow
  capacity, voltage range, short-circuit level, and an upstream design pressure/temperature above this point's design
  requires a declared protection;
- dates: the providing side's available_by is not later than the receiving side's needed_by;
- agreement: agreed/frozen points name at least two agreeing parties; draft points are listed as open;
- coverage: every Owner-supplied scope item with a terminal point is served by a tie-in.
Pressures are barg (HMB streams are bar(a): barg = bar(a) - 1.01325). Options: pdf=no."""
from __future__ import annotations

from datetime import date

from ..core import hmb
from ..core.runner import Context, Engine, pdf_from_office

P_ATM = 1.01325
CATEGORY_ORDER = ["process", "utility", "electrical", "control", "civil", "marine", "land_access", "permit_service"]
PIPED = ("size", "connection", "design_pressure", "design_temperature", "operating_pressure_min", "operating_pressure_max",
         "operating_temperature_min", "operating_temperature_max", "flow_max", "flow_unit", "isolation", "quality_refs")
REQUIRED = {
    "process": PIPED, "utility": PIPED,
    "electrical": ("voltage_nominal", "voltage_min", "voltage_max", "connection", "isolation", "protection"),
    "control": ("connection",), "civil": (), "marine": (), "land_access": (), "permit_service": ()}
DATED = {"process", "utility", "electrical", "control", "civil", "marine", "land_access"}
MW_UNITS = ("MW", "MVA")


def to_kgs(value, unit, density):
    """Flow in kg/s, or None if the unit is not a mass/volume flow or needs a density that is not given."""
    if value is None or unit is None:
        return None
    if unit == "kg/s":
        return value
    if unit == "t/h":
        return value / 3.6
    if density is None:
        return None
    return {"m3/h": value * density / 3600, "Sm3/h": value * density / 3600, "Nm3/h": value * density / 3600,
            "m3/day": value * density / 86400, "m3/s": value * density}.get(unit)


def envelope(cases, streams):
    """{stream number: {m/p/T: (min, case, max, case)}} over the given cases (streams with zero flow are ignored)."""
    ids = {c["id"] for c in cases}
    env = {}
    for s in streams:
        if s["hmb_case"] not in ids or not s.get("mass_flow"):
            continue
        e = env.setdefault(s["number"], {"cases": set()})
        e["cases"].add(s["hmb_case"])
        for k, v in (("m", s["mass_flow"]), ("p", s["pressure"] - P_ATM), ("T", s["temperature"])):
            lo = e.get(k)
            e[k] = (min((lo[0], lo[1]), (v, s["hmb_case"])) if lo else (v, s["hmb_case"])) + \
                   (max((lo[2], lo[3]), (v, s["hmb_case"])) if lo else (v, s["hmb_case"]))
    return env


def _fmt(v, d=2):
    return "-" if v is None else (f"{v:,.{d}f}" if isinstance(v, (int, float)) else str(v))


def _pfmt(v):
    """Pressure: 3 decimals near atmospheric (vacuum, draft), 1 decimal otherwise."""
    return "-" if v is None else f"{v:,.{3 if abs(v) < 1 else 1}f}"


def check_register(points, cases, streams, powers, scope_items, owner_party=None):
    """Return [(point id or '-', check, 'OK'|'WARN'|'INFO', text)]. powers = {case id: {metric: MW}}."""
    out = []
    by_id = {p["id"]: p for p in points}
    env = envelope(cases, streams)

    def add(pid, chk, st, txt):
        out.append((pid, chk, st, txt))

    for p in sorted(points, key=lambda r: r["id"]):
        pid, cat = p["id"], p.get("category")
        # completeness
        miss = [f for f in ("boundary",) + REQUIRED.get(cat, ()) if p.get(f) in (None, "", [])]
        if cat == "electrical" and (p.get("voltage_nominal") or 0) >= 1:
            miss += [f for f in ("short_circuit", "short_circuit_rating") if p.get(f) is None]
        if cat in DATED and not p.get("needed_by"):
            miss.append("needed_by")
        if cat in DATED and not p.get("available_by"):
            miss.append("available_by")
        add(pid, "completeness", "WARN" if miss else "OK", ("missing: " + ", ".join(miss)) if miss else f"{cat}: complete")

        # own consistency
        lo, hi, dp = p.get("operating_pressure_min"), p.get("operating_pressure_max"), p.get("design_pressure")
        if lo is not None and hi is not None and lo > hi:
            add(pid, "ranges", "WARN", f"operating pressure min {lo} > max {hi} barg")
        if hi is not None and dp is not None and hi > dp:
            add(pid, "ranges", "WARN", f"operating pressure max {hi} barg above design {dp} barg")
        tlo, thi, dt = p.get("operating_temperature_min"), p.get("operating_temperature_max"), p.get("design_temperature")
        if thi is not None and dt is not None and thi > dt:
            add(pid, "ranges", "WARN", f"operating temperature max {thi} degC above design {dt} degC")
        if p.get("flow_normal") is not None and p.get("flow_max") is not None and p["flow_normal"] > p["flow_max"]:
            add(pid, "ranges", "WARN", f"normal flow {p['flow_normal']} above capacity {p['flow_max']} {p.get('flow_unit')}")
        vn, vl, vh = p.get("voltage_nominal"), p.get("voltage_min"), p.get("voltage_max")
        if None not in (vn, vl, vh) and not vl <= vn <= vh:
            add(pid, "ranges", "WARN", f"voltage range {vl}-{vh} kV does not contain the nominal {vn} kV")
        if p.get("short_circuit") is not None and p.get("short_circuit_rating") is not None:
            ok = p["short_circuit_rating"] >= p["short_circuit"]
            add(pid, "short circuit", "OK" if ok else "WARN",
                f"prospective {p['short_circuit']} kA {'<=' if ok else '>'} rating {p['short_circuit_rating']} kA")

        # HMB envelope
        n = p.get("hmb_stream")
        if n is not None:
            e = env.get(n)
            if not e:
                add(pid, "HMB", "WARN", f"stream {n} not found in any HMB case")
            else:
                mode = p.get("hmb_check") or "full"
                ncase = len(e["cases"])
                cap = to_kgs(p.get("flow_max"), p.get("flow_unit"), p.get("density"))
                if cap is None:
                    add(pid, "HMB flow", "WARN", f"stream {n}: flow capacity not convertible to kg/s "
                                                 f"({p.get('flow_max')} {p.get('flow_unit')}, density {p.get('density')})")
                else:
                    m = e["m"]
                    ok = m[2] <= cap * 1.000001
                    add(pid, "HMB flow", "OK" if ok else "WARN",
                        f"stream {n} max {m[2]:.2f} kg/s ({m[3]}) {'<=' if ok else '>'} capacity {cap:.2f} kg/s"
                        f" (margin {100 * (cap - m[2]) / cap:+.1f} %), {ncase} cases")
                if mode == "full":
                    pr = e["p"]
                    bad = []
                    if lo is not None and pr[0] < lo - 1e-6:
                        bad.append(f"min {pr[0]:.3f} barg ({pr[1]}) below range {lo}")
                    if hi is not None and pr[2] > hi + 1e-6:
                        bad.append(f"max {pr[2]:.3f} barg ({pr[3]}) above range {hi}")
                    if dp is not None and pr[2] > dp:
                        bad.append(f"max {pr[2]:.3f} barg above design {dp}")
                    add(pid, "HMB pressure", "WARN" if bad else "OK", f"stream {n}: " + ("; ".join(bad) if bad else
                        f"{pr[0]:.3f} - {pr[2]:.3f} barg within {lo} - {hi} barg (design {dp})"))
                if mode in ("full", "flow_temperature"):
                    tr = e["T"]
                    bad = []
                    if tlo is not None and tr[0] < tlo - 1e-6:
                        bad.append(f"min {tr[0]:.1f} degC ({tr[1]}) below range {tlo}")
                    if thi is not None and tr[2] > thi + 1e-6:
                        bad.append(f"max {tr[2]:.1f} degC ({tr[3]}) above range {thi}")
                    if dt is not None and tr[2] > dt:
                        bad.append(f"max {tr[2]:.1f} degC above design {dt}")
                    add(pid, "HMB temperature", "WARN" if bad else "OK", f"stream {n}: " + ("; ".join(bad) if bad else
                        f"{tr[0]:.1f} - {tr[2]:.1f} degC within {tlo} - {thi} degC (design {dt})"))
                if mode != "full":
                    add(pid, "HMB scope", "INFO", f"stream {n}: {mode.replace('_', ' + ')} checked only - see remarks")

        metric = p.get("hmb_metric")
        if metric:
            vals = [(pw[metric], cid) for cid, pw in powers.items() if pw.get(metric) is not None]
            unit, cap = p.get("flow_unit"), p.get("flow_max")
            if not vals or unit not in MW_UNITS or cap is None:
                add(pid, "HMB power", "WARN", f"{metric}: no HMB values or no rating in MW/MVA")
            else:
                P, cid = max(vals)
                pf = p.get("power_factor") or 1.0
                load = P if unit == "MW" else P / pf
                ok = load <= cap
                txt = (f"max {metric} {P:.1f} MW ({cid})" + ("" if unit == "MW" else f" / pf {pf} = {load:.1f} MVA")
                       + f" {'<=' if ok else '>'} {cap:,.0f} {unit} (margin {100 * (cap - load) / cap:+.1f} %)")
                if not ok and p.get("flow_limited"):
                    add(pid, "HMB power", "INFO", txt + f"; limited to {cap:,.0f} {unit} at this point: {p.get('protection') or '?'}")
                else:
                    add(pid, "HMB power", "OK" if ok else "WARN", txt)

        # chain to the upstream point
        up = by_id.get(p.get("upstream")) if p.get("upstream") else None
        if p.get("upstream") and up is None:
            add(pid, "chain", "WARN", f"upstream {p['upstream']} not found")
        if up:
            uid = up["id"]
            if up.get("operating_pressure_min") is not None and lo is not None:
                cdp = p.get("chain_dp")
                if cdp is None:
                    add(pid, "chain pressure", "WARN", f"from {uid}: chain_dp missing")
                else:
                    avail = up["operating_pressure_min"] - cdp
                    ok = avail >= lo - 1e-6
                    add(pid, "chain pressure", "OK" if ok else "WARN",
                        f"{uid} min {up['operating_pressure_min']} barg - {cdp} bar = {avail:.2f} barg "
                        f"{'>=' if ok else '<'} required {lo} barg")
            if up.get("design_pressure") is not None and dp is not None and up["design_pressure"] > dp:
                add(pid, "chain design", "OK" if p.get("protection") else "WARN",
                    f"{uid} design {up['design_pressure']} barg > {dp} barg here: "
                    + (f"protected - {p['protection'][:120]}" if p.get("protection") else "no protection declared"))
            if up.get("design_temperature") is not None and dt is not None and up["design_temperature"] > dt:
                add(pid, "chain design", "OK" if p.get("protection") else "WARN",
                    f"{uid} design {up['design_temperature']} degC > {dt} degC here: "
                    + ("protected" if p.get("protection") else "no protection declared"))
            fu = to_kgs(up.get("flow_max"), up.get("flow_unit"), up.get("density"))
            fh = to_kgs(p.get("flow_max"), p.get("flow_unit"), p.get("density"))
            if fu is not None and fh is not None:
                ok = fh <= fu * 1.000001
                add(pid, "chain flow", "OK" if ok else "WARN",
                    f"capacity here {fh:.2f} kg/s {'<=' if ok else '>'} {uid} {fu:.2f} kg/s")
            if None not in (up.get("voltage_min"), up.get("voltage_max"), vl, vh):
                ok = vl <= up["voltage_min"] and up["voltage_max"] <= vh
                add(pid, "chain voltage", "OK" if ok else "WARN",
                    f"{uid} {up['voltage_min']}-{up['voltage_max']} kV {'within' if ok else 'outside'} {vl}-{vh} kV here")
            if up.get("short_circuit") is not None and p.get("short_circuit_rating") is not None:
                ok = p["short_circuit_rating"] >= up["short_circuit"]
                add(pid, "chain short circuit", "OK" if ok else "WARN",
                    f"{uid} {up['short_circuit']} kA {'<=' if ok else '>'} rating {p['short_circuit_rating']} kA here")

        # dates
        a, nb = p.get("available_by"), p.get("needed_by")
        if a and nb:
            fl = (date.fromisoformat(nb) - date.fromisoformat(a)).days
            add(pid, "dates", "WARN" if fl < 0 else ("INFO" if fl < 14 else "OK"),
                f"available {a}, needed {nb}: float {fl} d" + (" (late)" if fl < 0 else (" (tight)" if fl < 14 else "")))

        # agreement
        st = p.get("status")
        if st in ("agreed", "frozen"):
            ok = len(p.get("agreed_by") or []) >= 2
            add(pid, "agreement", "OK" if ok else "WARN",
                f"{st} by {', '.join(p.get('agreed_by') or []) or '?'}" + (f" ({p['agreed_ref']})" if p.get("agreed_ref") else ""))
        else:
            add(pid, "agreement", "WARN", f"status {st or '?'}: not yet agreed")

    # HMB streams that no tie-in maps: either internal to one party's scope, or a missing interface
    mapped = {p.get("hmb_stream") for p in points}
    desc = {}
    for s in streams:
        if s["number"] in env and s["number"] not in mapped:
            desc.setdefault(s["number"], s["description"])
    if desc:
        add("-", "HMB coverage", "INFO", "streams not mapped to a tie-in (internal to one scope, else an interface is "
                                         "missing): " + "; ".join(f"{n} {d}" for n, d in sorted(desc.items())))

    # coverage of Owner-supplied scope items with a terminal point
    served = {}
    for p in points:
        for sc in [p.get("scope_item")] + list(p.get("related_scope") or []):
            if sc:
                served.setdefault(sc, []).append(p["id"])
    for sc in sorted(scope_items, key=lambda r: r["id"]):
        if not sc.get("terminal_point"):
            continue
        if sc["id"] in served:
            add("-", "coverage", "OK", f"{sc['id']} ({sc['terminal_point'][:50]}): {', '.join(sorted(served[sc['id']]))}")
        elif owner_party and sc.get("supply") == owner_party:
            add("-", "coverage", "WARN", f"{sc['id']} Owner-supplied with terminal point '{sc['terminal_point']}' has no tie-in")
        else:
            add("-", "coverage", "INFO", f"{sc['id']} terminal point '{sc['terminal_point'][:60]}' not referenced by a tie-in")
    return out


class TieInRegister(Engine):
    name = "tie_in_register"
    title = "Tie-in / terminal point register (Excel + Word + PDF) with HMB-envelope, chain, date and agreement checks"
    version = "1.0.0"
    inputs = ["project", "tie_in", "contract", "party", "hmb_case", "process_stream", "aux_load", "scope_item", "clarification"]
    formats = ["xlsx", "docx", "pdf"]
    code_deps = ["engine/core/hmb.py", "engine/core/thermo.py", "templates/docx/datasheet_base.docx"]

    def run(self, ctx: Context):
        s = ctx.store
        points = s.records("tie_in")
        if not points:
            ctx.warnings.append("no tie_in records - nothing generated")
            return []
        cases = [c for c in s.records("hmb_case") if c.get("status") != "superseded"]
        streams = s.records("process_stream")
        powers = {}
        for c in cases:
            st = [x for x in streams if x["hmb_case"] == c["id"]]
            if not st:
                continue
            r = hmb.calculate(c, st, [x for x in s.records("aux_load") if x["hmb_case"] == c["id"]], {})
            powers[c["id"]] = {"gt_output": c["gt_output"] / 1000, "st_output": c["st_output"] / 1000,
                               "gross": r.summary["gross"] / 1000, "net": r.summary["net"] / 1000}
        owner = next((p["id"] for p in s.records("party") if p.get("role") == "owner"), None)
        owner_items = [x for x in s.records("scope_item") if x["contract"] in
                       {k["id"] for k in s.records("contract") if k.get("kind", "owner_contract") == "owner_contract"}]
        checks = check_register(points, cases, streams, powers, owner_items, owner)
        for pid, chk, st, txt in checks:
            if st == "WARN":
                ctx.warnings.append(f"{pid} {chk}: {txt}")
        env = envelope(cases, streams)
        stem = f"{ctx.project_record()['id']}_tie_in_register"
        files = [self._workbook(ctx, points, checks, env, cases, stem)]
        doc = self._document(ctx, points, checks, env, cases, stem)
        files.append(doc)
        if ctx.options.get("pdf", "yes") != "no":
            pdf = pdf_from_office(doc, ctx)
            if pdf:
                files.append(pdf)
        return files

    # ------------------------------------------------------------------ Excel
    COLS = [("ID", "id", 8), ("Contract", "contract", 12), ("Category", "category", 11), ("Status", "status", 8),
            ("Agreed by", "agreed_by", 12), ("Agreed in", "agreed_ref", 11), ("Service", "service", 24), ("Medium", "medium", 14),
            ("Location", "location_text", 28), ("Size", "size", 16), ("Connection", "connection", 30),
            ("Scope break (boundary)", "boundary", 50), ("Employer side (Owner / EPC)", "owner_side", 30), ("Contractor side (EPC / IEC)", "contractor_side", 30),
            ("Isolation", "isolation", 36), ("Protection", "protection", 36), ("Metering", "metering", 24),
            ("Design p barg", "design_pressure", 9), ("Design T degC", "design_temperature", 9),
            ("Op p min barg", "operating_pressure_min", 9), ("Op p max barg", "operating_pressure_max", 9),
            ("Op T min degC", "operating_temperature_min", 9), ("Op T max degC", "operating_temperature_max", 9),
            ("Flow normal", "flow_normal", 10), ("Flow max", "flow_max", 10), ("Flow unit", "flow_unit", 8),
            ("Density kg/m3", "density", 8), ("Limited", "flow_limited", 7),
            ("U nom kV", "voltage_nominal", 7), ("U min kV", "voltage_min", 7), ("U max kV", "voltage_max", 7),
            ("Isc kA", "short_circuit", 7), ("Isc rating kA", "short_circuit_rating", 8), ("pf", "power_factor", 5),
            ("HMB stream", "hmb_stream", 7), ("HMB check", "hmb_check", 10), ("HMB power", "hmb_metric", 10),
            ("Upstream", "upstream", 8), ("Chain dp bar", "chain_dp", 7), ("Quality", "quality_refs", 26),
            ("Available by", "available_by", 11), ("Needed by", "needed_by", 11), ("Scope", "scope_item", 8),
            ("Related scope", "related_scope", 14), ("Operating conditions (text)", "operating_conditions", 40),
            ("Basis", "basis_refs", 24), ("Remarks", "remarks", 40)]

    def _workbook(self, ctx, points, checks, env, cases, stem):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        head = PatternFill("solid", fgColor="1F3864")
        fills = {"WARN": PatternFill("solid", fgColor="F8CBAD"), "INFO": PatternFill("solid", fgColor="FFF2CC"),
                 "OK": PatternFill("solid", fgColor="E2EFDA")}
        wb = Workbook()

        def sheet(ws, title, cols):
            ws["A1"] = f"{ctx.project_record().get('name', '')} - {title}"
            ws["A1"].font = Font(bold=True, size=14)
            ws["A2"] = (f"engine {self.name} v{self.version} | generated {ctx.stamp['generated_at'][:19]}Z by "
                        f"{ctx.stamp['generated_by']} | data {ctx.stamp['inputs_hash']} | pressures barg, temperatures degC")
            ws["A2"].font = Font(italic=True, size=8, color="595959")
            for c, (n, w) in enumerate(cols, 1):
                cell = ws.cell(4, c, n)
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                ws.column_dimensions[get_column_letter(c)].width = w
            ws.freeze_panes = "B5"
            ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}4"
            ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True

        ws = wb.active
        ws.title = "Register"
        sheet(ws, "Tie-in / Terminal Point Register", [(n, w) for n, _, w in self.COLS])
        for i, p in enumerate(self._sorted(points), 5):
            for c, (_, k, _) in enumerate(self.COLS, 1):
                v = p.get(k)
                v = "\n".join(v) if isinstance(v, list) else v
                ws.cell(i, c, v).alignment = Alignment(wrap_text=isinstance(v, str) and len(v) > 12, vertical="top")

        wc = wb.create_sheet("Checks")
        sheet(wc, "Tie-in checks", [("Point", 8), ("Check", 18), ("Result", 8), ("Detail", 140)])
        for i, (pid, chk, st, txt) in enumerate(checks, 5):
            for c, v in enumerate((pid, chk, st, txt), 1):
                wc.cell(i, c, v).alignment = Alignment(vertical="top", wrap_text=c == 4)
            wc.cell(i, 3).fill = fills[st]

        we = wb.create_sheet("HMB envelope")
        sheet(we, f"HMB envelope at the tie-ins over {len(cases)} cases: {', '.join(sorted(c['id'] for c in cases))}",
              [("Point", 8), ("Stream", 7), ("Check", 12), ("Flow min kg/s", 10), ("Flow max kg/s", 10), ("at case", 13),
               ("Capacity kg/s", 10), ("p min barg", 10), ("p max barg", 10), ("Range barg", 14), ("Design barg", 9),
               ("T min degC", 9), ("T max degC", 9), ("Range degC", 12), ("Design degC", 9)])
        r = 5
        for p in self._sorted(points):
            e = env.get(p.get("hmb_stream"))
            if not e:
                continue
            cap = to_kgs(p.get("flow_max"), p.get("flow_unit"), p.get("density"))
            row = [p["id"], p["hmb_stream"], p.get("hmb_check") or "full", round(e["m"][0], 3), round(e["m"][2], 3), e["m"][3],
                   None if cap is None else round(cap, 2), round(e["p"][0], 4), round(e["p"][2], 4),
                   f"{_fmt(p.get('operating_pressure_min'))} - {_fmt(p.get('operating_pressure_max'))}", p.get("design_pressure"),
                   round(e["T"][0], 1), round(e["T"][2], 1),
                   f"{_fmt(p.get('operating_temperature_min'), 0)} - {_fmt(p.get('operating_temperature_max'), 0)}",
                   p.get("design_temperature")]
            for c, v in enumerate(row, 1):
                we.cell(r, c, v)
            r += 1

        sm = wb.create_sheet("Summary")
        sm.append(["Category", "Points", "Agreed / frozen", "Draft", "Checks WARN", "Checks INFO"])
        for cat in CATEGORY_ORDER:
            ps = [p for p in points if p.get("category") == cat]
            if not ps:
                continue
            ids = {p["id"] for p in ps}
            sm.append([cat, len(ps), sum(1 for p in ps if p.get("status") in ("agreed", "frozen")),
                       sum(1 for p in ps if p.get("status") not in ("agreed", "frozen")),
                       sum(1 for x in checks if x[0] in ids and x[2] == "WARN"),
                       sum(1 for x in checks if x[0] in ids and x[2] == "INFO")])
        sm.append(["Total", len(points), sum(1 for p in points if p.get("status") in ("agreed", "frozen")),
                   sum(1 for p in points if p.get("status") not in ("agreed", "frozen")),
                   sum(1 for x in checks if x[2] == "WARN"), sum(1 for x in checks if x[2] == "INFO")])
        for row in sm.iter_rows(max_row=1):
            for c in row:
                c.font = Font(bold=True)
        sm.column_dimensions["A"].width = 16
        out = ctx.out_dir / f"{stem}.xlsx"
        wb.save(out)
        return out

    @staticmethod
    def _sorted(points):
        return sorted(points, key=lambda p: (not p["id"].startswith("TP-"), p["id"]))

    # ------------------------------------------------------------------ Word / PDF
    def _document(self, ctx, points, checks, env, cases, stem):
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Cm, Pt, RGBColor

        s = ctx.store
        base = ctx.template("docx", "datasheet_base.docx")
        doc = Document(str(base)) if base.exists() else Document()
        for p in list(doc.paragraphs):
            p._element.getparent().remove(p._element)
        for t in list(doc.tables):
            t._element.getparent().remove(t._element)
        self.doc = doc
        doc.styles["Normal"].font.size = Pt(9)
        sec = doc.sections[0]
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(1.8)
        sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)
        proj = ctx.project_record()
        contracts = {k["id"]: k for k in s.records("contract")}
        parties = {p["id"]: p for p in s.records("party")}
        tqs = {q["id"]: q for q in s.records("clarification")}

        hdr = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
        hdr.text = f"{proj.get('name', '')} | {stem} | tie-in / terminal point register"
        hdr.runs[0].font.size = Pt(7)
        ftr = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        ftr.text = (f"engine {self.name} v{self.version} | data {ctx.stamp['inputs_hash']} | generated "
                    f"{ctx.stamp['generated_at'][:10]} - generated from the project database; do not edit")
        ftr.runs[0].font.size = Pt(7)

        doc.add_heading(f"{proj.get('name', '')}\nTie-in and Terminal Point Register", level=0)
        warn = sum(1 for c in checks if c[2] == "WARN")
        agreed = sum(1 for p in points if p.get("status") in ("agreed", "frozen"))
        doc.add_paragraph(
            f"{len(points)} points: {sum(1 for p in points if p['id'].startswith('TP-'))} Owner/Contractor terminal points (TP-) "
            f"under {', '.join(sorted({p['contract'] for p in points if p['id'].startswith('TP-')}))} and "
            f"{sum(1 for p in points if p['id'].startswith('IF-'))} interfaces between the consortium members (IF-) under "
            f"{', '.join(sorted({p['contract'] for p in points if p['id'].startswith('IF-')}))}. "
            f"{agreed} agreed or frozen, {len(points) - agreed} draft. Checks: {warn} warning(s), "
            f"{sum(1 for c in checks if c[2] == 'INFO')} note(s).")
        doc.add_heading("1. Conventions", level=1)
        for t in [
            "TP- points: Owner side / Contractor side under the EPC contract. IF- points: Istanbul EPC side / Imaginary Electric "
            "side under the consortium agreement. The scope break states who supplies the flange or weld end, counter-flange, "
            "gasket, bolts, the closing weld and its NDT/PWHT, supports and cable terminations; 'available by' is the date the "
            "providing side has its side ready, 'needed by' the date the receiving side needs it.",
            "Pressures in barg (vacuum as negative barg; HMB stream pressures in bar(a) are converted with 1.01325 bar), "
            "temperatures in degC. Design values are the mechanical design of the receiving connection.",
            "Operating envelope = the range the providing side guarantees at the point. It is checked against the HMB streams of "
            f"all {len(cases)} HMB cases ({', '.join(sorted(c['id'] for c in cases))}).",
            "Chain = point fed from an upstream point through passive equipment: the minimum upstream pressure less the agreed "
            "pressure loss (chain dp) must cover this point's minimum pressure; a higher upstream design pressure or temperature "
            "requires a declared protection (relief, slam-shut, safety valve set points).",
            "Status: draft (proposed), agreed (confirmed in writing by both sides, TQ reference), frozen (change only by TQ)."]:
            doc.add_paragraph(t, style="List Bullet")

        doc.add_heading("2. Overview", level=1)
        rows = [[p["id"], p.get("category", "").replace("_", " "), p["service"], p.get("size") or "-", self._env_short(p), p.get("available_by") or "-",
                 p.get("needed_by") or "-", p.get("status", "")] for p in self._sorted(points)]
        self._table(["ID", "Category", "Service", "Size / rating", "Operating envelope", "Available", "Needed", "Status"],
                    rows, [1.5, 1.8, 3.3, 2.6, 3.8, 1.7, 1.7, 1.3])

        doc.add_heading("3. Tie-in data cards", level=1)
        for p in self._sorted(points):
            self._card(p, contracts, parties, tqs, env, checks)

        doc.add_heading("4. Check results", level=1)
        doc.add_paragraph("WARN = to be resolved before the point is frozen; INFO = noted; OK rows are listed in the Excel register.")
        rows = [[a, b, c, d] for a, b, c, d in checks if c != "OK"]
        if rows:
            self._table(["Point", "Check", "Result", "Detail"], rows, [1.4, 2.6, 1.3, 12.1])
        else:
            doc.add_paragraph("No warnings or notes.")
        out = ctx.out_dir / f"{stem}.docx"
        doc.save(out)
        return out

    @staticmethod
    def _env_short(p):
        parts = []
        if p.get("operating_pressure_min") is not None:
            parts.append(f"{_fmt(p['operating_pressure_min'], 1)}-{_fmt(p.get('operating_pressure_max'), 1)} barg")
        if p.get("operating_temperature_min") is not None:
            parts.append(f"{_fmt(p['operating_temperature_min'], 0)}-{_fmt(p.get('operating_temperature_max'), 0)} degC")
        if p.get("voltage_nominal") is not None:
            parts.append(f"{_fmt(p['voltage_nominal'], 1)} kV ({_fmt(p.get('voltage_min'), 1)}-{_fmt(p.get('voltage_max'), 1)})")
        if p.get("flow_max") is not None:
            parts.append(f"max {_fmt(p['flow_max'], 2 if p['flow_max'] < 100 else 0)} {p.get('flow_unit') or ''}")
        return "; ".join(parts) or "-"

    def _card(self, p, contracts, parties, tqs, env, checks):
        from docx.shared import Pt
        k = contracts.get(p["contract"], {})
        pid = p["id"]
        h = self.doc.add_heading(f"{pid}  {p['service']}", level=2)
        h.paragraph_format.keep_with_next = True
        emp = parties.get(k.get("owner"), {}).get("short_name") or k.get("owner") or "Employer"
        con = ", ".join(parties.get(x, {}).get("short_name") or x for x in k.get("contractor_parties") or []) or "Contractor"
        if k.get("kind", "owner_contract") == "owner_contract":
            emp, con = "Owner", "Contractor"
        rows = [("Contract / category", f"{p['contract']} {k.get('title', '')} / {p.get('category', '-')}"),
                ("Medium / location", f"{p['medium']} / {p['location_text']}"
                 + (f" (E {p['location'][0] / 1000:.1f}, N {p['location'][1] / 1000:.1f}, EL {p['location'][2] / 1000:+.2f} m)"
                    if p.get("location") else "")),
                ("Size / connection", f"{p.get('size') or '-'} / {p.get('connection') or '-'}"),
                ("Scope break", p.get("boundary") or "-"),
                (f"{emp} side", p.get("owner_side") or "-"),
                (f"{con} side", p.get("contractor_side") or "-"),
                ("Isolation", p.get("isolation") or "-"), ("Protection", p.get("protection") or "-")]
        if p.get("metering"):
            rows.append(("Metering", p["metering"]))
        if p.get("design_pressure") is not None or p.get("design_temperature") is not None:
            rows.append(("Design", f"{_pfmt(p.get('design_pressure'))} barg / {_fmt(p.get('design_temperature'), 0)} degC"))
        if p.get("operating_pressure_min") is not None or p.get("operating_temperature_min") is not None:
            rows.append(("Operating envelope", f"{_pfmt(p.get('operating_pressure_min'))} to {_pfmt(p.get('operating_pressure_max'))} barg; "
                                              f"{_fmt(p.get('operating_temperature_min'), 0)} to {_fmt(p.get('operating_temperature_max'), 0)} degC"))
        if p.get("flow_max") is not None:
            rows.append(("Flow", f"normal {_fmt(p.get('flow_normal'), 2)} / max {_fmt(p['flow_max'], 2)} {p.get('flow_unit') or ''}"
                                 + (f" (density {p['density']} kg/m3)" if p.get("density") else "")
                                 + (" - actively limited" if p.get("flow_limited") else "")))
        if p.get("voltage_nominal") is not None:
            rows.append(("Electrical", f"{_fmt(p['voltage_nominal'], 2)} kV ({_fmt(p.get('voltage_min'), 1)}-{_fmt(p.get('voltage_max'), 1)} kV); "
                                       f"Isc {_fmt(p.get('short_circuit'), 1)} kA, rating {_fmt(p.get('short_circuit_rating'), 1)} kA"
                                       + (f"; pf {p['power_factor']}" if p.get("power_factor") else "")))
        e = env.get(p.get("hmb_stream"))
        if e:
            rows.append(("HMB envelope (all cases)", f"stream {p['hmb_stream']} ({p.get('hmb_check') or 'full'} check): "
                                                     f"{e['m'][0]:.2f}-{e['m'][2]:.2f} kg/s, {e['p'][0]:.3f}-{e['p'][2]:.3f} barg, "
                                                     f"{e['T'][0]:.1f}-{e['T'][2]:.1f} degC"))
        if p.get("hmb_metric"):
            rows.append(("HMB power", p["hmb_metric"]))
        if p.get("upstream"):
            rows.append(("Fed from", f"{p['upstream']} (chain dp {_fmt(p.get('chain_dp'), 1)} bar)"))
        if p.get("quality_refs"):
            rows.append(("Quality / specification", ", ".join(p["quality_refs"])))
        if p.get("operating_conditions"):
            rows.append(("Conditions (text)", p["operating_conditions"]))
        rows.append(("Dates", f"available by {p.get('available_by') or '-'}; needed by {p.get('needed_by') or '-'}"))
        scope = ", ".join(x for x in [p.get("scope_item")] + list(p.get("related_scope") or []) if x)
        if scope:
            rows.append(("Scope items", scope))
        ag = p.get("status", "-")
        if p.get("agreed_by"):
            ag += " by " + ", ".join(parties.get(x, {}).get("short_name") or x for x in p["agreed_by"])
        if p.get("agreed_ref"):
            q = tqs.get(p["agreed_ref"], {})
            ag += f" in {p['agreed_ref']}" + (f" ({q.get('response_date')})" if q.get("response_date") else "")
        rows.append(("Agreement", ag))
        rows.append(("Basis", ", ".join(p.get("basis_refs") or []) or "-"))
        if p.get("remarks"):
            rows.append(("Remarks", p["remarks"]))
        ver = [f"{c[2]} {c[1]}: {c[3]}" for c in checks if c[0] == pid and c[1] not in ("agreement",)]
        if ver:
            rows.append(("Checks (engine)", "\n".join(ver)))
        t = self._table(None, [list(r) for r in rows], [4.2, 13.2], size=8)
        for r in t.rows:
            r.cells[0].paragraphs[0].runs[0].bold = True
            for par in r.cells[1].paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8)

    def _table(self, head, rows, widths_cm, size=7.5):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
        t = self.doc.add_table(rows=1 if head else 0, cols=len(widths_cm))
        t.style = "Table Grid"
        if head:
            for c, h in zip(t.rows[0].cells, head):
                c.text = h
                run = c.paragraphs[0].runs[0]
                run.bold, run.font.size, run.font.color.rgb = True, Pt(size), RGBColor(0xFF, 0xFF, 0xFF)
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), "1F3864")
                c._element.get_or_add_tcPr().append(shd)
            rh = OxmlElement("w:tblHeader")
            rh.set(qn("w:val"), "true")
            t.rows[0]._tr.get_or_add_trPr().append(rh)
        for r in rows:
            cells = t.add_row().cells
            for c, v in zip(cells, r):
                c.text = "" if v is None else str(v)
                for par in c.paragraphs:
                    par.paragraph_format.space_after = Pt(0)
                    for run in par.runs:
                        run.font.size = Pt(size)
            cs = OxmlElement("w:cantSplit")
            cs.set(qn("w:val"), "true")
            t.rows[-1]._tr.get_or_add_trPr().append(cs)
        t.autofit = False
        tbl_grid = t._tbl.tblGrid
        for gc, w in zip(tbl_grid.findall(qn("w:gridCol")), widths_cm):
            gc.set(qn("w:w"), str(int(w * 567)))
        for row in t.rows:
            for cell, w in zip(row.cells, widths_cm):
                cell.width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t


ENGINE = TieInRegister()
