"""Heat and mass balance (HMB) diagram and calculation per hmb_case: DXF + PDF diagram (A1) with stream data, plant
power balance and checks; Excel workbook with streams, node balances, auxiliary loads and checks.

The power-island data (GT/ST outputs, stream conditions) come from the vendor records; the engine recomputes every
stream enthalpy (IAPWS-IF97 / ideal-gas mixtures), closes mass and energy balances per node, adds the balance-of-plant
auxiliary loads and checks the result against the guarantees and limits listed in the case's check_refs.
Layout: templates/hmb/<layout>.json. Options: case=<id> (default: all cases), layout=layout_1x1_3prh, pdf=no."""
from __future__ import annotations

import json
import math

from ..core import dxfkit, hmb
from ..core.runner import Context, Engine

FLUID_LAYER = {"air": ("HMB-AIR", 8), "fuel_gas": ("HMB-FUEL", 30), "flue_gas": ("HMB-GAS", 34), "steam": ("HMB-STEAM", 1),
               "water": ("HMB-WATER", 5), "seawater": ("HMB-CW", 4), "ldo": ("HMB-FUEL", 30), "other": ("HMB-WATER", 5)}


def refs_for(s, case):
    kinds = {c["id"]: c.get("kind", "owner_contract") for c in s.records("contract")}
    out = {}
    for r in case.get("check_refs", []):
        ent, rid = r.split(":", 1)
        rec = s.get(ent, rid)
        if rec is None:
            continue
        rec = dict(rec)
        if ent == "guarantee":
            rec["_kind"] = kinds.get(rec["contract"], "owner_contract")
        out[r] = rec
    return out


class HeatMassBalance(Engine):
    name = "hmb"
    title = "Heat and mass balance: diagram (DXF + PDF) and calculation workbook (Excel) per HMB case"
    version = "1.0.0"
    inputs = ["project", "hmb_case", "process_stream", "aux_load", "guarantee", "design_parameter", "contract"]
    formats = ["dxf", "pdf", "xlsx"]
    code_deps = ["engine/core/hmb.py", "engine/core/thermo.py", "engine/core/dxfkit.py", "templates/hmb"]

    def run(self, ctx: Context):
        s = ctx.store
        cases = [c for c in s.records("hmb_case") if c.get("status") != "superseded"]
        if ctx.options.get("case"):
            cases = [c for c in cases if c["id"] == ctx.options["case"]]
        if not cases:
            ctx.warnings.append("no hmb_case records - nothing generated")
            return []
        layout = json.loads(ctx.template("hmb", ctx.options.get("layout", "layout_1x1_3prh") + ".json").read_text())
        files = []
        for case in sorted(cases, key=lambda c: c["id"]):
            streams = [x for x in s.records("process_stream") if x["hmb_case"] == case["id"]]
            aux = [x for x in s.records("aux_load") if x["hmb_case"] == case["id"]]
            if not streams:
                ctx.warnings.append(f"{case['id']}: no process streams")
                continue
            res = hmb.calculate(case, streams, aux, refs_for(s, case))
            for st, text in res.checks:
                if st == "WARN":
                    ctx.warnings.append(f"{case['id']}: {text}")
            stem = f"{ctx.project_record()['id']}-HMB-{case['id']}"
            files += self._drawing(ctx, case, res, layout, stem)
            files.append(self._workbook(ctx, case, res, stem))
        return files

    # ------------------------------------------------------------------ drawing
    def _drawing(self, ctx, case, res, L, stem):
        doc = dxfkit.new_doc()
        for name, (color, _) in {k: (c, 0) for k, c in FLUID_LAYER.values()}.items():
            if name not in doc.layers:
                doc.layers.add(name, color=color)
        for name, color in (("HMB-EQPT", 7), ("HMB-TAG", 7), ("HMB-DATA", 8), ("HMB-TABLE", 7), ("HMB-SCOPE", 3), ("HMB-WARN", 1)):
            doc.layers.add(name, color=color)
        m = doc.modelspace()
        W, H = L["paper"]
        T = lambda txt, x, y, h, layer="HMB-EQPT", align="LEFT", rot=0.0: dxfkit.text(m, txt, x, y, h, layer, align, rot)  # noqa: E731
        line = lambda pts, layer="HMB-EQPT", close=False, **kw: m.add_lwpolyline(pts, close=close, dxfattribs={"layer": layer, **kw})  # noqa: E731,E501
        rect = lambda x0, y0, x1, y1, layer="HMB-EQPT", **kw: line([(x0, y0), (x1, y0), (x1, y1), (x0, y1)], layer, close=True, **kw)  # noqa: E731,E501
        # frame
        rect(0, 0, W, H, "BORDER")
        rect(10, 10, W - 10, H - 10, "BORDER")
        nodes = L["nodes"]
        warn_nodes = {t.split(":")[0].split()[-1] for st, t in res.checks if st == "WARN" and t.startswith(("energy", "mass"))}
        for n, nd in nodes.items():
            self._symbol(m, T, line, rect, n, nd, n in warn_nodes)
        for g in L.get("generators", []):
            x, y = g["at"]
            m.add_circle((x, y), 14, dxfattribs={"layer": "HMB-EQPT"})
            T(g["label"], x, y, 5, "HMB-EQPT", "MIDDLE")
            if "shaft_to" in g:
                line([(x + 14, y), tuple(g["shaft_to"])], "HMB-EQPT", lineweight=70)
            if "shaft_from" in g:
                line([tuple(g["shaft_from"]), (x - 14, y)], "HMB-EQPT", lineweight=70)
            p = res.summary["gt_output"] if g["node"] == "GT" else res.summary["st_output"]
            T(f"{p / 1000:,.2f} MW", x, y - 22, 3.2, "HMB-TAG", "CENTER")
        # streams
        by_no = {s["number"]: s for s in res.streams}
        routes = L["streams"]
        for no, st in sorted(by_no.items()):
            r = routes.get(str(no))
            layer = FLUID_LAYER[st["fluid"]][0]
            if r is None:
                ctx.warnings.append(f"{case['id']}: stream {no} has no route in the layout - not drawn on the diagram")
                continue
            if r.get("internal"):
                continue
            pts = [tuple(p) for p in r["route"]]
            line(pts, layer, lineweight=50)
            self._arrow(m, pts[-2], pts[-1], layer)
            tx, ty = r["tag"]
            self._tag(m, T, tx, ty, no)
            bx, by = r["box"]
            q = st.get("quality")
            rows = [f"{st['mass_flow']:,.2f} kg/s", f"{st['pressure']:.4g} bar(a)" if st["pressure"] < 0.1 else f"{st['pressure']:.2f} bar(a)",
                    f"{st['temperature']:.1f} degC" if q is None else f"x = {q:.4f}", f"{st['h']:,.1f} kJ/kg"]
            rect(bx, by, bx + 30, by + 13.5, "HMB-DATA")
            T(str(no), bx + 1, by + 10.2, 2.4, "HMB-TAG")
            for i, txt in enumerate(rows):
                T(txt, bx + 29, by + 10.2 - i * 3.1, 2.2, "HMB-DATA", "RIGHT")
        # PIPE internal connections (dashed)
        pipe = nodes.get("PIPE", {})
        for a, b in pipe.get("pairs", []):
            if str(a) in routes and str(b) in routes and "route" in routes[str(a)] and "route" in routes[str(b)]:
                (xa, ya), (xb, yb) = routes[str(a)]["route"][-1], routes[str(b)]["route"][0]
                pts = [(xa, ya), (xb, yb)] if abs(xa - xb) < 0.01 else [(xa, ya), (xa, (ya + yb) / 2), (xb, (ya + yb) / 2), (xb, yb)]
                line(pts, "HMB-SCOPE", linetype="DASHED")
        if "legend" in L:
            lx, ly = L["legend"]["at"]
            T("LINES", lx, ly, 2.8, "HMB-TABLE")
            seen = []
            for fl, (lay, _) in FLUID_LAYER.items():
                if lay in seen or not any(s["fluid"] == fl for s in res.streams):
                    continue
                seen.append(lay)
                ly -= 5
                line([(lx, ly + 1), (lx + 12, ly + 1)], lay, lineweight=50)
                T(fl.replace("_", " "), lx + 15, ly, 2.4, "HMB-TABLE")
            ly -= 6
            self._tag(m, T, lx + 6, ly + 1, 0)
            T("stream number; data box: flow / pressure / temperature (or steam quality x) / enthalpy", lx + 15, ly, 2.2, "HMB-TABLE")
        self._tables(m, T, line, rect, case, res, W, H)
        dxfkit.title_block(m, W - 12, 12, 1.0, [
            ("PROJECT", ctx.project_record().get("name", "")),
            ("TITLE", f"HEAT AND MASS BALANCE - {case['id']}"),
            ("CASE", case["title"][:60]),
            ("DOCUMENT", stem),
            ("STATUS", f"{case.get('status', '')} - generated by engine {self.name} v{self.version}"),
            ("DATA", f"{ctx.stamp['inputs_hash']} | {ctx.stamp['generated_at'][:10]}"),
            ("NOTE", "all values calculated from the database; do not edit")], 190, 6)
        out = []
        dxf = ctx.out_dir / f"{stem}.dxf"
        assert not doc.audit().has_errors
        doc.saveas(dxf)
        out.append(dxf)
        if ctx.options.get("pdf", "yes") != "no":
            out.append(dxfkit.save_pdf(doc, ctx.out_dir / f"{stem}.pdf", tuple(L["paper"])))
        return out

    @staticmethod
    def _arrow(m, p0, p1, layer):
        a = math.atan2(p1[1] - p0[1], p1[0] - p0[0])
        L, w = 3.5, 1.3
        bx, by = p1[0] - L * math.cos(a), p1[1] - L * math.sin(a)
        pts = [p1, (bx + w * math.sin(a), by - w * math.cos(a)), (bx - w * math.sin(a), by + w * math.cos(a))]
        h = m.add_hatch(dxfattribs={"layer": layer})
        from ezdxf import const
        h.set_solid_fill(color=const.BYLAYER)
        h.paths.add_polyline_path(pts, is_closed=True)

    @staticmethod
    def _tag(m, T, x, y, no):
        r = 2.8
        m.add_lwpolyline([(x - r, y), (x, y + r), (x + r, y), (x, y - r)], close=True, dxfattribs={"layer": "HMB-TAG"})
        h = m.add_hatch(color=7, dxfattribs={"layer": "HMB-TAG"})
        h.paths.add_polyline_path([(x - r, y), (x, y + r), (x + r, y), (x, y - r)], is_closed=True)
        h.set_solid_fill(color=255)          # white fill hides the line under the tag
        T(str(no), x, y, 2.3, "HMB-TAG", "MIDDLE")

    @staticmethod
    def _symbol(m, T, line, rect, n, nd, warn):
        sym = nd["symbol"]
        lay = "HMB-WARN" if warn else "HMB-EQPT"
        if sym == "gas_turbine":
            x, y = nd["at"]
            line([(x - 60, y + 25), (x - 5, y + 12), (x - 5, y - 12), (x - 60, y - 25)], lay, close=True)       # compressor
            rect(x - 5, y - 8, x + 13, y + 8, lay)                                                            # combustor
            line([(x + 13, y + 12), (x + 60, y + 25), (x + 60, y - 25), (x + 13, y - 12)], lay, close=True)     # turbine
            T("C", x - 32, y, 5, lay, "MIDDLE"); T("T", x + 38, y, 5, lay, "MIDDLE")
            T(nd["label"] + " (IEC)", x, y + 32, 3.2, lay, "CENTER")
        elif sym == "hrsg":
            x0, y0, x1, y1 = nd["box"]
            rect(x0, y0, x1, y1, lay, lineweight=50)
            secs = nd["sections"]
            ws = nd.get("section_widths") or [(x1 - x0) / len(secs)] * len(secs)
            edges = [x0]
            for w in ws:
                edges.append(edges[-1] + w * (x1 - x0) / sum(ws))
            for i, sname in enumerate(secs):
                if i:
                    line([(edges[i], y0), (edges[i], y1)], "HMB-DATA")
                T(sname, (edges[i] + edges[i + 1]) / 2 + 1, y0 + 5, 2.3, "HMB-DATA", "LEFT", 90)
            for name, i in nd.get("drums", []):
                cx = (edges[i] + edges[i + 1]) / 2
                m.add_ellipse((cx, y1 + 11), (10, 0), 0.42, dxfattribs={"layer": lay})
                line([(cx, y1), (cx, y1 + 6.8)], "HMB-DATA")
                T(f"{name} DRUM", cx, y1 + 19, 2.4, lay, "CENTER")
            T(nd["label"], x0 + 2, y1 + 25, 3.2, lay)
            T("GAS FLOW  -->", x0 + (x1 - x0) / 2, y1 - 6, 2.6, "HMB-DATA", "CENTER")
        elif sym == "steam_turbine":
            x, y = nd["at"]
            line([(x - 85, y + 10), (x - 45, y + 18), (x - 45, y - 18), (x - 85, y - 10)], lay, close=True)     # HP
            line([(x - 30, y + 10), (x + 10, y + 18), (x + 10, y - 18), (x - 30, y - 10)], lay, close=True)     # IP
            line([(x + 25, y + 26), (x + 55, y + 10), (x + 55, y - 10), (x + 25, y - 26)], lay, close=True)     # LP left flow
            line([(x + 55, y + 10), (x + 85, y + 26), (x + 85, y - 26), (x + 55, y - 10)], lay, close=True)     # LP right flow
            line([(x - 45, y), (x - 30, y)], lay, lineweight=70); line([(x + 10, y), (x + 25, y)], lay, lineweight=70)
            T("HP", x - 65, y, 4, lay, "MIDDLE"); T("IP", x - 10, y, 4, lay, "MIDDLE"); T("LP", x + 55, y + 30, 3.5, lay, "CENTER")
            T(nd["label"], x, y - 38, 3.2, lay, "CENTER")
        elif sym == "condenser":
            x0, y0, x1, y1 = nd["box"]
            rect(x0, y0, x1, y1, lay, lineweight=50)
            for k in range(3):
                yy = y0 + 12 + k * 9
                line([(x0 + 4, yy), (x1 - 4, yy)], "HMB-CW")
            T(nd["label"], (x0 + x1) / 2, y0 - 6, 2.8, lay, "CENTER")
        elif sym == "pump":
            x, y = nd["at"]
            m.add_circle((x, y), 6, dxfattribs={"layer": lay})
            line([(x - 3.5, y + 4.2), (x + 6, y), (x - 3.5, y - 4.2)], lay)
            T(nd["label"], x, y - 11, 2.6, lay, "CENTER")
        elif sym == "heat_exchanger":
            x0, y0, x1, y1 = nd["box"]
            rect(x0, y0, x1, y1, lay)
            line([(x0, y0), (x0 + 5, y1), (x0 + 10, y0), (x0 + 15, y1), (x1, y0)], "HMB-DATA")
            T(nd["label"], x0 - 2, y1 + 3, 2.6, lay, "RIGHT")
        elif sym == "junction":
            x, y = nd["at"]
            m.add_circle((x, y), 1.6, dxfattribs={"layer": lay})
        elif sym == "scope_zone":
            x0, y0, x1, y1 = nd["box"]
            rect(x0, y0, x1, y1, "HMB-SCOPE", linetype="DASHED")
            T(nd["label"], x0 + 3, y0 + 3, 2.8, "HMB-SCOPE")
        elif sym == "stack":
            x0, y0, x1, y1 = nd["box"]
            rect(x0, y0, x1, y1, lay)
            T(nd["label"], (x0 + x1) / 2 + 1, y0 + 30, 2.8, lay, "LEFT", 90)
        elif sym == "boundary":
            x, y = nd["at"]
            T(nd["label"], x, y, 2.8, lay, "CENTER")

    def _tables(self, m, T, line, rect, case, res, W, H):
        sm = res.summary
        x0, x1, y = 575, W - 12, H - 16
        T(f"STREAM TABLE - {case['id']}", x0, y, 3.5, "HMB-TABLE")
        cols = [("No", 8), ("Stream", 104), ("kg/s", 22), ("bar(a)", 20), ("degC / x", 20), ("kJ/kg", 20), ("MW*", 22)]
        y -= 4
        rh = 4.6
        xs = [x0]
        for _, w in cols:
            xs.append(xs[-1] + w)
        rect(x0, y - rh, xs[-1], y, "HMB-TABLE")
        for (c, _), xa, xb in zip(cols, xs, xs[1:]):
            T(c, xb - 1 if c not in ("Stream",) else xa + 1, y - rh + 1.3, 2.3, "HMB-TABLE", "RIGHT" if c not in ("Stream",) else "LEFT")
        for st in sorted(res.streams, key=lambda s: s["number"]):
            y -= rh
            q = st.get("quality")
            vals = [str(st["number"]), st["description"][:72], f"{st['mass_flow']:,.2f}",
                    f"{st['pressure']:.4f}" if st["pressure"] < 0.1 else f"{st['pressure']:.2f}",
                    f"{st['temperature']:.1f}" if q is None else f"x {q:.3f}", f"{st['h']:,.1f}", f"{st['E'] / 1000:,.1f}"]
            for i, (v, xa, xb) in enumerate(zip(vals, xs, xs[1:])):
                T(v, xa + 1 if i == 1 else xb - 1, y - rh + 1.3, 2.1, "HMB-TABLE", "LEFT" if i == 1 else "RIGHT")
            line([(x0, y - rh), (xs[-1], y - rh)], "HMB-DATA")
        for xv in xs:
            line([(xv, H - 20), (xv, y - rh)], "HMB-DATA")
        y -= rh + 3.5
        T("* energy flow m x h (reference: IAPWS-IF97 for water/steam, 25 degC sensible for gases, cp x T for seawater)", x0, y, 1.9, "HMB-DATA")
        # power balance
        y -= 8
        T("PLANT POWER BALANCE", x0, y, 3.5, "HMB-TABLE")
        rows = [("GT output (generator terminals)", sm["gt_output"] / 1000, "MW"), ("ST output (generator terminals)", sm["st_output"] / 1000, "MW"),
                ("Gross output", sm["gross"] / 1000, "MW"), ("Auxiliary loads", -sm["aux"] / 1000, "MW"),
                ("Transformer losses (GSU, UAT, MV/LV)", -sm["transformer_losses"] / 1000, "MW"),
                ("NET OUTPUT (HV terminals)", sm["net"] / 1000, "MW"), ("Heat input (LHV)", sm["heat_input"] / 1000, "MW"),
                ("Gross heat rate (LHV)", sm["gross_hr"], "kJ/kWh"), ("NET HEAT RATE (LHV)", sm["net_hr"], "kJ/kWh"),
                ("Net efficiency (LHV)", sm["net_eff"] * 100, "%"),
                ("Fuel LHV (ISO 6976, from composition)", sm["fuel_LHV"], "MJ/kg"), ("Fuel gas flow", sm["fuel_Sm3h"], "Sm3/h"),
                ("Circulating water flow", sm["cw_m3h"], "m3/h")]
        y -= 2
        for label, v, u in rows:
            if v is None:                      # e.g. no condenser in the case
                continue
            y -= 4.4
            bold = label.isupper()
            T(label, x0 + 1, y, 2.5 if bold else 2.3, "HMB-TABLE")
            T(f"{v:,.2f}" if abs(v) < 1000 else f"{v:,.0f}", x0 + 150, y, 2.5 if bold else 2.3, "HMB-TABLE", "RIGHT")
            T(u, x0 + 153, y, 2.3, "HMB-TABLE")
        # checks
        y -= 9
        T("CHECKS (guarantees, limits, balances)", x0, y, 3.5, "HMB-TABLE")
        y -= 2
        for st, text in res.checks:
            for k, chunk in enumerate(_wrap(text, 118)):
                y -= 3.6
                if k == 0:
                    T(st, x0 + 1, y, 2.1, "HMB-WARN" if st == "WARN" else "HMB-TABLE")
                T(chunk, x0 + 13, y, 2.1, "HMB-WARN" if st == "WARN" else "HMB-TABLE")

    # ------------------------------------------------------------------ workbook
    def _workbook(self, ctx, case, res, stem):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        head = PatternFill("solid", fgColor="1F3864")

        def sheet(ws, title, header, rows, widths):
            ws.append([title]); ws["A1"].font = Font(bold=True, size=12)
            ws.append([f"engine {self.name} v{self.version} | data {ctx.stamp['inputs_hash']} | {ctx.stamp['generated_at'][:19]}Z"])
            ws.append(header)
            for c in ws[3]:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), head
            for r in rows:
                ws.append(r)
            for i, w in enumerate(widths):
                ws.column_dimensions[chr(65 + i)].width = w
            ws.freeze_panes = "A4"

        ws = wb.active; ws.title = "Streams"
        sheet(ws, f"{stem} - streams", ["No", "Description", "Fluid", "From", "To", "Flow kg/s", "p bar(a)", "T degC", "Quality",
                                        "h kJ/kg", "Energy flow MW", "Status", "Basis"],
              [[s["number"], s["description"], s["fluid"], s["from_node"], s["to_node"], s["mass_flow"], s["pressure"], s["temperature"],
                s.get("quality"), round(s["h"], 2), round(s["E"] / 1000, 3), s.get("status"), ", ".join(s.get("basis_refs", []))]
               for s in sorted(res.streams, key=lambda s: s["number"])], [5, 55, 10, 8, 8, 11, 10, 9, 9, 10, 13, 12, 26])
        ws = wb.create_sheet("Node balances")
        sheet(ws, f"{stem} - node balances", ["Node", "Boundary", "Mass in kg/s", "Mass out kg/s", "Energy in MW", "Energy out MW",
                                              "Work MW", "Residual MW", "Expected loss MW", "Note"],
              [[n, b["boundary"], round(b["m_in"], 3), round(b["m_out"], 3), round(b["E_in"] / 1000, 3), round(b["E_out"] / 1000, 3),
                round(b["work"] / 1000, 3), round(b["residual"] / 1000, 3), round(b.get("expected_loss", 0) / 1000, 3), b["note"]]
               for n, b in sorted(res.nodes.items())], [8, 10, 13, 13, 13, 13, 10, 12, 15, 60])
        ws = wb.create_sheet("Auxiliary loads")
        sheet(ws, f"{stem} - auxiliary loads", ["Id", "Consumer", "Supplier", "Method", "kW", "Counts as", "Calculation", "Status", "Basis"],
              [[a["id"], a["description"], a["supplier"], a["method"], round(a["kW"], 1), a.get("counts_as", "auxiliary"), a["how"],
                a.get("status"), ", ".join(a.get("basis_refs", []))] for a in res.aux], [7, 60, 9, 11, 10, 16, 50, 12, 24])
        ws = wb.create_sheet("Summary")
        sm = res.summary
        sheet(ws, f"{stem} - summary", ["Item", "Value", "Unit"],
              [["GT output", sm["gt_output"] / 1000, "MW"], ["ST output", sm["st_output"] / 1000, "MW"], ["Gross output", sm["gross"] / 1000, "MW"],
               ["Auxiliary loads", sm["aux"] / 1000, "MW"], ["Transformer losses", sm["transformer_losses"] / 1000, "MW"],
               ["Net output", sm["net"] / 1000, "MW"], ["Heat input (LHV)", sm["heat_input"] / 1000, "MW"],
               ["Gross heat rate", sm["gross_hr"], "kJ/kWh"], ["Net heat rate", sm["net_hr"], "kJ/kWh"],
               ["Net efficiency", sm["net_eff"] * 100, "%"], ["Auxiliary + transformer losses", sm["aux_pct_gross"], "% of gross"],
               ["Fuel LHV (ISO 6976)", sm["fuel_LHV"], "MJ/kg"], ["Fuel LHV volumetric", sm["fuel_LHV_vol"], "MJ/Sm3"],
               ["Fuel flow", sm["fuel_flow"], "kg/s"], ["Fuel flow", sm["fuel_Sm3h"], "Sm3/h"],
               ["HRSG gas-side duty", (sm["hrsg_gas_duty"] or 0) / 1000, "MW"], ["CW flow", sm["cw_m3h"], "m3/h"],
               ["CW temperature rise", sm["cw_rise"], "K"], ["Condenser pressure", sm["condenser_pressure_mbar"], "mbar(a)"]]
              + [["Flue gas " + k + " (wet)", round(v * 100, 3), "mol %"] for k, v in sorted(res.flue_x.items())], [40, 14, 12])
        ws = wb.create_sheet("Checks")
        sheet(ws, f"{stem} - checks", ["Status", "Check"], [[a, b] for a, b in res.checks], [8, 150])
        for row in ws.iter_rows(min_row=4):
            if row[0].value == "WARN":
                row[0].font = Font(bold=True, color="C00000")
        out = ctx.out_dir / f"{stem}.xlsx"
        wb.save(out)
        return out


def _wrap(text, n):
    words, lines, cur = text.split(), [], ""
    for w in words:
        if len(cur) + len(w) + 1 > n:
            lines.append(cur)
            cur = w
        else:
            cur = f"{cur} {w}".strip()
    return lines + [cur]


ENGINE = HeatMassBalance()
