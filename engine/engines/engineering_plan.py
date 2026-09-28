"""Engineering plan: master document list (MDL) with the workflow network and timeline, the AWP structure
(CWA -> CWP -> EWP / PWP) and the coverage checks (Excel + timeline PDF).

Everything is derived from the records: documents (numbered per the KKS numbering rule, with inputs, EWP, MR, equipment,
requirements and scope items), doc_type (workflow rules), kks_key, system, equipment, cwa, cwp, ewp, mr and the CPM of the
CWP activities (engine/core/workflow.py). Checks:
- numbering and KKS: every document number follows ALP-ORG-KKS-DISC-TYPE-NNNN and matches its record; every system code,
  equipment tag and structure resolves against the project KKS key list;
- workflow: input network without loops, typical input types present, construction documents in an EWP, datasheets and
  requisitions in an MR;
- coverage: every system has its required document types (by category) and at least one document, every requirement and
  every EPC/IEC scope item is answered by a document, every equipment item has a datasheet, an MR and a CWP, every MR its
  specification and requisition, every CWP an EWP with documents;
- completeness: the MDL rules (mdl_rule) are in sync with the documents, and the quantities per discipline are within the
  indicative benchmarks (mdl_benchmark: below = WARN, above = INFO);
- plant data: every equipment item, system and document carries an AVEVA class (ER-01.06); the KKS key list and the AWP
  definitions are agreed (INFO while proposed);
- timeline: EWP float (IFC + lead before the CWP start), PWP float (delivery before the CWP start), key dates.
Options: pdf=no."""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, timedelta

from ..core import kks, mdl, workflow
from ..core.runner import Context, Engine

MONTHS = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()


def _short(text, n=58):
    """Gantt row label: whole words, ellipsis when shortened."""
    if len(text) <= n:
        return text
    cut = text[:n - 1].rsplit(" ", 1)[0].rstrip(" ,(-/")
    return cut + "\u2026"


def check_engineering(s, res) -> list[tuple[str, str, str, str]]:
    out = []

    def add(item, chk, st, txt):
        out.append((item, chk, st, txt))

    kk = kks.keys(s)
    types = {t["id"]: t for t in s.records("doc_type")}
    systems = {x["id"]: x for x in s.records("system")}
    docs = [d for d in s.records("document") if d.get("status") != "cancelled"]
    eqs = s.records("equipment")
    mrs = {m["id"]: m for m in s.records("mr")}
    cwps = {c["id"]: c for c in s.records("cwp")}
    ewps = {e["id"]: e for e in s.records("ewp")}
    # --- numbering and KKS
    bad = 0
    for d in docs:
        for e in kks.check_document(d, types, systems):
            add(d["id"], "numbering", "WARN", e)
            bad += 1
    add("-", "numbering", "OK" if not bad else "WARN", f"{len(docs)} document numbers checked, {bad} finding(s)")
    for code, sy in sorted(systems.items()):
        p, errs = kks.parse_system(code, kk)
        for e in errs:
            add(code, "KKS system", "WARN", e)
    for e in eqs:
        p, errs = kks.parse_tag(e["id"], kk)
        for x in errs:
            add(e["id"], "KKS tag", "WARN", x)
        if p and e.get("system") and not p["system"].startswith(e["system"]):
            add(e["id"], "KKS tag", "WARN", f"tag system {p['system']} differs from system {e['system']}")
    unver = [k for k, v in {**kk["F"], **kk["A"], **kk["B"]}.items() if not v.get("verified")]
    if unver:
        add("-", "KKS key list", "INFO", f"{len(unver)} key titles not yet verified against the licensed VGB key list")
    for ent, recs in (("equipment", eqs), ("system", list(systems.values())), ("document", docs)):
        miss = [r["id"] for r in recs if not r.get("aveva_class")]
        for i in miss:
            add(i, "AVEVA class", "WARN", f"{ent} without aveva_class (ER-01.06)")
        add("-", "AVEVA class", "OK" if not miss else "WARN", f"{len(recs) - len(miss)} of {len(recs)} {ent} records classified")
    for ent in ("kks_key", "cwa", "cwp", "ewp", "mr"):
        recs = s.records(ent)
        open_ = [r["id"] for r in recs if r.get("status") != "agreed" and r.get("status") != "superseded"]
        add("-", "agreement", "INFO" if open_ else "OK",
            f"{ent}: {len(recs) - len(open_)} of {len(recs)} agreed" + (f"; not agreed: {', '.join(open_[:8])}"
                                                                        + (" ..." if len(open_) > 8 else "") if open_ else ""))
    # --- completeness: MDL rules in sync, quantities against the benchmarks
    if "mdl_rule" in s.schemas and s.records("mdl_rule"):
        plan = mdl.plan_sync(s)
        for e in plan.errors:
            add("-", "MDL rules", "WARN", e)
        add("-", "MDL rules", "OK" if not (plan.create or plan.update or plan.orphans) else "WARN",
            f"{len(s.records('mdl_rule'))} rules: {len(plan.create)} required document(s) missing, {len(plan.update)} to update, "
            f"{len(plan.orphans)} orphan(s)" + ("" if not (plan.create or plan.update or plan.orphans)
                                                 else " - run python -m engine mdl sync"))
    if "mdl_benchmark" in s.schemas:
        for b in sorted(s.records("mdl_benchmark"), key=lambda x: x["id"]):
            sel = [d for d in docs if (b["discipline"] == "all" or d.get("discipline") == b["discipline"])
                   and ((d.get("originator") == "EPC") == (b["originator"] == "EPC") or b["originator"] == "ALL")]
            act = {"docs": len(sel), "sheets": sum(d.get("sheets") or 1 for d in sel), "hours": sum(d.get("weight") or 0 for d in sel)}
            for k in ("docs", "sheets", "hours"):
                lo, hi = b.get(f"{k}_min"), b.get(f"{k}_max")
                if lo is None and hi is None:
                    continue
                st = "WARN" if lo is not None and act[k] < lo else ("INFO" if hi is not None and act[k] > hi else "OK")
                rng = (f"{lo:,}" if lo is not None else "-") + " - " + (f"{hi:,}" if hi is not None else "-")
                add(b["id"], "benchmark", st, f"{b['discipline']} {b['originator']} {k} {act[k]:,.0f} vs indicative {rng}")
    # --- workflow rules
    ids = {d["id"] for d in docs}
    for d in docs:
        t = types.get(d.get("type_code"), {})
        if t.get("construction") and not d.get("ewp"):
            add(d["id"], "AWP", "WARN", f"construction document ({d.get('type_code')}) without EWP")
        if d.get("type_code") in ("DSH", "MRQ") and not d.get("mr") and d.get("originator") == "EPC":
            add(d["id"], "AWP", "WARN", f"{d['type_code']} without MR")
        want = set(t.get("input_types") or [])
        have = {next((x.get("type_code") for x in docs if x["id"] == i), None) for i in d.get("inputs", [])} if want else set()
        if want and d.get("originator") == "EPC" and not d.get("tender") and not (want & have):
            add(d["id"], "inputs", "INFO", f"none of the typical inputs {sorted(want)} of a {d.get('type_code')}")
        for i in d.get("inputs", []):
            if i not in ids:
                add(d["id"], "inputs", "WARN", f"input {i} is cancelled or missing")
    # --- coverage: systems
    covered = defaultdict(set)
    for d in docs:
        for sy in [d.get("system")] + list(d.get("systems") or []):
            if sy:
                covered[sy].add(d.get("type_code"))
    req_by_cat = defaultdict(set)
    for t in types.values():
        for c in t.get("required_for") or []:
            req_by_cat[c].add(t["id"])
    for code, sy in sorted(systems.items()):
        if len(code) < 5:
            continue
        have = covered.get(code, set())
        need = req_by_cat.get(sy.get("category"), set()) if sy.get("scope") == "IEPC" else set()
        miss = sorted(need - have)
        if not have:
            add(code, "coverage system", "WARN", f"{sy['title']}: no document")
        elif miss:
            add(code, "coverage system", "WARN", f"{sy['title']}: missing {', '.join(miss)}")
    # --- coverage: requirements and scope items
    cited = defaultdict(list)
    for d in docs:
        for r in d.get("basis_refs", []):
            cited[r].append(d["id"])
    reqs = s.records("requirement")
    miss = [r["id"] for r in reqs if f"requirement:{r['id']}" not in cited]
    for r in miss:
        add(r, "coverage requirement", "WARN", "no document answers this requirement")
    add("-", "coverage requirement", "OK" if not miss else "WARN", f"{len(reqs) - len(miss)} of {len(reqs)} requirements answered")
    scope = [x for x in s.records("scope_item") if x.get("design") in ("IEPC", "IEC", None) and x.get("supply") != "OWNER"]
    miss = [x["id"] for x in scope if f"scope_item:{x['id']}" not in cited]
    for x in miss:
        add(x, "coverage scope", "WARN", "EPC / IEC scope item without a document")
    add("-", "coverage scope", "OK" if not miss else "WARN", f"{len(scope) - len(miss)} of {len(scope)} EPC/IEC scope items covered")
    # --- coverage: equipment, MRs, CWPs, EWPs
    dsh = defaultdict(list)
    for d in docs:
        if d.get("type_code") == "DSH":
            for e in d.get("equipment", []):
                dsh[e].append(d["id"])
    for e in eqs:
        m = [k for k, v in (("datasheet", dsh.get(e["id"])), ("MR", e.get("mr")), ("CWP", e.get("cwp"))) if not v]
        if m:
            add(e["id"], "coverage equipment", "WARN", f"{e['description']}: missing {', '.join(m)}")
    by_mr = defaultdict(set)
    for d in docs:
        if d.get("mr"):
            by_mr[d["mr"]].add(d.get("type_code"))
    for m in mrs.values():
        if m["package_type"] != "iec_supply":
            miss = {"SPC", "MRQ"} - by_mr.get(m["id"], set())
            if miss:
                add(m["id"], "coverage MR", "WARN", f"missing {', '.join(sorted(miss))}")
        for c in m.get("cwps", []):
            if c not in cwps:
                add(m["id"], "coverage MR", "WARN", f"CWP {c} unknown")
    ewp_of = defaultdict(list)
    for e in ewps.values():
        ewp_of[e["cwp"]].append(e["id"])
    for c in cwps:
        if not ewp_of.get(c):
            add(c, "coverage CWP", "WARN", "CWP without EWP")
    for e, v in res.ewps.items():
        if not v["docs"]:
            add(e, "coverage EWP", "WARN", "EWP without documents")
    # --- timeline
    d = res.d
    for e, v in sorted(res.ewps.items()):
        if v["float"] is not None:
            st = "WARN" if v["float"] < 0 else ("INFO" if v["float"] < 10 else "OK")
            add(e, "EWP float", st, f"IFC {d(v['ready'])} ({v['driver']}), needed {d(v['need'])} for {v['cwp']} start "
                                    f"{d(v['cwp_start'])}: float {v['float']} wd")
        if v["vendor_float"] is not None and v["vendor_float"] < 0:
            add(e, "EWP vendor data", "WARN", f"certified vendor data of {v['vendor_mr']} {d(v['vendor_data'])} after the need "
                                              f"date {d(v['need'])}")
    for m, v in sorted(res.mrs.items()):
        if v["float"] is not None:
            st = "WARN" if v["float"] < 0 else ("INFO" if v["float"] < 10 else "OK")
            add(m, "PWP float", st, f"PO {d(v['po'])}, on site {d(v['ros'])}, needed {d(v['need'])} ({v['need_cwp']}): "
                                    f"float {v['float']} wd")
    for k in ("KD-010", "KD-020", "KD-030", "KD-040"):
        a = res.acts.get(k)
        if a:
            add(k, "key date", "WARN" if a.tf < 0 else "OK",
                f"{a.title}: forecast {res.cal.finish_date(a.ef, a.es)}, float {a.tf} wd")
    return out


class EngineeringPlan(Engine):
    name = "engineering_plan"
    title = "Engineering plan: MDL, workflow network and timeline, AWP (CWA/CWP/EWP/PWP), coverage checks (Excel + PDF)"
    version = "1.3.0"
    inputs = ["project", "document", "document_revision", "doc_type", "kks_key", "system", "equipment", "cwa", "cwp", "ewp",
              "mr", "activity", "wbs", "requirement", "scope_item", "party", "eng_resource", "mdl_rule", "mdl_benchmark",
              "instrument", "line", "decision"]
    formats = ["xlsx", "pdf"]
    code_deps = ["engine/core/workflow.py", "engine/core/kks.py", "engine/core/planning.py", "engine/core/mdl.py"]

    def run(self, ctx: Context):
        s = ctx.store
        if not s.records("document"):
            ctx.warnings.append("no document records - nothing generated")
            return []
        res = workflow.compute(s)
        checks = check_engineering(s, res)
        for item, chk, st, txt in checks:
            if st == "WARN":
                ctx.warnings.append(f"{item} {chk}: {txt}")
        stem = ctx.project_record()["id"]
        files = [self._workbook(ctx, res, checks, f"{stem}_MDL")]
        if ctx.options.get("pdf", "yes") != "no":
            files.append(self._timeline(ctx, res, f"{stem}_engineering_timeline"))
        return files

    # ------------------------------------------------------------------ Excel
    def _workbook(self, ctx, res, checks, stem):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        s = ctx.store
        head = PatternFill("solid", fgColor="1F3864")
        fills = {"WARN": PatternFill("solid", fgColor="F8CBAD"), "INFO": PatternFill("solid", fgColor="FFF2CC"),
                 "OK": PatternFill("solid", fgColor="E2EFDA")}
        wb = Workbook()
        d = res.d

        def sheet(ws, title, cols, rows, wrap=()):
            ws["A1"] = f"{ctx.project_record().get('name', '')} - {title}"
            ws["A1"].font = Font(bold=True, size=14)
            ws["A2"] = (f"engine {self.name} v{self.version} | generated {ctx.stamp['generated_at'][:19]}Z by "
                        f"{ctx.stamp['generated_by']} | data {ctx.stamp['inputs_hash']} | dates computed from the "
                        f"workflow network and the CPM (working days)")
            ws["A2"].font = Font(italic=True, size=8, color="595959")
            for c, (n, w) in enumerate(cols, 1):
                cell = ws.cell(4, c, n)
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                ws.column_dimensions[get_column_letter(c)].width = w
            for i, row in enumerate(rows, 5):
                for c, v in enumerate(row, 1):
                    v = "\n".join(map(str, v)) if isinstance(v, (list, tuple, set)) else v
                    ws.cell(i, c, v).alignment = Alignment(wrap_text=c in wrap, vertical="top")
            ws.freeze_panes = "B5"
            ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}4"
            ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            return ws

        types = {t["id"]: t for t in s.records("doc_type")}
        systems = {x["id"]: x for x in s.records("system")}
        cwps = {c["id"]: c for c in s.records("cwp")}
        ewps = {e["id"]: e for e in s.records("ewp")}
        latest = {}
        for r in s.records("document_revision"):
            if r["document"] not in latest or r["issue_date"] >= latest[r["document"]]["issue_date"]:
                latest[r["document"]] = r
        rows = []
        for k in sorted(res.docs):
            x = res.docs[k]
            r = x.rec
            t = types.get(r.get("type_code"), {})
            e = res.ewps.get(r.get("ewp"), {})
            need = e.get("need")
            lr = latest.get(k)
            rows.append([k, r["title"], r.get("originator"), (r.get("system") or "00000"),
                         systems.get(r.get("system"), {}).get("title", "plant general"), r["discipline"], r.get("type_code"),
                         t.get("dcc"), r.get("review") or t.get("review"), "yes" if r.get("bdp") else "",
                         "yes" if r.get("tender") else "", r.get("sheets"), r.get("weight"), r.get("ewp"),
                         ewps.get(r.get("ewp"), {}).get("cwp"), r.get("cwa") or cwps.get(ewps.get(r.get("ewp"), {}).get("cwp"), {}).get("cwa"),
                         r.get("mr"), r.get("equipment"), len(r.get("inputs", [])), d(x.start), d(x.ifr), d(x.ifa), d(x.ifc),
                         d(need), None if need is None else need - x.ifc, x.driver,
                         r.get("vendor_ref"), f"{lr['revision']} {lr['purpose']} {lr['issue_date']}" if lr else ""])
        sheet(wb.active, "Master document list (MDL)",
              [("Document number", 30), ("Title", 55), ("Org", 6), ("KKS", 7), ("System", 28), ("Discipline", 12), ("Type", 6),
               ("DCC", 5), ("Owner review", 10), ("BDP", 5), ("Tender", 6), ("Sheets", 6), ("Hours", 7), ("EWP", 14),
               ("CWP", 14), ("CWA", 8), ("MR", 11), ("Equipment", 16), ("Inputs", 6), ("Start", 11), ("IFR", 11), ("IFA", 11),
               ("IFC", 11), ("EWP need", 11), ("Float wd", 8), ("Driving input", 30), ("Vendor ref", 18), ("Latest issue", 18)],
              rows, wrap=(2, 18))
        wb.active.title = "MDL"
        edges = []
        for k in sorted(res.docs):
            x = res.docs[k]
            mat = types.get(x.rec.get("type_code"), {}).get("input_maturity", "IFR")
            for i in x.rec.get("inputs", []):
                y = res.docs.get(i)
                if y:
                    edges.append([k, x.rec.get("type_code"), i, y.rec.get("type_code"), mat, d(y.ifr if mat == "IFR" else y.ifc),
                                  "driving" if x.driver == i else ""])
        sheet(wb.create_sheet("Workflow"), "Workflow network (document <- input)",
              [("Document", 30), ("Type", 6), ("Input", 30), ("Input type", 8), ("Needs input at", 9), ("Input ready", 11),
               ("Driving", 8)], edges)
        rows = [[e, v["cwa"], v["cwp"], ewps[e]["title"], len(v["docs"]), d(v["ready"]), v["driver"], d(v["cwp_start"]),
                 d(v["need"]), v["float"], v["vendor_mr"], d(v["vendor_data"]), v["vendor_float"]]
                for e, v in sorted(res.ewps.items())]
        sheet(wb.create_sheet("EWP"), "Engineering work packages",
              [("EWP", 14), ("CWA", 8), ("CWP", 14), ("Title", 45), ("Docs", 6), ("IFC complete", 11), ("Last document", 30),
               ("CWP start", 11), ("Needed", 11), ("Float wd", 8), ("Vendor data MR", 11), ("Vendor data", 11),
               ("Vendor float", 8)], rows, wrap=(4,))
        mrr = {m["id"]: m for m in s.records("mr")}
        rows = [[m, mrr[m]["title"], mrr[m]["package_type"], mrr[m].get("supplier") or "", len(v["docs"]), d(v["issue"]),
                 d(v["po"]), d(v["vdr"]), d(v["ros"]), v["need_cwp"], d(v["need"]), v["float"], mrr[m].get("cwps")]
                for m, v in sorted(res.mrs.items())]
        sheet(wb.create_sheet("PWP (MR)"), "Procurement work packages (material requisitions)",
              [("MR", 11), ("Title", 45), ("Type", 11), ("Supplier", 9), ("Docs", 6), ("MR issue", 11), ("PO", 11),
               ("Vendor data", 11), ("On site (ROS)", 11), ("Needed by CWP", 14), ("Needed", 11), ("Float wd", 8), ("CWPs", 16)],
              rows, wrap=(2, 13))
        rows = []
        for c in sorted(cwps.values(), key=lambda c: c["id"]):
            a = res.acts.get(c.get("activity"))
            rows.append([c["id"], c["cwa"], c["discipline"], c["title"], c.get("contractor"),
                         d(a.es) if a else None, res.cal.finish_date(a.ef, a.es) if a else None, a.tf if a else None,
                         ", ".join(sorted(e for e, v in res.ewps.items() if v["cwp"] == c["id"])),
                         ", ".join(sorted(m for m in mrr if c["id"] in mrr[m].get("cwps", [])))])
        sheet(wb.create_sheet("CWP"), "Construction work packages (CPM dates)",
              [("CWP", 14), ("CWA", 8), ("Discipline", 11), ("Title", 45), ("Contractor", 22), ("Start", 11), ("Finish", 11),
               ("Total float wd", 8), ("EWP", 14), ("PWP (MR)", 40)], rows, wrap=(4, 10))
        rows = [[c["id"], c["title"], c["path_seq"], f"E {c.get('e_min')}-{c.get('e_max')}, N {c.get('n_min')}-{c.get('n_max')}",
                 c.get("structures"), c.get("description"), ", ".join(sorted(x for x in cwps if cwps[x]["cwa"] == c["id"]))]
                for c in sorted(s.records("cwa"), key=lambda c: c["id"])]
        sheet(wb.create_sheet("CWA"), "Construction work areas (path of construction)",
              [("CWA", 8), ("Title", 40), ("Path", 6), ("Extent (plant grid m)", 26), ("Structures", 16), ("Description", 60),
               ("CWPs", 50)], rows, wrap=(5, 6, 7))
        # coverage
        by_sys = defaultdict(Counter)
        for x in res.docs.values():
            for sy in [x.rec.get("system")] + list(x.rec.get("systems") or []):
                if sy:
                    by_sys[sy][x.rec.get("type_code")] += 1
        req_by_cat = defaultdict(set)
        for t in types.values():
            for c in t.get("required_for") or []:
                req_by_cat[c].add(t["id"])
        rows = []
        for code, sy in sorted(systems.items()):
            if len(code) < 5:
                continue
            need = req_by_cat.get(sy.get("category"), set()) if sy.get("scope") == "IEPC" else set()
            rows.append([code, sy["title"], sy.get("category"), sy.get("scope"), sum(by_sys[code].values()),
                         ", ".join(f"{k} {v}" for k, v in sorted(by_sys[code].items())), ", ".join(sorted(need - set(by_sys[code])))])
        sheet(wb.create_sheet("Coverage systems"), "Coverage: documents per KKS system",
              [("System", 8), ("Title", 40), ("Category", 18), ("Scope", 7), ("Docs", 6), ("Types", 60), ("Missing required", 20)],
              rows, wrap=(6,))
        cited = defaultdict(list)
        for x in res.docs.values():
            for r in x.rec.get("basis_refs", []):
                cited[r].append(x.id)
        rows = [[r["id"], r["title"], len(cited.get(f"requirement:{r['id']}", [])), cited.get(f"requirement:{r['id']}", [])]
                for r in sorted(s.records("requirement"), key=lambda r: r["id"])]
        sheet(wb.create_sheet("Coverage requirements"), "Coverage: Employer's Requirements answered by documents",
              [("Requirement", 10), ("Title", 40), ("Docs", 6), ("Documents", 60)], rows, wrap=(4,))
        rows = [[x["id"], x["area"], x["item"], x.get("design"), len(cited.get(f"scope_item:{x['id']}", [])),
                 cited.get(f"scope_item:{x['id']}", [])] for x in sorted(s.records("scope_item"), key=lambda x: x["id"])]
        sheet(wb.create_sheet("Coverage scope"), "Coverage: scope items (Owner contract SC-, consortium split DR-)",
              [("Item", 8), ("Area", 20), ("Scope", 55), ("Design", 7), ("Docs", 6), ("Documents", 50)], rows, wrap=(3, 6))
        dsh = defaultdict(list)
        for x in res.docs.values():
            if x.rec.get("type_code") == "DSH":
                for e in x.rec.get("equipment", []):
                    dsh[e].append(x.id)
        rows = [[e["id"], e["description"], e.get("system"), e.get("supply"), dsh.get(e["id"]), e.get("mr"), e.get("cwp"),
                 e.get("cwa")] for e in sorted(s.records("equipment"), key=lambda e: e["id"])]
        sheet(wb.create_sheet("Plant assets"), "Plant assets: KKS tag, datasheet, PWP, CWP, CWA",
              [("Tag", 16), ("Description", 45), ("System", 8), ("Supply", 7), ("Datasheet", 30), ("MR", 11), ("CWP", 14),
               ("CWA", 8)], rows, wrap=(2, 5))
        wc = sheet(wb.create_sheet("Checks"), "Checks", [("Item", 30), ("Check", 20), ("Result", 8), ("Detail", 120)], checks,
                   wrap=(4,))
        for i, c in enumerate(checks, 5):
            wc.cell(i, 3).fill = fills[c[2]]
        # summary
        rows = []
        for (org, disc), grp in sorted(Counter((x.rec.get("originator"), x.rec["discipline"]) for x in res.docs.values()).items()):
            ds = [x for x in res.docs.values() if (x.rec.get("originator"), x.rec["discipline"]) == (org, disc)]
            rows.append([org, disc, len(ds), sum(x.rec.get("sheets") or 1 for x in ds),
                         round(sum(x.rec.get("weight") or 0 for x in ds)), sum(1 for x in ds if x.rec.get("bdp")),
                         sum(1 for x in ds if x.rec.get("ewp")), sum(1 for x in ds if x.rec.get("mr")),
                         d(max(x.ifc for x in ds))])
        rows.append(["Total", "", len(res.docs), sum(x.rec.get("sheets") or 1 for x in res.docs.values()),
                     round(sum(x.rec.get("weight") or 0 for x in res.docs.values())),
                     sum(1 for x in res.docs.values() if x.rec.get("bdp")), sum(1 for x in res.docs.values() if x.rec.get("ewp")),
                     sum(1 for x in res.docs.values() if x.rec.get("mr")), d(max(x.ifc for x in res.docs.values()))])
        sheet(wb.create_sheet("Summary"), "Summary by originator and discipline",
              [("Originator", 10), ("Discipline", 14), ("Documents", 10), ("Sheets", 8), ("Hours", 9), ("BDP", 6),
               ("In EWP", 8), ("In MR", 7), ("Last IFC", 11)], rows)
        void = sorted((d for d in s.records("document") if d.get("status") == "cancelled"), key=lambda d: d["id"])
        sheet(wb.create_sheet("Void numbers"), "Cancelled document numbers (kept, never reused)",
              [("Document number", 30), ("Title", 60), ("Rule", 34), ("Remarks", 60)],
              [[d["id"], d.get("title"), d.get("rule"), d.get("remarks")] for d in void], wrap=(1, 3))
        out = ctx.out_dir / f"{stem}.xlsx"
        wb.save(out)
        return out

    # ------------------------------------------------------------------ timeline PDF (A3 landscape, several pages)
    def _timeline(self, ctx, res, stem):
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.dates as mdates
        import matplotlib.pyplot as plt
        from matplotlib.backends.backend_pdf import PdfPages
        s = ctx.store
        d = res.d
        cwas = sorted(s.records("cwa"), key=lambda c: (c["path_seq"], c["id"]))
        cwps = {c["id"]: c for c in s.records("cwp")}
        mrr = {m["id"]: m for m in s.records("mr")}
        kds = [(k, res.acts[k]) for k in ("KD-010", "KD-020", "KD-030", "KD-040") if k in res.acts]
        x0, x1 = d(res.dd - 5), max(d(max(x.ifc for x in res.docs.values())), res.cal.finish_date(
            max(a.ef for a in res.acts.values()), 0)) if res.acts else d(res.dd + 600)
        foot = (f"{ctx.project_record().get('name', '')} | engine {self.name} v{self.version} | data {ctx.stamp['inputs_hash']} "
                f"| generated {ctx.stamp['generated_at'][:10]} - dates computed from the MDL workflow network and the CPM")
        out = ctx.out_dir / f"{stem}.pdf"
        with PdfPages(out) as pdf:
            # page 1: engineering progress curve and issues per month
            fig = plt.figure(figsize=(420 / 25.4, 297 / 25.4))
            ax = fig.add_axes([0.06, 0.40, 0.88, 0.50])
            for key, lab, col in (("ifr", "first issue (IFR), cumulative hours", "#4472C4"),
                                  ("ifc", "issued for construction / final (IFC), cumulative hours", "#C00000")):
                pts = sorted((getattr(x, key), x.rec.get("weight") or 0) for x in res.docs.values())
                tot, xs, ys = 0, [], []
                for i, w in pts:
                    tot += w
                    xs.append(d(i))
                    ys.append(tot)
                ax.step(xs, ys, where="post", color=col, label=lab)
            for k, a in kds:
                ax.axvline(res.cal.finish_date(a.ef, a.es), color="grey", ls="--", lw=0.8)
                ax.text(res.cal.finish_date(a.ef, a.es), ax.get_ylim()[1] * 0.02, f" {k}", rotation=90, fontsize=7,
                        color="grey", va="bottom")
            tender = [x for x in res.docs.values() if x.rec.get("tender")]
            ax.set_title(f"Engineering progress curve - {len(res.docs)} documents, "
                         f"{round(sum(x.rec.get('weight') or 0 for x in res.docs.values())):,} h (step at NTP: "
                         f"{len(tender)} tender-stage documents, {round(sum(x.rec.get('weight') or 0 for x in tender)):,} h; "
                         f"supplier documents at their VDRL dates)" + (" - levelled to the discipline capacities"
                                                                       if res.levelled else ""), fontsize=11, loc="left")
            ax.set_ylabel("hours")
            ax.legend(loc="upper left", fontsize=8)
            ax.grid(alpha=0.3)
            ax.xaxis.set_major_formatter(mdates.DateFormatter("%b %Y"))
            ax2 = fig.add_axes([0.06, 0.08, 0.88, 0.24], sharex=ax)
            months = Counter((d(x.ifc).year, d(x.ifc).month) for x in res.docs.values())
            ks = sorted(months)
            ax2.bar([date(y, m, 15) for y, m in ks], [months[k] for k in ks], width=24, color="#A5A5A5")
            for (y, m) in ks:
                ax2.text(date(y, m, 15), months[(y, m)], str(months[(y, m)]), ha="center", va="bottom", fontsize=6)
            ax2.xaxis.set_major_formatter(mdates.DateFormatter("%b %Y"))
            ax2.set_title("Documents reaching IFC per month", fontsize=10, loc="left")
            ax2.grid(axis="y", alpha=0.3)
            fig.text(0.06, 0.015, foot, fontsize=6, color="#595959")
            pdf.savefig(fig)
            plt.close(fig)
            # AWP Gantt pages: per CWA -> EWP (engineering), MR (procurement), CWP (construction)
            rows = []
            for c in cwas:
                rows.append(("hdr", f"{c['id']} {c['title']} (path {c['path_seq']})", None))
                for cid in sorted(x for x in cwps if cwps[x]["cwa"] == c["id"]):
                    for e, v in sorted(res.ewps.items()):
                        if v["cwp"] == cid and v["docs"]:
                            first = min(res.docs[i].start for i in v["docs"])
                            rows.append(("ewp", e, (first, v["ready"], v["need"], v["float"])))
                    for m, v in sorted(res.mrs.items()):
                        if v["need_cwp"] == cid:
                            rows.append(("mr", _short(f"{m} {mrr[m]['title']}"), (v["po"], v["ros"], v["need"], v["float"])))
                    a = res.acts.get(cwps[cid].get("activity"))
                    if a:
                        rows.append(("cwp", _short(f"{cid} {cwps[cid]['title']}"), (a.es, a.ef, None, a.tf)))
            per, pages, p0 = 46, [], 0
            while p0 < len(rows):
                n = min(per, len(rows) - p0)
                while n > 1 and p0 + n < len(rows) and rows[p0 + n - 1][0] == "hdr":   # no header at the page end
                    n -= 1
                pages.append(rows[p0:p0 + n])
                p0 += n
            x1 = x1 + timedelta(days=30)
            for pno, chunk in enumerate(pages):
                fig = plt.figure(figsize=(420 / 25.4, 297 / 25.4))
                ax = fig.add_axes([0.27, 0.07, 0.70, 0.86])
                for i, (kind, lab, v) in enumerate(chunk):
                    y = len(chunk) - i
                    if kind == "hdr":
                        ax.axhspan(y - 0.5, y + 0.5, color="#D9E1F2")
                        ax.text(-0.003, y, lab, transform=ax.get_yaxis_transform(), ha="right", va="center", fontsize=7,
                                weight="bold")
                        continue
                    a0, a1, need, fl = v
                    col = {"ewp": "#4472C4", "mr": "#ED7D31", "cwp": "#70AD47"}[kind]
                    late = fl is not None and fl < 0 and kind != "cwp"
                    ax.barh(y, (d(a1) - d(a0)).days or 1, left=d(a0), height=0.6, color="#C00000" if late else col)
                    if need is not None:
                        ax.plot([d(need)], [y], marker="|", color="black", markersize=9, mew=2)
                    ax.text(-0.003, y, lab, transform=ax.get_yaxis_transform(), ha="right", va="center", fontsize=6.5)
                    if fl is not None:
                        ax.text(max(d(a1), d(need)) if need is not None else d(a1), y, f"  {fl:+d}", va="center",
                                fontsize=5.5, color="#C00000" if late else "#404040")
                for k, a in kds:
                    ax.axvline(res.cal.finish_date(a.ef, a.es), color="grey", ls="--", lw=0.8)
                    ax.text(res.cal.finish_date(a.ef, a.es), len(chunk) + 0.6, k, fontsize=6.5, color="grey", ha="center")
                ax.set_ylim(0.3, len(chunk) + 1.1)
                ax.set_xlim(x0, x1)
                ax.set_yticks([])
                ax.xaxis.set_major_locator(mdates.MonthLocator(bymonth=(1, 4, 7, 10)))
                ax.xaxis.set_major_formatter(mdates.DateFormatter("%b %Y"))
                ax.tick_params(axis="x", labelsize=7)
                ax.grid(axis="x", alpha=0.3)
                ax.set_title(f"AWP timeline - EWP (blue: first document start to all IFC), PWP (orange: PO to on site), "
                             f"CWP (green: CPM)\n| = first need (EWP: CWP start - lead; progressive documents later); "
                             f"numbers = float in working days (EWP: smallest document float); red = late   "
                             f"[page {pno + 2}]", fontsize=8, loc="left")
                fig.text(0.02, 0.015, foot, fontsize=6, color="#595959")
                pdf.savefig(fig)
                plt.close(fig)
        return out


ENGINE = EngineeringPlan()
