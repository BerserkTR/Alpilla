"""Equipment list (.xlsx) from system + equipment records. Reference example for new engines."""
from __future__ import annotations

from ..core.runner import Context, Engine

COLUMNS = ["id", "description", "system", "aveva_class", "equipment_type", "service", "quantity", "redundancy", "rated_power",
           "voltage", "design_flow", "design_pressure", "design_temperature", "material", "location",
           "manufacturer", "model", "status", "datasheet", "basis_refs", "remarks"]


class EquipmentList(Engine):
    name = "equipment_list"
    title = "Equipment list (Excel) grouped by system"
    version = "1.2.0"
    inputs = ["project", "system", "equipment"]
    formats = ["xlsx"]

    def run(self, ctx: Context):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        s = ctx.store
        fields = s.schema("equipment").fields
        prj = ctx.project_record()
        systems = {r["id"]: r for r in s.records("system")}
        items = sorted(s.records("equipment"), key=lambda r: (r["system"], r["id"]))
        if ctx.options.get("include_deleted", "no") != "yes":
            items = [r for r in items if r.get("status") != "deleted"]

        wb = Workbook()
        ws = wb.active
        ws.title = "Equipment List"
        bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="1F3864")
        thin = Side(style="thin", color="A6A6A6")
        grid = Border(left=thin, right=thin, top=thin, bottom=thin)

        ws["A1"] = f"{prj.get('name')} - Equipment List"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = (f"Project {prj['id']} | generated {ctx.stamp['generated_at'][:19]}Z by {ctx.stamp['generated_by']} "
                    f"| engine {self.name} v{self.version} | data {ctx.stamp['inputs_hash']} | git {ctx.stamp['git_commit']}")
        ws["A2"].font = Font(italic=True, size=8, color="595959")

        hdr_row = 4
        for c, f in enumerate(COLUMNS, 1):
            unit = fields.get(f, {}).get("unit")
            label = "Tag" if f == "id" else f.replace("_", " ").title() + (f" [{unit}]" if unit else "")
            cell = ws.cell(hdr_row, c, label)
            cell.font, cell.fill, cell.border = Font(bold=True, color="FFFFFF"), head_fill, grid
            cell.alignment = Alignment(wrap_text=True, vertical="center")

        row, current = hdr_row + 1, None
        for r in items:
            if r["system"] != current:
                current = r["system"]
                sysrec = systems.get(current, {})
                ws.cell(row, 1, f"{current} - {sysrec.get('title', '')}").font = bold
                for c in range(1, len(COLUMNS) + 1):
                    ws.cell(row, c).fill = PatternFill("solid", fgColor="D9E1F2")
                row += 1
            for c, f in enumerate(COLUMNS, 1):
                v = r.get(f)
                cell = ws.cell(row, c, ", ".join(v) if isinstance(v, list) else v)
                cell.border = grid
                cell.alignment = Alignment(vertical="top", wrap_text=f in ("description", "remarks", "basis_refs"))
            row += 1
        if not items:
            ws.cell(row, 1, "No equipment records in the database yet.")

        widths = {"aveva_class": 22, "id": 16, "description": 34, "remarks": 30, "basis_refs": 26, "service": 18, "datasheet": 20}
        for c, f in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(c)].width = widths.get(f, 13)
        ws.freeze_panes = ws.cell(hdr_row + 1, 2)
        ws.auto_filter.ref = f"A{hdr_row}:{get_column_letter(len(COLUMNS))}{max(row - 1, hdr_row)}"
        ws.print_title_rows = f"{hdr_row}:{hdr_row}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        ss = wb.create_sheet("Systems")
        for c, h in enumerate(["System", "Title", "Category", "Unit", "Parent", "Equipment count"], 1):
            ss.cell(1, c, h).font = bold
        counts = {}
        for r in items:
            counts[r["system"]] = counts.get(r["system"], 0) + 1
        for i, sy in enumerate(sorted(systems.values(), key=lambda x: x["id"]), 2):
            for c, v in enumerate([sy["id"], sy["title"], sy["category"], sy.get("unit_no"), sy.get("parent"),
                                   counts.get(sy["id"], 0)], 1):
                ss.cell(i, c, v)
        for col, w in zip("ABCDEF", (14, 40, 22, 8, 12, 16)):
            ss.column_dimensions[col].width = w
        ss.page_setup.fitToWidth, ss.page_setup.fitToHeight = 1, 0
        ss.sheet_properties.pageSetUpPr.fitToPage = True

        out = ctx.out_dir / f"{prj['id']}_equipment_list.xlsx"
        wb.save(out)
        return [out]


ENGINE = EquipmentList()
