"""Builds the default templates/docx/datasheet_base.docx (A4, margins, styles, header/footer).

Run: python -m engine.tools.make_docx_template
After that, the file can be restyled in Word; engines only rely on the built-in style names
(Heading 1/2, List Bullet, Table Grid) and on one paragraph in header and footer.
"""
from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor

from ..core.project import Project


def build(path):
    doc = Document()
    for sec in doc.sections:
        sec.page_height, sec.page_width = Mm(297), Mm(210)
        sec.left_margin = sec.right_margin = Mm(20)
        sec.top_margin, sec.bottom_margin = Mm(22), Mm(18)
        sec.header.paragraphs[0].text = "PROJECT | DOCUMENT NO. | REV"
        sec.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        sec.footer.paragraphs[0].text = "generation stamp"
        for p in (sec.header.paragraphs[0], sec.footer.paragraphs[0]):
            for r in p.runs:
                r.font.size = Pt(8)
                r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    st = doc.styles
    st["Normal"].font.name = "Arial"
    st["Normal"].font.size = Pt(10)
    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        rpr = st[name].element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is not None:
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                fonts.attrib.pop(qn(a), None)
        st[name].font.name = "Arial"
        st[name].font.size = Pt(size)
        st[name].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    doc.add_paragraph("Template body - replaced by the engine.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build(Project.locate().templates / "docx" / "datasheet_base.docx"))
