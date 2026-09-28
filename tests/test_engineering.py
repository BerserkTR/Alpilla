"""KKS identification, document numbering, engineering workflow timeline (incl. levelling), AWP, engineering_plan and
procedures engines, HMB export limiter."""
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from engine.cli import main
from engine.core import hmb, kks, planning, workflow
from engine.core.store import Store
from engine.engines import registry
from engine.core.runner import run_engine
from tests.conftest import who
from tests.test_planning import act, base

R = "test"
KK = {"F": {"PAC": {"title": "CW pumps"}, "MBA": {"title": "GT"}},
      "A": {"AP": {"title": "Pump units"}}, "B": {"KP": {"title": "Pump"}}}


# ---------------------------------------------------------------- KKS and numbering
def test_parse_tag_levels_and_keys():
    p, errs = kks.parse_tag("00PAC10AP001KP01", KK)
    assert not errs and p["system"] == "00PAC" and p["A"] == "AP" and p["AN"] == "001" and p["B"] == "KP"
    p, errs = kks.parse_tag("10MBA10", KK)
    assert not errs and p["FN"] == "10" and "A" not in p
    _, errs = kks.parse_tag("30XYZ10AA001", KK)
    assert any("unit '3'" in e for e in errs) and any("XYZ" in e for e in errs) and any("AA" in e for e in errs)
    assert kks.parse_tag("PAC-10", KK)[0] is None


def test_parse_system_group_and_function():
    assert kks.parse_system("00PAC", KK)[1] == []
    assert kks.parse_system("10MB", KK)[1] == []             # group of MBA
    assert kks.parse_system("10MX", KK)[1]                   # no key of group MX


def test_document_number_rules():
    p, errs = kks.parse_docno("ALP-EPC-00PAC-ME-DSH-0001")
    assert not errs and p == {"ORG": "EPC", "KKS": "00PAC", "DISC": "ME", "TYPE": "DSH", "SEQ": "0001"}
    assert kks.parse_docno("ALP-EPC-00PAC-XX-DSH-0001")[1]            # unknown discipline
    assert kks.parse_docno("ALP-DS-001")[0] is None
    types, systems = {"DSH": {}}, {"00PAC": {}}
    ok = {"id": "ALP-EPC-00PAC-ME-DSH-0001", "type_code": "DSH", "discipline": "mechanical", "originator": "EPC",
          "system": "00PAC"}
    assert kks.check_document(ok, types, systems) == []
    bad = dict(ok, type_code="SPC", discipline="piping", originator="IEC", system="00PAB")
    errs = " ".join(kks.check_document(bad, types, systems))
    assert "type_code SPC" in errs and "discipline piping" in errs and "originator IEC" in errs and "system 00PAB" in errs
    gen = {"id": "ALP-EPC-00000-GE-DSH-0001", "type_code": "DSH", "discipline": "general", "system": "00PAC"}
    assert "plant-general" in " ".join(kks.check_document(gen, types, systems))
    assert kks.next_number(["ALP-EPC-00PAC-ME-DSH-0001", "ALP-EPC-00PAC-ME-DSH-0007", "ALP-EPC-00PAC-ME-SPC-0009"],
                           "EPC", "00PAC", "ME", "DSH") == "ALP-EPC-00PAC-ME-DSH-0008"
    assert kks.next_number([], "IEC", "00000", "GE", "LST") == "ALP-IEC-00000-GE-LST-0001"
    assert any("KP01" in x for x in kks.explain("00PAC10AP001KP01", KK))


# ---------------------------------------------------------------- data set
def dtype(s, code, review, prep=10, upd=5, mat="IFR", **kw):
    s.create("doc_type", {"id": code, "title": code, "dcc": "F", "category": "drawing", "review": review, "prep_days": prep,
                          "update_days": upd, "input_maturity": mat, **kw}, R)


def doc(s, no, disc, t, **kw):
    kks_part = no.split("-")[2]
    s.create("document", {"id": no, "title": no, "discipline": disc, "doc_type": "drawing", "type_code": t,
                          "originator": no.split("-")[1], **({} if kks_part == "00000" else {"system": kks_part}), **kw}, R)


def data(project, capacity=None):
    s = base(project, start="2026-11-02", data_date="2026-11-02")
    s.create("wbs", {"id": "ALP.C", "title": "Construction", "parent": "ALP"}, R)
    s.create("party", {"id": "IEPC", "name": "Istanbul EPC", "role": "consortium_leader"}, R)
    s.create("system", {"id": "00PAC", "title": "CW pumps", "category": "cooling", "aveva_class": "System"}, R)
    for k, lvl, g in (("PAC", "function", "cooling water"), ("AP", "equipment_unit", "equipment unit")):
        s.create("kks_key", {"id": f"{'F' if lvl == 'function' else 'A'}-{k}", "level": lvl, "key": k, "main_group": g,
                             "title": k, "verified": False, "status": "agreed"}, R)
    dtype(s, "DBR", "approval", prep=20)
    dtype(s, "PID", "approval", prep=25, upd=15, input_types=["DBR"])
    dtype(s, "LST", "review", prep=15, upd=10)
    dtype(s, "ISO", "information", prep=20, mat="IFC", input_types=["PID"], construction=True)
    dtype(s, "DSH", "approval", prep=10, input_types=["PID"], procurement=True)
    dtype(s, "MRQ", "information", prep=5, upd=3, input_types=["DSH"], procurement=True)
    s.create("cwa", {"id": "CWA-05", "title": "CW pump house", "path_seq": 1, "status": "agreed"}, R)
    s.create("activity", {"id": "CWP-05-PI-01", "title": "CW piping", "wbs": "ALP.C", "phase": "construction", "duration": 60,
                          "constraint": "start_no_earlier_than", "constraint_date": "2027-06-01"}, R)
    s.create("cwp", {"id": "CWP-05-PI-01", "cwa": "CWA-05", "discipline": "piping", "title": "CW piping", "scope": "piping",
                     "activity": "CWP-05-PI-01", "ewp_lead_days": 20, "status": "agreed"}, R)
    s.create("ewp", {"id": "EWP-05-PI-01", "cwp": "CWP-05-PI-01", "discipline": "piping", "title": "IFC piping",
                     "status": "agreed"}, R)
    s.create("mr", {"id": "MR-ME-003", "title": "CW pumps", "discipline": "mechanical", "package_type": "equipment",
                    "cwps": ["CWP-05-PI-01"], "bid_days": 30, "award_days": 10, "manufacture_weeks": 30,
                    "transport_weeks": 4, "vdr_weeks": 8, "status": "agreed"}, R)
    s.create("equipment", {"id": "00PAC10AP001", "description": "CW pump A", "system": "00PAC", "equipment_type": "pump",
                           "aveva_class": "Centrifugal Pump", "cwa": "CWA-05", "cwp": "CWP-05-PI-01", "mr": "MR-ME-003"}, R)
    doc(s, "ALP-EPC-00000-PR-DBR-0001", "process", "DBR", tender=True, weight=80, aveva_class="Deliverables Documents")
    doc(s, "ALP-EPC-00PAC-PR-PID-0001", "process", "PID", weight=120, bdp=True, aveva_class="Drawing")
    doc(s, "ALP-EPC-00PAC-PR-LST-0001", "process", "LST", weight=150, aveva_class="Deliverables Documents")
    doc(s, "ALP-EPC-00PAC-PI-ISO-0001", "piping", "ISO", weight=200, ewp="EWP-05-PI-01", aveva_class="Drawing")
    doc(s, "ALP-EPC-00PAC-ME-DSH-0001", "mechanical", "DSH", weight=30, mr="MR-ME-003", equipment=["00PAC10AP001"],
        aveva_class="Specification Documents")
    doc(s, "ALP-EPC-00PAC-ME-MRQ-0001", "mechanical", "MRQ", weight=10, mr="MR-ME-003", aveva_class="Specification Documents")
    doc(s, "ALP-IEC-00000-GE-LST-0001", "general", "LST", planned_ifr="2026-12-14", planned_ifc="2027-01-11",
        aveva_class="Deliverables Documents")
    links = {"ALP-EPC-00PAC-PR-PID-0001": ["ALP-EPC-00000-PR-DBR-0001", "ALP-IEC-00000-GE-LST-0001"],
             "ALP-EPC-00PAC-PR-LST-0001": ["ALP-EPC-00000-PR-DBR-0001"],
             "ALP-EPC-00PAC-PI-ISO-0001": ["ALP-EPC-00PAC-PR-PID-0001"],
             "ALP-EPC-00PAC-ME-DSH-0001": ["ALP-EPC-00PAC-PR-PID-0001"],
             "ALP-EPC-00PAC-ME-MRQ-0001": ["ALP-EPC-00PAC-ME-DSH-0001"]}
    for k, v in links.items():
        s.update("document", k, {"inputs": v}, R)
    for disc, fte in (capacity or {}).items():
        s.create("eng_resource", {"id": f"RES-{kks.DISC_CODE[disc]}", "discipline": disc, "fte": fte}, R)
    return s


# ---------------------------------------------------------------- workflow timeline
def test_workflow_dates_review_classes_and_awp(project):
    s = data(project)
    r = workflow.compute(s)
    d, dd = r.docs, r.dd
    dbr, pid = d["ALP-EPC-00000-PR-DBR-0001"], d["ALP-EPC-00PAC-PR-PID-0001"]
    iec = d["ALP-IEC-00000-GE-LST-0001"]
    assert dbr.ifr == dd and dbr.ifa == dd + 10 + 5 and dbr.ifc == dbr.ifa + workflow.APPROVAL_CODE      # tender: IFR at NTP
    assert iec.fixed and r.d(iec.ifr).isoformat() == "2026-12-14"                                       # VDRL date kept
    assert pid.start == iec.ifr and pid.driver == iec.id                                               # latest input (IFR)
    assert pid.ifr == pid.start + 25 and pid.ifa == pid.ifr + workflow.OWNER_REVIEW_BDP + 15            # BDP review 21 d
    iso = d["ALP-EPC-00PAC-PI-ISO-0001"]
    assert iso.start == pid.ifc and iso.ifc == iso.ifr + 5                                             # IFC maturity, info
    lst = d["ALP-EPC-00PAC-PR-LST-0001"]
    assert lst.ifa is None and lst.ifc == lst.ifr + workflow.OWNER_REVIEW + 10                        # review class
    e = r.ewps["EWP-05-PI-01"]
    assert e["ready"] == iso.ifc and e["need"] == e["cwp_start"] - 20 and e["float"] == e["need"] - e["ready"]
    m = r.mrs["MR-ME-003"]
    mrq = d["ALP-EPC-00PAC-ME-MRQ-0001"]
    assert m["issue"] == mrq.ifc and m["po"] == mrq.ifc + 40 and m["ros"] == m["po"] + 5 * 34
    assert m["need_cwp"] == "CWP-05-PI-01" and m["float"] == m["need"] - m["ros"]
    assert not r.levelled


def test_workflow_input_loop_is_reported(project):
    s = data(project)
    s.update("document", "ALP-EPC-00000-PR-DBR-0001", {"inputs": ["ALP-EPC-00PAC-PI-ISO-0001"]}, R)
    with pytest.raises(planning.PlanningError, match="input loop"):
        workflow.compute(s)


def test_levelling_by_discipline_capacity(project):
    free = workflow.compute(data(project))
    s = Store(project, who())
    s.create("eng_resource", {"id": "RES-PR", "discipline": "process", "fte": 1}, R)      # 7.5 h / day for PID + LST
    r = workflow.compute(s)
    assert r.levelled and r.capacity == {"process": 7.5}
    pid, lst = r.docs["ALP-EPC-00PAC-PR-PID-0001"], r.docs["ALP-EPC-00PAC-PR-LST-0001"]
    for x in (pid, lst):
        assert abs(sum(x.work.values()) - x.rec["weight"]) < 1e-6
    days = set(pid.work) | set(lst.work)
    assert all(sum(x.work.get(i, 0) for x in (pid, lst)) <= 7.5 + 1e-6 for i in days)          # capacity never exceeded
    # PID has a need date (feeds the EWP and the requisition) and goes first; LST has none and is pushed back
    assert pid.late_start < lst.late_start
    assert lst.ifr > free.docs[lst.id].ifr and lst.driver == "capacity process"
    assert pid.ifr >= pid.start + 25


# ---------------------------------------------------------------- engines
def test_engineering_plan_engine(project):
    data(project, capacity={"process": 4, "piping": 4, "mechanical": 2})
    m = run_engine(project, Store(project, who()), registry()["engineering_plan"], {})
    assert sorted(m["files"]) == ["ALP_MDL.xlsx", "ALP_engineering_timeline.pdf"]
    wb = load_workbook(project.output / "engineering_plan" / "ALP_MDL.xlsx")
    assert {"MDL", "Workflow", "EWP", "PWP (MR)", "Checks"} <= set(wb.sheetnames)
    checks = [tuple(c.value for c in r[:4]) for r in wb["Checks"].iter_rows(min_row=2)]
    assert ("-", "numbering", "OK", "7 document numbers checked, 0 finding(s)") in checks
    assert any(c[1] == "AVEVA class" and c[2] == "OK" and "1 of 1 equipment" in c[3] for c in checks)


def test_engineering_plan_reports_awp_gaps(project):
    s = data(project)
    s.update("document", "ALP-EPC-00PAC-PI-ISO-0001", {"ewp": None}, R)
    s.update("equipment", "00PAC10AP001", {"aveva_class": None}, R)
    m = run_engine(project, Store(project, who()), registry()["engineering_plan"], {"pdf": "no"})
    w = " | ".join(m["warnings"])
    assert "construction document (ISO) without EWP" in w and "EWP without documents" in w
    assert "00PAC10AP001 AVEVA class" in w


def test_procedures_engine(project):
    s = data(project, capacity={"process": 4})
    for no, t in (("ALP-EPC-00000-GE-PRC-0001", "KKS identification manual"),
                  ("ALP-EPC-00000-GE-PRC-0002", "Document numbering"),
                  ("ALP-EPC-00000-GE-PLN-0001", "Engineering execution plan"),
                  ("ALP-EPC-00000-GE-PLN-0002", "AWP execution plan")):
        s.create("doc_type", {"id": no.split("-")[4], "title": t, "dcc": "B", "category": "procedure", "review": "approval",
                              "prep_days": 15, "update_days": 10, "input_maturity": "IFR"}, R) \
            if not s.get("doc_type", no.split("-")[4]) else None
        doc(s, no, "general", no.split("-")[4], weight=40)
    m = run_engine(project, Store(project, who()), registry()["procedures"], {"pdf": "no"})
    assert sorted(m["files"]) == ["ALP-EPC-00000-GE-PLN-0001.docx", "ALP-EPC-00000-GE-PLN-0002.docx",
                                  "ALP-EPC-00000-GE-PRC-0001.docx", "ALP-EPC-00000-GE-PRC-0002.docx"]
    from docx import Document
    text = "\n".join(p.text for p in Document(str(project.output / "procedures" / "ALP-EPC-00000-GE-PRC-0002.docx")).paragraphs)
    assert "ALP-<ORG>-<KKS>-<DISC>-<TYPE>-<NNNN>" in text and "10 working days" in text
    kd = Document(str(project.output / "procedures" / "ALP-EPC-00000-GE-PRC-0001.docx"))
    cells = [c.text for t in kd.tables for r in t.rows for c in r.cells]
    assert "00PAC10AP001" in cells and "PAC" in cells
    eep = Document(str(project.output / "procedures" / "ALP-EPC-00000-GE-PLN-0001.docx"))
    assert any("levelled" in p.text for p in eep.paragraphs)


def test_kks_cli(project, capsys):
    data(project)
    assert main(["kks", "00PAC10AP001"]) == 0
    assert "unit 0" in capsys.readouterr().out
    assert main(["kks", "--next", "EPC", "00PAC", "ME", "DSH"]) == 0
    assert "ALP-EPC-00PAC-ME-DSH-0002" in capsys.readouterr().out


# ---------------------------------------------------------------- HMB export limiter
@pytest.mark.parametrize("limiter,expected", [(None, "WARN"), ("GT load limiter per DEC-EPCE-0003", "INFO")])
def test_hmb_export_limiter(limiter, expected):
    case = {"id": "C1", "fuel": "natural_gas", **({"export_limiter": limiter} if limiter else {})}
    res = SimpleNamespace(checks=[], summary={"net": 610000.0}, case=case, nodes={}, aux=[])
    hmb._checks(res, {"tie_in:TP-E1": {"id": "TP-E1", "unit": "MW export", "value": 600}})
    st, txt = next(c for c in res.checks if "net export" in c[1])
    assert st == expected and ("export limited" in txt) == bool(limiter)
    res = SimpleNamespace(checks=[], summary={"net": 590000.0}, case=case, nodes={}, aux=[])
    hmb._checks(res, {"tie_in:TP-E1": {"id": "TP-E1", "unit": "MW export", "value": 600}})
    assert next(c for c in res.checks if "net export" in c[1])[0] == "OK"
