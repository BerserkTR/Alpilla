"""Heat and mass balance: node balances, auxiliary power and checks computed from the records; diagram and workbook
produced by the hmb engine. The test case is built so that its balances close exactly (known answers)."""
import re

import ezdxf
import pytest
from openpyxl import load_workbook

from engine.core import hmb, thermo as th
from engine.core.runner import run_engine
from engine.core.store import Store
from engine.engines import registry
from tests.conftest import who

R = "t"
FUEL = {"CH4": 95.0, "C2H6": 3.0, "N2": 1.5, "CO2": 0.5}


def closed_case():
    """GT node closing exactly: exhaust temperature solved from the energy balance; pump node P; boundaries."""
    case = {"id": "T1", "title": "Test case", "fuel": "natural_gas", "fuel_composition": "; ".join(f"{k}={v}" for k, v in FUEL.items()),
            "load_pct": 100, "ambient_temperature": 15.0, "relative_humidity": 60.0, "barometric_pressure": 1013.25,
            "seawater_temperature": 16.0, "gt_output": 300000.0, "gt_generator_efficiency": 0.985, "gt_other_losses": 2000.0,
            "st_output": 0.0, "st_mechanical_efficiency": 1.0, "st_generator_efficiency": 1.0, "status": "preliminary",
            "basis_refs": ["source:SRC-AAA-0001"]}
    f = th.fuel_properties(FUEL)
    air = th.humid_air(15.0, 60.0, 1.01325)
    m_a, m_f, T_f = 600.0, 16.0, 25.0
    flue = th.mole_fractions(th.combustion(air, m_a / th.mixture_M(air) * 1000, f["x"], m_f / f["M"] * 1000))
    E_in = m_a * th.h_gas(air, 15.0) + m_f * (f["LHV_mass"] * 1000 + th.h_fuel(f["x"], T_f))
    W = 300000.0 / 0.985 + 2000.0
    T_exh = th.T_from_h_gas(flue, (E_in - W) / (m_a + m_f))
    streams = [
        {"number": 1, "description": "Air", "fluid": "air", "from_node": "AMB", "to_node": "GT", "mass_flow": m_a, "pressure": 1.01325, "temperature": 15.0},
        {"number": 2, "description": "Fuel", "fluid": "fuel_gas", "from_node": "FGS", "to_node": "GT", "mass_flow": m_f, "pressure": 35.0, "temperature": T_f},
        {"number": 3, "description": "Exhaust", "fluid": "flue_gas", "from_node": "GT", "to_node": "STK", "mass_flow": m_a + m_f, "pressure": 1.05, "temperature": T_exh},
        {"number": 4, "description": "Feed in", "fluid": "water", "from_node": "TANK", "to_node": "P", "mass_flow": 10.0, "pressure": 2.0, "temperature": 50.0},
        {"number": 5, "description": "Feed out", "fluid": "water", "from_node": "P", "to_node": "BOIL", "mass_flow": 10.0, "pressure": 100.0, "temperature": 51.0},
    ]
    aux = [{"id": "AL-01", "description": "Fixed", "supplier": "LEAD", "method": "fixed", "power": 1500.0},
           {"id": "AL-02", "description": "Pump", "supplier": "LEAD", "method": "pump_node", "node": "P", "motor_efficiency": 0.95},
           {"id": "AL-03", "description": "GSU", "supplier": "MEMB", "method": "fixed", "power": 500.0, "counts_as": "transformer_loss"}]
    return case, streams, aux, f


def test_calculation_closes_and_computes_net():
    case, streams, aux, f = closed_case()
    refs = {"guarantee:PG-01": {"id": "PG-01", "contract": "K", "_kind": "owner_contract", "parameter": "Net output",
                                "guaranteed_value": 299.0, "unit": "MW", "direction": "min"},
            "design_parameter:DP-1": {"id": "DP-1", "parameter": "Reference LHV", "value": 49.5, "unit": "MJ/kg"}}
    res = hmb.calculate(case, streams, aux, refs)
    assert abs(res.nodes["GT"]["residual"]) < 1.0                       # kW: exact closure by construction
    pump = (th.h_pT(100.0, 51.0) - th.h_pT(2.0, 50.0)) * 10.0
    assert res.aux[1]["kW"] == pytest.approx(pump / 0.95)
    sm = res.summary
    assert sm["net"] == pytest.approx(300000.0 - 1500.0 - pump / 0.95 - 500.0)
    assert sm["heat_input"] == pytest.approx(16.0 * f["LHV_mass"] * 1000)
    assert sm["net_hr"] == pytest.approx(sm["heat_input"] * 3600 / sm["net"])
    status = {text.split()[0] + " " + text.split()[1]: st for st, text in res.checks}
    assert status["energy balance"] == "OK"
    warn = [t for s, t in res.checks if s == "WARN"]
    assert any("net output" in t and "PG-01" in t for t in warn)         # 297.x MW < 299 MW
    assert any("fuel LHV" in t for t in warn)                            # composition LHV differs from 49.5 MJ/kg


def test_energy_imbalance_is_reported():
    case, streams, aux, _ = closed_case()
    streams[2]["temperature"] += 10.0                                    # vendor data inconsistent by 10 K
    res = hmb.calculate(case, streams, aux, {})
    assert any(s == "WARN" and t.startswith("energy balance GT") for s, t in res.checks)
    streams[2]["mass_flow"] += 1.0
    res = hmb.calculate(case, streams, aux, {})
    assert any(s == "WARN" and t.startswith("mass balance GT") for s, t in res.checks)


def test_hmb_engine_outputs(project):
    s = Store(project, who())
    s.create("project", {"id": "ALP", "name": "Alpilla CCGT", "plant_type": "CCGT"}, R)
    for pid, role in (("OWNER", "owner"), ("LEAD", "consortium_leader"), ("MEMB", "consortium_member")):
        s.create("party", {"id": pid, "name": pid, "role": role}, R)
    s.create("source", {"id": "SRC-AAA-0001", "title": "Vendor heat balance", "originator": "MEMB"}, R)
    case, streams, aux, _ = closed_case()
    s.create("hmb_case", case, R)
    for st in streams:
        s.create("process_stream", dict(st, hmb_case="T1", basis_refs=["source:SRC-AAA-0001"]), R)
    for a in aux:
        s.create("aux_load", dict(a, hmb_case="T1", basis_refs=["source:SRC-AAA-0001"]), R)
    m = run_engine(project, Store(project, who()), registry()["hmb"], {})
    assert sorted(m["files"]) == ["ALP-HMB-T1.dxf", "ALP-HMB-T1.pdf", "ALP-HMB-T1.xlsx"]
    doc = ezdxf.readfile(project.output / "hmb" / "ALP-HMB-T1.dxf")
    assert not doc.audit().has_errors
    texts = [e.dxf.text for e in doc.modelspace().query("TEXT")]
    assert "NET OUTPUT (HV terminals)" in texts and "HEAT AND MASS BALANCE - T1" in texts
    pdf = (project.output / "hmb" / "ALP-HMB-T1.pdf").read_bytes()
    box = [float(v) for v in re.search(rb"/MediaBox \[\s*([\d.\s]+)\]", pdf).group(1).split()]
    assert abs(box[2] - 841 / 25.4 * 72) < 1 and abs(box[3] - 594 / 25.4 * 72) < 1          # A1 landscape
    wb = load_workbook(project.output / "hmb" / "ALP-HMB-T1.xlsx")
    assert wb.sheetnames == ["Streams", "Node balances", "Auxiliary loads", "Summary", "Checks"]
    assert [r[0] for r in wb["Streams"].iter_rows(min_row=4, values_only=True)] == [1, 2, 3, 4, 5]


def test_code_or_template_change_makes_output_stale(project):
    """Outputs record a fingerprint of the engine code and its declared code_deps (here the HMB layout template)."""
    from engine.core.runner import check_outputs
    test_hmb_engine_outputs(project)
    s = Store(project, who())
    assert not [m for lvl, m in check_outputs(project, s) if "STALE" in m]
    tpl = project.root / "templates" / "hmb" / "layout_1x1_3prh.json"
    tpl.write_text(tpl.read_text() + "\n")                                # any change to a declared dependency
    assert [m for lvl, m in check_outputs(project, s) if "output/hmb/: STALE" in m]


def test_ldo_case_with_water_injection_closes():
    """Liquid fuel with injected water: the water enters on the IAPWS basis and leaves as vapour in the flue gas."""
    fuel = th.liquid_fuel("C12H23", 42.9)
    air = th.humid_air(15.0, 70.0, 1.0115)
    m_a, m_f, m_w = 776.0, 22.4, 13.4
    flue = th.mole_fractions(th.combustion(air, m_a / th.mixture_M(air) * 1000, fuel["x"], m_f / fuel["M"] * 1000,
                                           m_w / th.M["H2O"] * 1000, fuel["atoms"]))
    E_in = m_a * th.h_gas(air, 15.0) + m_f * (42900 + th.h_liquid_fuel(fuel, 30.0)) + m_w * th.h_water_gas_basis(60.0, 30.0)
    W = 400500.0 / 0.989 + 3800.0
    T_exh = th.T_from_h_gas(flue, (E_in - W) / (m_a + m_f + m_w))
    case = {"id": "L1", "title": "LDO", "fuel": "ldo", "fuel_formula": "C12H23", "fuel_lhv": 42.9, "load_pct": 100,
            "ambient_temperature": 15.0, "relative_humidity": 70.0, "barometric_pressure": 1011.5, "seawater_temperature": 16.0,
            "gt_output": 400500.0, "gt_generator_efficiency": 0.989, "gt_other_losses": 3800.0, "st_output": 0.0,
            "st_mechanical_efficiency": 1.0, "st_generator_efficiency": 1.0}
    streams = [
        {"number": 1, "description": "Air", "fluid": "air", "from_node": "AMB", "to_node": "GT", "mass_flow": m_a, "pressure": 1.0115, "temperature": 15.0},
        {"number": 27, "description": "LDO", "fluid": "ldo", "from_node": "LDS", "to_node": "GT", "mass_flow": m_f, "pressure": 5.0, "temperature": 30.0},
        {"number": 28, "description": "Water", "fluid": "water", "from_node": "DMW", "to_node": "GT", "mass_flow": m_w, "pressure": 60.0, "temperature": 30.0},
        {"number": 4, "description": "Exhaust", "fluid": "flue_gas", "from_node": "GT", "to_node": "STK", "mass_flow": m_a + m_f + m_w,
         "pressure": 1.05, "temperature": T_exh}]
    res = hmb.calculate(case, streams, [], {})
    assert abs(res.nodes["GT"]["residual"]) < 1.0
    assert res.summary["heat_input"] == pytest.approx(m_f * 42900)
    assert res.summary["water_injection"] == pytest.approx(m_w) and res.summary["fuel_Sm3h"] is None
    assert 580 < T_exh < 600                            # injected water lowers the exhaust temperature
