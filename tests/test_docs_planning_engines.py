"""Document control + planning engines: MDR, schedule (xlsx/html/MSPDI), progress report, plan CLI."""
import xml.etree.ElementTree as ET

from engine.cli import main
from engine.core import validate
from engine.core.runner import run_engine
from engine.core.store import Store
from engine.engines import registry
from tests.conftest import who
from tests.test_planning import act, base

NS = {"p": "http://schemas.microsoft.com/project"}


def setup_data(project):
    s = base(project, data_date="2026-02-16", holidays=["2026-02-02"])
    for rid, pct in (("STARTED", 10), ("IFR", 50), ("IFA", 70), ("IFC", 100)):
        s.create("progress_rule", {"id": rid, "percent": pct}, "t")
    s.create("document", {"id": "ALP-EPC-10LAC-ME-DSH-0001", "title": "BFP datasheet", "discipline": "mechanical", "doc_type": "datasheet",
                          "weight": 30, "planned_ifr": "2026-01-19", "planned_ifc": "2026-02-09"}, "t")
    s.create("document", {"id": "ALP-EPC-00000-EL-LST-0001", "title": "Load list", "discipline": "electrical", "doc_type": "list",
                          "weight": 10, "planned_ifc": "2026-03-30"}, "t")
    s.create("document_revision", {"document": "ALP-EPC-10LAC-ME-DSH-0001", "revision": "A", "purpose": "IFR", "issue_date": "2026-01-21"}, "t")
    s.create("review_comment", {"revision": "ALP-EPC-10LAC-ME-DSH-0001_A", "originator": "Client", "comment": "Add NPSH margin"}, "t")
    act(s, "A-10", 5)
    act(s, "A-20", 10, ["A-10"], deliverables=["ALP-EPC-10LAC-ME-DSH-0001"])
    act(s, "A-30", 0, ["A-20:FS+3"], type="finish_milestone", constraint="finish_no_later_than", constraint_date="2026-02-06")
    return s


def run(project, name):
    return run_engine(project, Store(project, who()), registry()[name], {})


def test_schedule_outputs(project):
    setup_data(project)
    m = run(project, "schedule")
    assert sorted(m["files"]) == ["ALP_schedule.html", "ALP_schedule.xlsx", "ALP_schedule.xml"]
    assert any("negative float" in w and "A-30" in w for w in m["warnings"])
    root = ET.parse(project.output / "schedule" / "ALP_schedule.xml").getroot()
    tasks = root.findall("p:Tasks/p:Task", NS)
    names = [t.findtext("p:Name", namespaces=NS) for t in tasks]
    assert "A-20 A-20" in names and "Engineering" in names
    uids = [t.findtext("p:UID", namespaces=NS) for t in tasks]
    assert len(uids) == len(set(uids))
    by = {t.findtext("p:Name", namespaces=NS): t for t in tasks}
    link = by["A-30 A-30"].find("p:PredecessorLink", NS)
    assert link.findtext("p:Type", namespaces=NS) == "1" and link.findtext("p:LinkLag", namespaces=NS) == str(3 * 4800)
    assert link.findtext("p:PredecessorUID", namespaces=NS) == by["A-20 A-20"].findtext("p:UID", namespaces=NS)
    assert by["A-30 A-30"].findtext("p:ConstraintType", namespaces=NS) == "7"
    assert by["A-30 A-30"].findtext("p:Milestone", namespaces=NS) == "1"
    holidays = [w for w in root.findall("p:Calendars/p:Calendar/p:WeekDays/p:WeekDay", NS)
                if w.findtext("p:DayType", namespaces=NS) == "0"]
    assert holidays[0].findtext("p:TimePeriod/p:FromDate", namespaces=NS) == "2026-02-02T00:00:00"
    html = (project.output / "schedule" / "ALP_schedule.html").read_text()
    assert "<svg" in html and "CRIT" in html and "data-tip" in html


def test_register_and_progress(project):
    from openpyxl import load_workbook
    setup_data(project)
    m = run(project, "document_register")
    assert any("1 document(s) late" in w for w in m["warnings"])          # DS-001 planned IFC 09-Feb passed
    ws = load_workbook(project.output / "document_register" / "ALP_document_register.xlsx")["MDR"]
    rows = {r[0]: r for r in ws.iter_rows(min_row=5, values_only=True)}
    assert rows["ALP-EPC-10LAC-ME-DSH-0001"][9:11] == ("A", "IFR") and rows["ALP-EPC-10LAC-ME-DSH-0001"][18] == "LATE"
    assert rows["ALP-EPC-10LAC-ME-DSH-0001"][19] == 1 and rows["ALP-EPC-10LAC-ME-DSH-0001"][20] == 50
    run(project, "progress_report")
    html = (project.output / "progress_report" / "ALP_progress_report.html").read_text()
    assert "37.5%" in html        # earned: DS-001 IFR 50% x 30/40
    assert "Table view" in html and "ALP-EPC-10LAC-ME-DSH-0001" in html


def test_validate_catches_loops_and_plan_cli(project, capsys):
    s = setup_data(project)
    assert main(["plan"]) == 0
    out = capsys.readouterr().out
    assert "critical" in out and "A-20" in out
    s.update("activity", "A-10", {"predecessors": ["A-30"]}, "loop")
    assert any("logic loop" in e for e in validate.run(project).errors)
    assert main(["plan"]) == 1
