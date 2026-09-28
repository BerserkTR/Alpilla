"""Document release control (engine/core/release.py, engine doc / gate CLI, release_control engine) and the gate wiring of
the workflow (PO prerequisites, gate need dates)."""
from engine.cli import main
from engine.core import release, workflow
from engine.core.runner import run_engine
from engine.core.store import Store
from engine.engines import registry
from tests.conftest import who
from tests.test_engineering import R, data

PID, DBR, IEC = "ALP-EPC-00PAC-PR-PID-0001", "ALP-EPC-00000-PR-DBR-0001", "ALP-IEC-00000-GE-LST-0001"
DSH, MRQ, LST = "ALP-EPC-00PAC-ME-DSH-0001", "ALP-EPC-00PAC-ME-MRQ-0001", "ALP-EPC-00PAC-PR-LST-0001"


class FakeDocs:
    """Stand-in producing engine: writes <document number>.txt for every document produced_by 'fake_docs'."""
    name, title, version, inputs, formats, code_deps = "fake_docs", "test documents", "1.0.0", ["document", "document_revision"], ["txt"], []

    def run(self, ctx):
        from engine.core import docshell
        out = []
        for d in ctx.store.records("document"):
            if d.get("produced_by") == self.name:
                f = ctx.out_dir / f"{d['id']}.txt"
                f.write_text(docshell.state(ctx.store, d["id"])["status"] + "\n", encoding="utf-8")
                out.append(f)
        return out


def producing(project, monkeypatch):
    """All fixture documents produced by the fake engine."""
    import engine.engines as eng
    real = eng.registry
    monkeypatch.setattr(eng, "registry", lambda: {**real(), "fake_docs": FakeDocs()})
    s = Store(project, who())
    for d in s.records("document"):
        s.update("document", d["id"], {"produced_by": "fake_docs"}, R)


def issue(doc, purpose, date, *extra):
    return main(["doc", "issue", doc, "--purpose", purpose, "--date", date, "--reason", "t", *extra])


def review(doc, rev, code, date):
    return main(["doc", "review", doc, "--rev", rev, "--code", code, "--date", date, "--reason", "t"])


def gate(s, gid, scope, requires, need, **kw):
    s.create("gate_rule", {"id": gid, "process": kw.pop("process", "procurement"), "scope": scope, "title": gid,
                           "requires": requires, "need": need, "status": "agreed", **kw}, R)


def test_release_rules_revisions_and_check_required(project, capsys, monkeypatch):
    data(project)
    producing(project, monkeypatch)
    assert issue(PID, "IFR", "2026-11-10") == 1                         # inputs not issued: blocked
    assert "BLOCKED" in capsys.readouterr().out
    assert issue(DBR, "IFR", "2026-11-03") == 0 and issue(IEC, "IFR", "2026-11-04") == 0
    assert issue(PID, "IFR", "2026-11-10") == 0
    s = Store(project, who())
    r = s.get("document_revision", f"{PID}_A")
    assert r["purpose"] == "IFR" and sorted(r["based_on"]) == [f"{DBR}_A", f"{IEC}_A"]
    st = release.states(s)
    assert release.NAME[st[PID].level] == "IFR" and release.next_purpose(s.get("document", PID), st[PID],
                                                                           {t["id"]: t for t in s.records("doc_type")}) is None
    why = release.can_issue(s, PID, "IFC", st)                          # inputs not IFC / accepted, PID not approved
    assert any(DBR in w and "IFC" in w for w in why) and any(IEC in w and "ACCEPTED" in w for w in why)
    assert any("not yet accepted by the Owner" in w for w in why)
    # DBR approved and issued IFC: revision numbers start at 0
    assert review(DBR, "A", "1", "2026-11-20") == 0 and issue(DBR, "IFC", "2026-11-25") == 0
    s = Store(project, who())
    assert s.get("document_revision", f"{DBR}_0")["purpose"] == "IFC"
    # a later DBR revision makes the PID (prepared from DBR rev A) CHECK REQUIRED, which blocks its users
    st = release.states(s)
    assert (DBR, f"{DBR}_0") in st[PID].suspect
    assert any("CHECK REQUIRED" in w for w in release.can_issue(s, DSH, "IFR", st))
    # override: released with the justification recorded
    assert issue(DSH, "IFR", "2026-11-26", "--override", "vendor enquiry on preliminary data") == 0
    s = Store(project, who())
    assert "vendor enquiry" in s.get("document_revision", f"{DSH}_A")["override"]


def test_next_revision_labels():
    st = release.DocState("X")
    assert release.next_rev(st, "IFR") == "A"
    st.revs = [{"revision": "A"}, {"revision": "B"}]
    st.level = release.LEVEL["ACCEPTED"]
    assert release.next_rev(st, "IFA") == "C" and release.next_rev(st, "IFC") == "0"
    st.revs.append({"revision": "0"})
    st.level = release.LEVEL["IFC"]
    assert release.next_rev(st, "IFC") == "1" and release.next_rev(st, "AB") == "1"


def test_po_prerequisites_from_the_gate_set_the_po(project):
    s = data(project)
    s.update("mr", "MR-ME-003", {"bid_days": 0, "award_days": 0}, R)
    r = workflow.compute(s)
    assert r.mrs["MR-ME-003"]["po"] == r.docs[MRQ].ifc                  # no gate: requisition + bid + award
    gate(s, "G-PRC-PO", "mr", ["type:DSH@epc IFC", "type:MRQ@epc IFC", "type:ITP@epc? IFR"], "po")
    r = workflow.compute(s)
    assert r.docs[DSH].ifc > r.docs[MRQ].ifc
    assert r.mrs["MR-ME-003"]["po"] == r.docs[DSH].ifc                  # the datasheet must be IFC before the PO
    g = [x for x in release.gates(s, r) if x.gate == "G-PRC-PO"][0]
    assert g.scope_key == "MR-ME-003" and not g.empty and g.ready == r.docs[DSH].ifc
    assert g.need == r.mrs["MR-ME-003"]["po_late"] and not g.ok_now


def test_gate_need_prioritises_required_documents(project):
    s = data(project)
    r = workflow.compute(s)
    assert r.docs[LST].late_start == workflow.NO_NEED                   # nothing needs the process list
    gate(s, "G-COM-X", "plant", [f"id:{LST} ACCEPTED"], "activity:CWP-05-PI-01", process="commissioning")
    r = workflow.compute(s)
    act = r.acts["CWP-05-PI-01"]
    assert r.docs[LST].late_start < workflow.NO_NEED
    t = s.get("doc_type", "LST")
    accepted_by = r.docs[LST].late_start + workflow._prep(r.docs[LST], t) + workflow.level_offset(r.docs[LST], t, "ACCEPTED")
    assert accepted_by == act.ef
    assert ("ALP-EPC-00PAC-PR-LST-0001", "ACCEPTED", act.ef, "G-COM-X") in r.gate_needs


def test_gate_requirement_syntax(project):
    s = data(project)
    gate(s, "G-TST-A", "plant", ["type:XYZ IFR"], "activity:CWP-05-PI-01")          # mandatory, nothing selected
    gate(s, "G-TST-B", "plant", ["type:XYZ? IFR"], "activity:CWP-05-PI-01")         # optional only: gate not applicable
    gate(s, "G-TST-C", "cwp", ["any:ewp-all IFC", "type:DBR|type:LST@plant? IFR"], "cwp_start", process="construction")
    gate(s, "G-TST-D", "mr", ["type:DSH@supplier? ACCEPTED", "type:DSH@mr,epc ACCEPTED"], "fabrication")
    r = workflow.compute(s)
    gs = {x.gate: x for x in release.gates(s, r)}
    assert gs["G-TST-A"].empty == ["type:XYZ IFR"] and not gs["G-TST-A"].ok_now
    assert "G-TST-B" not in gs
    c = gs["G-TST-C"]
    # bindings apply per alternative: type:DBR is looked for in the EWP of the CWP (none), type:LST@plant in the plant
    assert {i for i, _ in c.members} == {"ALP-EPC-00PAC-PI-ISO-0001", IEC}
    assert gs["G-TST-D"].members == [(DSH, "ACCEPTED")]


def test_release_control_engine_and_gate_cli(project, capsys):
    s = data(project)
    gate(s, "G-OWN-BDP", "plant", ["bdp ACCEPTED"], "activity:CWP-05-PI-01", offset_days=-200, process="owner_acceptance")
    assert main(["gate", "--detail"]) == 0
    out = capsys.readouterr().out
    assert "G-OWN-BDP" in out and "bdp ACCEPTED" in out and "1 planned late" in out
    m = run_engine(project, Store(project, who()), registry()["release_control"])
    assert any("gate G-OWN-BDP 00000: planned" in w for w in m["warnings"])
    from openpyxl import load_workbook
    wb = load_workbook(next((project.output / "release_control").glob("*.xlsx")))
    assert wb.sheetnames == ["Rules", "Gates", "Gate requirements", "Release plan", "Relations", "Impacts", "Gate rules"]
    plan = [r for r in wb["Release plan"].iter_rows(min_row=5, values_only=True)]
    order = [r[1] for r in plan]
    assert order.index(DBR) < order.index(PID) < order.index("ALP-EPC-00PAC-PI-ISO-0001")
    row = plan[order.index(DBR)]
    assert row[12] == "IFR" and row[13] == "A" and row[14] == "yes"          # next purpose, revision, can issue now
    imp = {r[0]: r for r in wb["Impacts"].iter_rows(min_row=5, values_only=True)}
    assert imp[DBR][3] >= 4 and "G-OWN-BDP" in imp[DBR][6]


def test_after_activity_and_as_late_as_possible(project):
    s = data(project)
    s.update("document", LST, {"after_activity": "CWP-05-PI-01"}, R)
    r = workflow.compute(s)
    assert r.docs[LST].start == r.acts["CWP-05-PI-01"].ef and r.docs[LST].driver == "after CWP-05-PI-01"
    s.update("document", LST, {"after_activity": None, "start_after": "alap"}, R)
    gate(s, "G-COM-X", "plant", [f"id:{LST} IFR"], "activity:CWP-05-PI-01", offset_days=300, process="commissioning")
    r = workflow.compute(s)
    d = r.docs[LST]
    assert d.late_start < workflow.NO_NEED and d.start == d.late_start - workflow.ALAP_MARGIN
    assert d.driver == "as late as possible (need)"


def test_titles_subject_type_and_description():
    from engine.core import titles
    assert titles.make("Technical specification - MV switchgear 10.5 kV", {"technical specification"}) == \
        ("MV Switchgear 10.5 kV - Technical Specification", None)
    assert titles.make("Process design criteria (design margins, API 520/521)") == \
        ("Process Design Criteria", "design margins, API 520/521")
    assert titles.make("GT generator gas system (H2/CO2) - System turnover dossier") == \
        ("GT Generator Gas System - System Turnover Dossier", "H2/CO2")
    assert titles.proper("stack 65 m with CEMS platform and silencer") == "Stack 65 m with CEMS Platform and Silencer"


def test_issue_regenerates_and_freezes_the_document(project, capsys, monkeypatch):
    data(project)
    assert issue(DBR, "IFR", "2026-11-03") == 1                         # no producing engine: blocked (R-001)
    assert "no engine produces this document" in capsys.readouterr().out
    producing(project, monkeypatch)
    assert issue(DBR, "IFR", "2026-11-03", "--checked", "ABC", "--approved", "XYZ") == 0
    out = capsys.readouterr().out
    assert "frozen in internal_deliveries/" in out
    s = Store(project, who())
    rev = s.get("document_revision", f"{DBR}_A")
    assert rev["prepared_by"] == "AAA" and rev["checked_by"] == "ABC"
    dl = [x for x in s.records("delivery") if f"{DBR}_A" in (x.get("revisions") or [])]
    assert len(dl) == 1 and dl[0]["recipients"] == "Owner (approval)"
    frozen = next((project.root / dl[0]["folder"]).glob(f"{DBR}.txt"))
    assert frozen.read_text().startswith("Rev A - IFR, issued 2026-11-03")        # cover of the issued revision
