"""Engines are the only output path: manifests, tamper/stale detection, deliveries, governance files."""
import json
import shutil

import pytest

from engine.cli import main
from engine.core import validate
from engine.core.delivery import deliver
from engine.core.runner import run_engine
from engine.core.store import Store
from engine.engines import registry
from tests.conftest import seed, who

E = registry()


def run(project, name, **opts):
    return run_engine(project, Store(project, who()), E[name], {k: str(v) for k, v in opts.items()})


def test_equipment_list_xlsx(project):
    from openpyxl import load_workbook
    seed(Store(project, who()))
    m = run(project, "equipment_list")
    f = project.output / "equipment_list" / "ALP_equipment_list.xlsx"
    assert list(m["files"]) == ["ALP_equipment_list.xlsx"] and f.exists()
    ws = load_workbook(f)["Equipment List"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Rated Power [kW]" in values and "10LAC10AP001" in values and 3200 in values
    assert any(str(v).startswith("10LAC - Feedwater pumps") for v in values)
    assert validate.run(project, Store(project, who())).ok


def test_design_basis_md_html(project):
    s = Store(project, who())
    seed(s)
    s.create("design_parameter", {"category": "site", "parameter": "Soil class | seismic zone", "value_text": "B | 2",
                                  "basis_refs": ["source:SRC-AAA-0001"]}, "r")
    run(project, "design_basis")
    md = (project.output / "design_basis" / "ALP_design_basis.md").read_text()
    html = (project.output / "design_basis" / "ALP_design_basis.html").read_text()
    assert "| DP-AAA-0001 | Maximum dry bulb temperature | 45 | degC |" in md
    assert "SRC-AAA-0001 Client site data" in md and "IEC 60034-1" in md
    assert "Open assumptions (2 of 3)" in md
    assert "<table>" in html and "Grid frequency" in html
    assert "<td>Soil class | seismic zone</td>" in html and "<td>B | 2</td>" in html   # pipes do not break tables


def test_datasheets_docx_pdf(project):
    from docx import Document
    seed(Store(project, who()))
    m = run(project, "equipment_datasheets", ids="10LAC10AP001")
    docx = [f for f in m["files"] if f.endswith(".docx")]
    assert docx == ["ALP-DS-10LAC10AP001_datasheet.docx"]
    text = "\n".join(p.text for p in Document(project.output / "equipment_datasheets" / docx[0]).paragraphs)
    assert "Equipment Datasheet - 10LAC10AP001" in text and "Client site data" in text
    if shutil.which("soffice"):
        assert "ALP-DS-10LAC10AP001_datasheet.pdf" in m["files"]
    with pytest.raises(ValueError, match="unknown"):
        run(project, "equipment_datasheets", ids="NOPE")


def test_hand_made_and_edited_outputs_detected(project):
    seed(Store(project, who()))
    run(project, "design_basis")
    d = project.output / "design_basis"
    (d / "ALP_design_basis.md").write_text("tampered")
    (d / "extra.md").write_text("hand made")
    (project.output / "loose.xlsx").write_text("x")
    (project.output / "manual").mkdir()
    errs = validate.run(project, Store(project, who())).errors
    assert any("hand-edited" in e for e in errs)
    assert any("extra.md" in e and "hand-made" in e for e in errs)
    assert any("loose.xlsx" in e for e in errs)
    assert any("manual/: no manifest" in e for e in errs)


def test_stale_output_reported(project):
    s = Store(project, who())
    seed(s)
    run(project, "equipment_list")
    run(project, "design_basis")
    s.update("equipment", "10MBV10AP001", {"rated_power": 90}, "vendor update")
    warns = validate.run(project, Store(project, who())).warnings
    assert any("equipment_list/: STALE" in w for w in warns)
    assert not any("design_basis/: STALE" in w for w in warns)   # its inputs did not change


def test_engine_refuses_invalid_database(project):
    seed(Store(project, who()))
    f = project.records_dir / "system" / "10MBA.json"
    f.write_text(f.read_text().replace("Gas turbine unit 1", "edited"))
    with pytest.raises(RuntimeError, match="not valid"):
        run(project, "equipment_list")


def test_delivery_frozen_and_registered(project):
    seed(Store(project, who()))
    run(project, "equipment_list")
    s = Store(project, who())
    dest = deliver(project, s, "equipment_list", "Equipment list for review", "internal review", "Mech lead", "first issue")
    assert (dest / "ALP_equipment_list.xlsx").exists() and (dest / "TRANSMITTAL.md").exists()
    t = json.loads((dest / "transmittal.json").read_text())
    assert t["id"] == "DLV-AAA-0001" and t["engine"] == "equipment_list"
    assert Store(project, who()).get("delivery", "DLV-AAA-0001")["file_count"] == 1
    assert validate.run(project, Store(project, who())).ok
    (dest / "ALP_equipment_list.xlsx").write_bytes(b"changed")
    (dest / "notes.txt").write_text("x")
    errs = validate.run(project, Store(project, who())).errors
    assert any("changed after issue" in e for e in errs) and any("notes.txt" in e for e in errs)


def test_delivery_refuses_stale_output(project):
    s = Store(project, who())
    seed(s)
    run(project, "equipment_list")
    s.update("equipment", "10MBV10AP001", {"rated_power": 90}, "vendor update")
    with pytest.raises(RuntimeError, match="STALE"):
        deliver(project, Store(project, who()), "equipment_list", "x", "y", "z", "r")


def test_governance_files_generated_and_protected(project):
    assert main(["db", "add", "rule", "--id", "R-001", "--set", "title=Engines only",
                 "--set", "statement=Outputs come from engines.", "--reason", "t"]) == 0
    assert main(["db", "add", "lesson", "--set", "title=Check units", "--set", "lesson=kW not MW",
                 "--set", "date=2026-09-27", "--set", "tags=units,data", "--reason", "t"]) == 0
    rules, ledger = project.rules_md.read_text(), project.ledger_md.read_text()
    assert "## R-001 [MANDATORY] Engines only" in rules
    assert "### L-AAA-0001 (2026-09-27) Check units [units, data]" in ledger
    assert validate.run(project).ok
    project.ledger_md.write_text(ledger + "\nhand-written lesson\n")
    assert any("ledger.md is out of sync" in e for e in validate.run(project).errors)
    assert main(["run", "governance"]) == 0
    assert validate.run(project).ok


def test_cli_roundtrip(project, capsys, tmp_path):
    seed(Store(project, who()))
    csv = tmp_path / "eq.csv"
    csv.write_text("id,description,system,equipment_type,rated_power,basis_refs\n"
                   "10MBV10AP002,GT lube oil pump B,10MBA,pump,75,source:SRC-AAA-0001\n")
    assert main(["db", "import", "equipment", str(csv), "--reason", "OEM list"]) == 0
    assert main(["db", "update", "equipment", "10MBV10AP002", "--set", "rated_power=80", "--unset", "basis_refs",
                 "--reason", "vendor"]) == 0
    capsys.readouterr()
    assert main(["db", "list", "equipment", "--where", "rated_power < 100", "--fields", "rated_power"]) == 0
    out = capsys.readouterr().out
    assert "10MBV10AP002  80" in out and "10LAC10AP001" not in out
    assert main(["run", "--all"]) == 0
    assert main(["validate"]) == 0
    assert main(["db", "add", "system", "--id", "X1", "--set", "title=x", "--set", "category=bad", "--reason", "r"]) == 1


def test_cli_read_commands(project, capsys):
    seed(Store(project, who()))
    assert main(["db", "query", "SELECT id, rated_power FROM equipment ORDER BY id", "--limit", "1"]) == 0
    out = capsys.readouterr().out
    assert "10LAC10AP001  3200" in out and "more rows" in out
    assert main(["db", "get", "system", "10MBA"]) == 0
    assert main(["db", "schema", "equipment"]) == 0
    assert "rated_power" in capsys.readouterr().out
    assert main(["status"]) == 0
    assert main(["engines"]) == 0


def test_run_stale_only_reruns_outdated(project, capsys):
    s = Store(project, who())
    seed(s)
    assert main(["run", "--all"]) == 0
    capsys.readouterr()
    assert main(["run", "--stale"]) == 0 and "all outputs are current" in capsys.readouterr().out
    s.update("design_parameter", "DP-AAA-0001", {"value": 46}, "site data rev 2")
    assert main(["run", "--stale"]) == 0
    out = capsys.readouterr().out
    assert "design_basis:" in out and "equipment_list" not in out
    assert validate.run(project).warnings == []
