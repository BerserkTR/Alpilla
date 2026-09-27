"""AVEVA class library: tags are classed and their attributes validated against the library."""
import json

import pytest

from engine.cli import main
from engine.core import classlib, validate
from engine.core.store import Store, StoreError
from tests.conftest import seed, who

PUMP = "Centrifugal Pump"


def test_library_compiled_and_consistent(project):
    lib = classlib.get(project)
    m = json.loads((project.database / "classlib" / "manifest.json").read_text())
    assert m["counts"] == {"classes": 1533, "attributes": 10218, "enums": 850, "quantities": 154, "associations": 1146}
    cls = lib.resolve(PUMP)
    assert [c["label"] for c in lib.ancestors(cls)][:4] == [PUMP, "Dynamic Pump", "Pump", "Rotating Equipment"]
    assert lib.effective(cls)["Rated Power"]["quantity"] == "Power"
    assert "kW" in lib.quantities["Power"]["units"]
    assert lib.quantities["Bore"]["units_from"] == "Length"          # shared unit family resolved
    with pytest.raises(KeyError, match="ambiguous"):
        lib.resolve("Pressure Relief Device")
    with pytest.raises(KeyError, match="Did you mean"):
        lib.resolve("Centrifugal Pmp")


def test_classified_equipment(project):
    s = Store(project, who())
    seed(s)
    s.update("equipment", "10LAC10AP001", {"aveva_class": PUMP, "aveva_attrs": {
        "Rated Power": {"value": 3200, "unit": "kW"}, "Casing Material": "A216 WCB"}}, "classify")
    assert validate.run(project, Store(project, who())).ok
    bad = [({"aveva_class": "Centrifugal Pmp"}, "Did you mean"),
           ({"aveva_class": "Power Transformer"}, None),                        # valid class, under Item
           ({"aveva_class": "Document"}, "not under"),                           # wrong branch
           ({"aveva_class": PUMP, "aveva_attrs": {"Rated Powr": 1}}, "did you mean Rated Power"),
           ({"aveva_class": PUMP, "aveva_attrs": {"Rated Power": {"value": 1, "unit": "barg"}}}, "not valid for Power"),
           ({"aveva_class": PUMP, "aveva_attrs": {"Rated Power": 3200}}, "needs"),
           ({"aveva_attrs": {"Rated Power": {"value": 1, "unit": "kW"}}, "aveva_class": None}, "set aveva_class first")]
    for change, msg in bad:
        if msg is None:
            s.update("equipment", "10MBV10AP001", change, "ok class"); continue
        with pytest.raises(StoreError, match=msg):
            s.update("equipment", "10LAC10AP001", change, "bad")


def test_closed_value_list_enforced(project):
    lib = classlib.get(project)
    cls = lib.resolve(PUMP)
    closed = next(a for k, a in lib.effective(cls).items() if not k.startswith("AVEVA-") and a.get("lov")
                  and lib.enums[a["lov"]]["closed"])
    ok = lib.enums[closed["lov"]]["values"][0]
    assert lib.attr_errors(cls, {closed["label"]: ok}) == []
    assert "not in" in lib.attr_errors(cls, {closed["label"]: "definitely-not-a-listed-value"})[0]


def test_cli_attrs_merge_and_unset(project, capsys):
    seed(Store(project, who()))
    assert main(["db", "update", "equipment", "10LAC10AP001", "--set", f"aveva_class={PUMP}",
                 "--attr", "Rated Power=3200 kW", "--attr", "Casing Material=A216 WCB", "--reason", "classify"]) == 0
    assert main(["db", "update", "equipment", "10LAC10AP001", "--attr", "Shaft Power=2950.5 kW",
                 "--unset-attr", "Casing Material", "--reason", "vendor data"]) == 0
    r = Store(project, who()).get("equipment", "10LAC10AP001")
    assert r["aveva_attrs"] == {"Rated Power": {"value": 3200, "unit": "kW"}, "Shaft Power": {"value": 2950.5, "unit": "kW"}}
    capsys.readouterr()
    assert main(["db", "query", "SELECT json_extract(aveva_attrs, '$.\"Rated Power\".value') FROM equipment "
                 "WHERE aveva_class='Centrifugal Pump'"]) == 0
    assert "3200" in capsys.readouterr().out
    assert main(["db", "update", "equipment", "10MBV10AP001", "--attr", "Rated Power=75 kW", "--reason", "x"]) == 1


def test_lib_cli(project, capsys):
    assert main(["lib", "find", "transformer", "--root", "Electrical Component"]) == 0
    assert "Power Transformer" in capsys.readouterr().out
    assert main(["lib", "show", PUMP, "--grep", "power"]) == 0
    assert "Rated Power" in capsys.readouterr().out
    assert main(["lib", "attr", PUMP, "Rated Power"]) == 0
    assert "base kW" in capsys.readouterr().out
    assert main(["lib", "tree", "Document", "--depth", "2"]) == 0


def test_classlib_tamper_and_source_change_detected(project):
    f = project.database / "classlib" / "enums.jsonl"
    f.write_text(f.read_text().replace('"Aluminium"', '"Alu"', 1))
    assert any("enums.jsonl: missing or edited by hand" in e for e in validate.run(project).errors)
    src = next((project.references / "aveva").glob("*.ttl"))
    src.unlink()                                       # drop the link, then put a different file in its place
    src.write_text("# changed")
    assert any("changed since the class library was compiled" in e for e in validate.run(project).errors)


def test_aveva_export_engine(project):
    from openpyxl import load_workbook
    from engine.core.runner import run_engine
    from engine.engines import registry
    s = Store(project, who())
    seed(s)
    s.update("equipment", "10LAC10AP001", {"aveva_class": PUMP, "aveva_attrs": {
        "Rated Power": {"value": 3200, "unit": "kW"}, "Casing Material": "A216 WCB"}}, "classify")
    s.update("equipment", "10MBV10AP001", {"aveva_class": PUMP, "aveva_attrs": {
        "Rated Power": {"value": 100, "unit": "hp"}}}, "classify")
    m = run_engine(project, Store(project, who()), registry()["aveva_export"], {})
    wb = load_workbook(project.output / "aveva_export" / "ALP_aveva_tag_export.xlsx")
    ws = wb[PUMP]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0][:4] == ["Name", "Entity", "Class", "Description"]
    assert rows[1][2] == "AVEVA-1.0.1-CLA-000201"
    # mixed units -> one column per unit, never converted
    i_kw = [j for j, (n, u) in enumerate(zip(rows[0], rows[2])) if n == "Rated Power" and u == "kW"][0]
    i_hp = [j for j, (n, u) in enumerate(zip(rows[0], rows[2])) if n == "Rated Power" and u == "hp"][0]
    data = {r[0]: r for r in rows[3:]}
    assert data["10LAC10AP001"][i_kw] == 3200 and data["10LAC10AP001"][i_hp] is None
    assert data["10MBV10AP001"][i_hp] == 100
    assert "Unclassified" in wb.sheetnames and any("without aveva_class" in w for w in m["warnings"])
