"""MDL rules (engine/core/mdl.py, engine mdl sync), supplier document timing, progressive EWP release, mobilisation ramp
and the completeness benchmark check."""
from engine.cli import main
from engine.core import mdl, workflow
from engine.core.runner import run_engine
from engine.core.store import Store
from engine.engines import registry
from tests.conftest import who
from tests.test_engineering import R, data, dtype


def rule(s, rid, typ, disc, scope, title, **kw):
    s.create("mdl_rule", {"id": rid, "type_code": typ, "discipline": disc, "scope": scope, "title": title,
                          "originator": kw.pop("org", "EPC"), "status": "agreed", **kw}, R)


def setup(project):
    s = data(project)
    s.update("system", "00PAC", {"cwas": ["CWA-05"], "mr": "MR-ME-003", "q_pid_sheets": 3, "q_instruments": 250,
                                 "q_equipment": 0, "q_iso_sheets": 0}, R)
    s.update("mr", "MR-ME-003", {"supplier_code": "M03", "system": "00PAC"}, R)
    dtype(s, "GAD", "review", prep=10)
    dtype(s, "LAY", "approval", prep=20)
    s.update("doc_type", "DSH", {"category": "datasheet"}, R)
    return s


def test_rule_expansion_parts_register_and_zero_quantity(project):
    s = setup(project)
    rule(s, "DL-IC-IDS", "DSH", "i_and_c", "system", "Instrument datasheets - {name}", quantity="q_instruments",
         max_sheets=100, hours_base=8, hours_per_sheet=1.5)
    rule(s, "DL-PI-ISO", "ISO", "piping", "system", "Isometrics - {name}", quantity="q_iso_sheets", hours_per_sheet=10)
    rule(s, "DL-PI-LAY", "LAY", "piping", "cwa", "Layout - {name}", quantity="q_equipment", per_sheet=12, min_sheets=2)
    inst, errs = mdl.required(s)
    assert not errs
    ids = sorted(k for k in inst if k.startswith("DL-IC-IDS"))
    assert ids == ["DL-IC-IDS|00PAC|1", "DL-IC-IDS|00PAC|2", "DL-IC-IDS|00PAC|3"]          # 250 sheets / max 100
    assert sorted(inst[k].sheets for k in ids) == [83, 83, 84] and inst[ids[0]].hours == 8 + 1.5 * 84
    assert "part 1 of 3" in inst[ids[0]].title
    assert not any(k.startswith("DL-PI-ISO") for k in inst)                                  # quantity 0: not required
    lay = inst["DL-PI-LAY|CWA-05|1"]
    assert lay.sheets == 2                                                                  # min_sheets applies at q = 0 ...
    # ... and the register count replaces a smaller estimate (1 equipment record vs estimate 0)
    assert mdl.Ctx(s).q("00PAC", "q_equipment") == 1


def test_eqgroup_and_supplier_rules(project):
    s = setup(project)
    rule(s, "DL-ME-DSH", "DSH", "mechanical", "eqgroup", "Datasheet - {name}", supply=["IEPC"], equipment_types=["pump"],
         hours_base=6, hours_per_sheet=10)
    rule(s, "DL-V-GAD", "GAD", "mechanical", "mr", "Supplier GA - {name}", org="SUPPLIER", package_types=["equipment"],
         quantity="q_pid_sheets", po_weeks_ifr=6, po_weeks_final=12)
    s.update("equipment", "00PAC10AP001", {"supply": "IEPC"}, R)
    inst, errs = mdl.required(s)
    assert not errs
    g = inst["DL-ME-DSH|00PAC:MR-ME-003:pump|1"]
    assert g.equipment == ["00PAC10AP001"] and g.mr == "MR-ME-003" and g.sheets == 1
    v = inst["DL-V-GAD|MR-ME-003|1"]
    assert v.org == "M03" and v.kks == "00PAC" and v.sheets == 3
    s.update("mr", "MR-ME-003", {"supplier_code": None}, R)
    _, errs = mdl.required(s)
    assert any("no supplier_code" in e for e in errs)


def test_sync_create_update_orphan_and_idempotent(project, capsys):
    s = setup(project)
    rule(s, "DL-PR-SDS2", "PID", "process", "system", "P&ID - {name}", quantity="q_pid_sheets", hours_per_sheet=100,
         inputs=["doc:ALP-EPC-00000-PR-DBR-0001"])
    rule(s, "DL-IC-IDS", "DSH", "i_and_c", "system", "Instrument datasheets - {name}", quantity="q_instruments",
         hours_per_sheet=1.5, inputs=["DL-PR-SDS2"])
    # adopt the existing P&ID: sheets and hours follow the rule
    s.update("document", "ALP-EPC-00PAC-PR-PID-0001", {"rule": "DL-PR-SDS2|00PAC|1"}, R)
    assert main(["mdl", "sync"]) == 0
    out = capsys.readouterr().out
    assert "create 1 document(s)" in out and "update 1" in out and "dry run" in out
    assert main(["mdl", "sync", "--apply", "--reason", "t"]) == 0
    s = Store(project, who())
    pid = s.get("document", "ALP-EPC-00PAC-PR-PID-0001")
    assert pid["sheets"] == 3 and pid["weight"] == 300 and "ALP-EPC-00000-PR-DBR-0001" in pid["inputs"]
    ids = s.get("document", "ALP-EPC-00PAC-IC-DSH-0001")
    assert ids["rule"] == "DL-IC-IDS|00PAC|1" and ids["sheets"] == 250 and ids["inputs"] == ["ALP-EPC-00PAC-PR-PID-0001"]
    assert ids["aveva_class"] == "Specification Documents"
    capsys.readouterr()
    assert main(["mdl", "sync"]) == 0
    assert "create 0 document(s), 0 sheets, 0 h; update 0; orphans 0" in capsys.readouterr().out     # idempotent
    # rule removed -> orphan; cancelled with --cancel-orphans and removed from the inputs of others
    s.update("document", "ALP-EPC-00PAC-PI-ISO-0001", {"inputs": ["ALP-EPC-00PAC-IC-DSH-0001"]}, R)
    s.update("mdl_rule", "DL-IC-IDS", {"status": "superseded"}, R)
    plan = mdl.plan_sync(s)
    assert plan.orphans == ["ALP-EPC-00PAC-IC-DSH-0001"]
    assert main(["mdl", "sync", "--apply", "--cancel-orphans", "--reason", "t"]) == 0
    s = Store(project, who())
    assert s.get("document", "ALP-EPC-00PAC-IC-DSH-0001")["status"] == "cancelled"
    assert "inputs" not in s.get("document", "ALP-EPC-00PAC-PI-ISO-0001")


def test_sync_retype_and_input_pruning(project):
    s = setup(project)
    rule(s, "DL-A", "PID", "process", "system", "P&ID - {name}", quantity="q_pid_sheets")
    rule(s, "DL-B", "DSH", "i_and_c", "system", "Datasheets - {name}", quantity="q_instruments", inputs=["DL-A"])
    main(["mdl", "sync", "--apply", "--reason", "t"])
    s = Store(project, who())
    b = next(d for d in s.records("document") if d.get("rule") == "DL-B|00PAC|1")
    # a manual input survives, a rule-generated input of a rule no longer listed is removed
    s.update("document", b["id"], {"inputs": b["inputs"] + ["ALP-EPC-00000-PR-DBR-0001"]}, R)
    s.update("mdl_rule", "DL-B", {"inputs": None}, R)
    plan = mdl.plan_sync(s)
    assert plan.dropped == [(b["id"], "ALP-EPC-00PAC-PR-PID-0002")]
    assert dict(plan.update)[b["id"]]["inputs"] == ["ALP-EPC-00000-PR-DBR-0001"]
    # a type change renumbers: the old document is an orphan, a new one is created
    s.update("mdl_rule", "DL-B", {"type_code": "LST"}, R) if s.get("doc_type", "LST") else None
    plan = mdl.plan_sync(s)
    assert b["id"] in plan.orphans and any(rec["type_code"] == "LST" for _, rec in plan.create)


def test_curated_ewp_kept_within_discipline(project):
    s = setup(project)
    s.create("cwp", {"id": "CWP-05-PI-02", "cwa": "CWA-05", "discipline": "piping", "title": "CW piping 2", "scope": "x",
                     "status": "agreed"}, R)
    s.create("ewp", {"id": "EWP-05-PI-02", "cwp": "CWP-05-PI-02", "discipline": "piping", "title": "x", "status": "agreed"}, R)
    s.update("system", "00PAC", {"q_iso_sheets": 20}, R)
    rule(s, "DL-PI-ISO", "ISO", "piping", "system", "Isometrics - {name}", quantity="q_iso_sheets", ewp_disc="PI")
    s.update("document", "ALP-EPC-00PAC-PI-ISO-0001", {"rule": "DL-PI-ISO|00PAC|1"}, R)      # curated EWP-05-PI-01
    plan = mdl.plan_sync(s)
    ch = dict(plan.update)["ALP-EPC-00PAC-PI-ISO-0001"]
    assert "ewp" not in ch and ch["sheets"] == 20                        # rule would pick PI-02: same discipline, kept
    assert mdl.Ctx(s).ewp_for("CWA-05", "HV", None) is None              # no HV CWP and no fallback CWP of HV's list
    assert mdl.Ctx(s).ewp_for("CWA-05", "ME", None) == "EWP-05-PI-02"    # ME falls back to PI


def test_supplier_document_timing(project):
    s = setup(project)
    doc_base = {"title": "x", "discipline": "mechanical", "doc_type": "drawing", "type_code": "GAD", "originator": "M03",
                "system": "00PAC", "mr": "MR-ME-003", "inputs": ["ALP-EPC-00PAC-ME-MRQ-0001"]}
    s.create("document", {"id": "ALP-M03-00PAC-ME-GAD-0001", **doc_base, "po_weeks_ifr": 6, "po_weeks_final": 12}, R)
    s.create("document", {"id": "ALP-M03-00PAC-ME-GAD-0002", **doc_base, "po_weeks_ifr": -8, "po_weeks_final": 12}, R)
    s.update("mr", "MR-ME-003", {"manufacture_weeks": 40}, R)
    r = workflow.compute(s)
    m = r.mrs["MR-ME-003"]
    a, b = r.docs["ALP-M03-00PAC-ME-GAD-0001"], r.docs["ALP-M03-00PAC-ME-GAD-0002"]
    assert a.fixed and a.ifr == m["po"] + 30 and a.ifc == m["po"] + 60 and a.driver == "PO MR-ME-003"
    ship = m["ros"] - 20                                                     # on site - 4 weeks transport
    assert b.ifr == ship - 40 and b.ifc == ship + 60                        # FAT / O&M type: relative to shipment
    s.update("mr", "MR-ME-003", {"manufacture_weeks": 12}, R)               # short lead: weeks after PO x 0.5 (minimum)
    r = workflow.compute(s)
    m = r.mrs["MR-ME-003"]
    a = r.docs["ALP-M03-00PAC-ME-GAD-0001"]
    assert a.ifr == m["po"] + 15 and a.ifc == m["po"] + 30


def test_progressive_need_and_ramp(project):
    s = setup(project)
    s.update("doc_type", "ISO", {"progressive_share": 0.5}, R)
    r = workflow.compute(s)
    iso = r.docs["ALP-EPC-00PAC-PI-ISO-0001"]
    a = r.acts["CWP-05-PI-01"]
    e = r.ewps["EWP-05-PI-01"]
    assert e["float"] == a.es - 20 + round(0.5 * (a.ef - a.es)) - iso.ifc and e["driver"] == iso.id
    # ramp: 30 % of the capacity at NTP
    s.create("eng_resource", {"id": "RES-PR", "discipline": "process", "fte": 2, "ramp_weeks": 10}, R)
    r = workflow.compute(s)
    pid = r.docs["ALP-EPC-00PAC-PR-PID-0001"]
    first = min(pid.work)
    assert r.levelled and abs(sum(pid.work.values()) - 120) < 1e-6
    assert max(pid.work.values()) <= 15 + 1e-6
    assert pid.work[first] <= 15 * (0.3 + 0.7 * (first - r.dd) / 50) + 1e-6


def test_benchmark_and_rule_checks(project):
    s = setup(project)
    s.create("mdl_benchmark", {"id": "BM-PR", "discipline": "process", "originator": "EPC", "sheets_min": 100}, R)
    s.create("mdl_benchmark", {"id": "BM-PR-MAX", "discipline": "process", "originator": "EPC", "docs_max": 0}, R)
    rule(s, "DL-X", "LST", "general", "plant", "Some list") if s.get("doc_type", "LST") else None
    m = run_engine(project, Store(project, who()), registry()["engineering_plan"], {"pdf": "no"})
    w = " | ".join(m["warnings"])
    assert "BM-PR benchmark: process EPC sheets" in w and "vs indicative 100 - -" in w
    assert "required document(s) missing" in w and "mdl sync" in w
    from openpyxl import load_workbook
    rows = [tuple(c.value for c in r[:4]) for r in load_workbook(project.output / "engineering_plan" / "ALP_MDL.xlsx")["Checks"]
            .iter_rows(min_row=2)]
    assert any(r[0] == "BM-PR-MAX" and r[2] == "INFO" for r in rows)


def test_inputs_never_fall_back_across_instances_of_one_scope(project):
    """Regression: a supplier document of an MR without requisition must not take the requisition of another MR in the
    same area; a system document must not take the P&ID of another system in the same area."""
    s = setup(project)
    s.create("mr", {"id": "MR-ME-101", "title": "GT (IEC)", "discipline": "mechanical", "package_type": "iec_supply",
                    "cwps": ["CWP-05-PI-01"], "po_date": "2026-11-02", "status": "agreed"}, R)
    s.create("system", {"id": "00PAB", "title": "CW piping", "category": "cooling", "cwas": ["CWA-05"], "q_pid_sheets": 1}, R)
    rule(s, "DL-MR-MRQ", "MRQ", "mechanical", "mr", "MRQ - {name}", package_types=["equipment"])
    rule(s, "DL-V-DSH", "DSH", "mechanical", "mr", "Supplier DSH - {name}", org="SUPPLIER",
         package_types=["equipment", "iec_supply"], inputs=["DL-MR-MRQ"], po_weeks_ifr=4)
    rule(s, "DL-A", "PID", "process", "system", "P&ID - {name}", quantity="q_pid_sheets")
    rule(s, "DL-B", "CMP", "commissioning", "system", "Commissioning - {name}", inputs=["DL-A"]) \
        if s.get("doc_type", "CMP") else None
    inst, errs = mdl.required(s)
    assert not errs
    assert inst["DL-V-DSH|MR-ME-101|1"].inputs == []                          # IEC package: no MRQ of MR-ME-003
    assert inst["DL-V-DSH|MR-ME-003|1"].inputs == ["DL-MR-MRQ|MR-ME-003|1"]


def test_structure_input_mode_and_design_criteria_check(project):
    s = setup(project)
    s.create("system", {"id": "00UPC", "title": "CW pump house", "category": "civil_structural", "cwas": ["CWA-05"],
                        "q_fnd_sheets": 4}, R)
    s.create("system", {"id": "00UQA", "title": "Utility building", "category": "civil_structural", "cwas": ["CWA-05"],
                        "q_fnd_sheets": 2}, R)
    s.update("system", "00PAC", {"structures": ["00UPC"]}, R)
    dtype(s, "FND", "approval", prep=20)
    rule(s, "DL-V-FDL", "DSH", "civil", "mr", "Loads - {name}", org="SUPPLIER", package_types=["equipment"], po_weeks_ifr=6)
    rule(s, "DL-CV-EFD", "FND", "civil", "structure", "Equipment foundations - {name}", quantity="q_fnd_sheets",
         inputs=["DL-V-FDL@structure"])
    inst, _ = mdl.required(s)
    assert inst["DL-CV-EFD|00UPC|1"].inputs == ["DL-V-FDL|MR-ME-003|1"]       # CW pumps stand in the pump house
    assert inst["DL-CV-EFD|00UQA|1"].inputs == []                             # nothing of MR-ME-003 in the utility building
    m = run_engine(project, Store(project, who()), registry()["engineering_plan"], {"pdf": "no"})
    w = " | ".join(m["warnings"])
    assert "no design criteria document for mechanical" in w and "no design criteria document for process" not in w


def test_requisition_priority_follows_vendor_data_need(project):
    """The backward pass carries a supplier document's PO chain to its requisition (latest start)."""
    s = setup(project)
    s.update("mr", "MR-ME-003", {"bid_days": 30, "award_days": 10}, R)
    s.create("document", {"id": "ALP-M03-00PAC-CV-CAL-0001", "title": "Loads", "discipline": "civil", "doc_type": "calculation",
                          "type_code": "CAL" if s.get("doc_type", "CAL") else "DSH", "originator": "M03", "system": "00PAC",
                          "mr": "MR-ME-003", "po_weeks_ifr": 6, "inputs": ["ALP-EPC-00PAC-ME-MRQ-0001"]}, R)
    s.update("document", "ALP-EPC-00PAC-PI-ISO-0001", {"inputs": ["ALP-EPC-00PAC-PR-PID-0001", "ALP-M03-00PAC-CV-CAL-0001"]}, R)
    r = workflow.compute(s)
    mrq, loads, iso = r.docs["ALP-EPC-00PAC-ME-MRQ-0001"], r.docs["ALP-M03-00PAC-CV-CAL-0001"], r.docs["ALP-EPC-00PAC-PI-ISO-0001"]
    _, moff = workflow._review_offsets(mrq, s.get("doc_type", "MRQ"))
    t = s.get("doc_type", loads.rec["type_code"])
    _, loff = workflow._review_offsets(loads, t)
    loads_late_ifc = loads.late_start + loff + workflow._prep(loads, t)
    assert mrq.late_start + workflow._prep(mrq, s.get("doc_type", "MRQ")) + moff <= loads_late_ifc - loff - 30 - 40
