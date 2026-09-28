"""Tie-in register checks: completeness, HMB envelope over all cases, upstream chain, dates, agreement, coverage; and the
engine outputs (Excel, Word)."""
from openpyxl import load_workbook

from engine.core.runner import run_engine
from engine.core.store import Store
from engine.engines import registry
from engine.engines.tie_in_register import check_register, envelope, to_kgs
from tests.conftest import who


def _st(case, n, m, p, t):
    return {"hmb_case": case, "number": n, "description": f"stream {n}", "mass_flow": m, "pressure": p, "temperature": t}


CASES = [{"id": "A"}, {"id": "B"}]
STREAMS = [_st("A", 2, 20.0, 38.0, 25.0), _st("B", 2, 22.4, 39.5, 30.0), _st("A", 9, 1.0, 5.0, 50.0)]
BASE = dict(category="process", boundary="flange", size="12 in", connection="RF", design_pressure=63.0, design_temperature=80.0,
            isolation="valve", quality_refs=["source:S"], flow_unit="kg/s", status="agreed", agreed_by=["X", "Y"],
            available_by="2028-08-31", needed_by="2028-09-18")


def pt(pid, **kw):
    return dict(BASE, id=pid, **kw)


def run(points, powers=None, scope=()):
    return check_register(points, CASES, STREAMS, powers or {}, list(scope), "OWNER")


def results(checks, pid, chk):
    return [(st, txt) for p, c, st, txt in checks if p == pid and c == chk]


def test_units_and_envelope():
    assert to_kgs(115000, "Sm3/h", 0.7325) == 115000 * 0.7325 / 3600
    assert to_kgs(36, "t/h", None) == 10.0 and to_kgs(1, "m3/h", None) is None
    e = envelope(CASES, STREAMS)[2]
    assert e["m"][2:] == (22.4, "B") and round(e["p"][0], 5) == round(38.0 - 1.01325, 5) and e["T"][2:] == (30.0, "B")


def test_hmb_envelope_flags_flow_pressure_and_temperature():
    ok = pt("IF-01", hmb_stream=2, flow_max=23.0, operating_pressure_min=36.5, operating_pressure_max=41.0,
            operating_temperature_min=5.0, operating_temperature_max=40.0)
    c = run([ok])
    assert all(st == "OK" for st, _ in results(c, "IF-01", "HMB flow") + results(c, "IF-01", "HMB pressure")
               + results(c, "IF-01", "HMB temperature"))
    bad = dict(ok, flow_max=22.0, operating_pressure_max=37.5, operating_temperature_max=28.0)
    c = run([bad])
    assert results(c, "IF-01", "HMB flow")[0][0] == "WARN"
    assert "above range 37.5" in results(c, "IF-01", "HMB pressure")[0][1]
    assert results(c, "IF-01", "HMB temperature")[0][0] == "WARN"
    partial = dict(bad, hmb_check="flow")                 # pressure/temperature not applicable at this point
    c = run([partial])
    assert not results(c, "IF-01", "HMB pressure") and results(c, "IF-01", "HMB scope")[0][0] == "INFO"


def test_chain_pressure_design_protection_and_flow():
    up = pt("TP-G1", design_pressure=75.0, operating_pressure_min=45.0, operating_pressure_max=70.0, flow_max=115000.0,
            flow_unit="Sm3/h", density=0.7325, operating_temperature_min=5, operating_temperature_max=25)
    down = pt("IF-01", upstream="TP-G1", chain_dp=4.7, operating_pressure_min=36.5, operating_pressure_max=41.0, flow_max=23.0,
              operating_temperature_min=5, operating_temperature_max=40)
    c = run([up, down])
    assert results(c, "IF-01", "chain pressure")[0][0] == "OK"          # 45 - 4.7 = 40.3 >= 36.5
    assert results(c, "IF-01", "chain design")[0][0] == "WARN"          # 75 > 63 barg without protection
    assert results(c, "IF-01", "chain flow")[0][0] == "OK"              # 23.0 <= 23.40 kg/s
    c = run([up, dict(down, chain_dp=9.0, protection="slam-shut 46 barg", flow_max=24.0)])
    assert results(c, "IF-01", "chain pressure")[0][0] == "WARN"
    assert results(c, "IF-01", "chain design")[0][0] == "OK"
    assert results(c, "IF-01", "chain flow")[0][0] == "WARN"


def test_electrical_power_rating_limiter_and_voltage_chain():
    grid = dict(category="electrical", id="TP-E1", boundary="gantry", connection="clamps", isolation="GIS", protection="87L",
                voltage_nominal=380.0, voltage_min=342.0, voltage_max=420.0, short_circuit=50.0, short_circuit_rating=50.0,
                flow_max=640.0, flow_unit="MW", hmb_metric="net", status="draft", available_by="2028-07-15",
                needed_by="2028-07-22")
    gsu = dict(grid, id="IF-34", upstream="TP-E1", voltage_min=360.0, flow_max=560.0, flow_unit="MVA", power_factor=0.85,
               hmb_metric="gt_output", status="agreed", agreed_by=["IEPC", "IEC"])
    powers = {"A": {"net": 665.3, "gt_output": 470.0}}
    c = run([grid, gsu], powers)
    assert results(c, "TP-E1", "HMB power")[0][0] == "WARN"             # export above capacity
    assert results(c, "IF-34", "HMB power")[0][0] == "OK"               # 470 / 0.85 = 552.9 <= 560 MVA
    assert results(c, "IF-34", "chain voltage")[0][0] == "WARN"         # 342 kV not covered
    assert results(c, "TP-E1", "dates")[0][0] == "INFO"                 # 7 d float
    assert results(c, "TP-E1", "agreement")[0][0] == "WARN"             # draft
    c = run([dict(grid, flow_limited=True), dict(gsu, voltage_min=342.0)], powers)
    assert results(c, "TP-E1", "HMB power")[0][0] == "INFO"
    assert results(c, "IF-34", "chain voltage")[0][0] == "OK"


def test_completeness_dates_and_coverage():
    c = run([pt("TP-W1", design_pressure=None, needed_by="2028-08-01", scope_item="SC-043")],
            scope=[{"id": "SC-043", "supply": "OWNER", "terminal_point": "boundary"},
                   {"id": "SC-001", "supply": "OWNER", "terminal_point": "site boundary"},
                   {"id": "SC-006", "supply": "IEC", "terminal_point": "GT exhaust"}])
    comp = results(c, "TP-W1", "completeness")[0]
    assert comp[0] == "WARN" and "design_pressure" in comp[1] and "operating_pressure_min" in comp[1]
    assert results(c, "TP-W1", "dates")[0][0] == "WARN"                 # available after needed
    cov = {txt.split()[0]: st for p, chk, st, txt in c if chk == "coverage"}
    assert cov == {"SC-043": "OK", "SC-001": "WARN", "SC-006": "INFO"}


def test_engine_outputs(project):
    s = Store(project, who())
    s.create("party", {"id": "OWNER", "name": "Owner", "short_name": "Owner", "role": "owner", "country": "TR"}, "t")
    s.create("contract", {"id": "K-1", "title": "EPC contract", "owner": "OWNER", "contractor_parties": ["OWNER"],
                       "currency": "EUR", "price": 1.0, "status": "signed"}, "t")
    s.create("tie_in", {"id": "TP-W1", "contract": "K-1", "service": "Potable water", "medium": "water", "category": "utility",
                     "status": "draft", "location_text": "north gate", "owner_side": "meter", "contractor_side": "storage",
                     "boundary": "meter pit outlet flange", "size": "DN150"}, "t")
    m = run_engine(project, Store(project, who()), registry()["tie_in_register"], {"pdf": "no"})
    names = sorted(f.split("/")[-1] for f in m["files"])
    assert names == ["PROJECT-TBD_tie_in_register.docx", "PROJECT-TBD_tie_in_register.xlsx"]
    wb = load_workbook(project.output / "tie_in_register" / "PROJECT-TBD_tie_in_register.xlsx")
    assert wb.sheetnames == ["Register", "Checks", "HMB envelope", "Summary"]
    assert wb["Register"]["A5"].value == "TP-W1"
    assert any("TP-W1 completeness" in w for w in m["warnings"])
