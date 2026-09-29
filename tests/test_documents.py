"""Document-producing engines: Plant Design Basis Report (design_basis), technical specifications (specifications),
topographic survey report (site_survey) and the titles helper."""
from engine.core.runner import run_engine
from engine.core.store import Store
from engine.engines import registry
from tests.conftest import who
from tests.test_engineering import R, data, doc, dtype


def _text(path):
    from docx import Document
    d = Document(str(path))
    return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])


def _params(s):
    s.create("source", {"id": "SRC-TEST-0001", "title": "Survey", "originator": "SURV", "doc_ref": "SUR-1", "revision": "0"}, R)
    for i, (cat, par, v, vt, u) in enumerate([("site", "Site grade level", 15.0, None, "m a.s.l."),
                                              ("site", "Rotation plant grid to UTM grid", -0.3281, None, "deg"),
                                              ("site", "Scale factor plant grid to UTM", 0.999612, None, "-"),
                                              ("site", "Plant grid origin in UTM zone 35N (E 0 / N 0)", None,
                                               "E 541,250.000 m / N 4,532,480.000 m", "m"),
                                              ("ambient", "Summer design dry bulb (0.4 %)", 32.8, None, "degC")], 1):
        rec = {"id": f"DP-TEST-{i:04d}", "category": cat, "parameter": par, "status": "confirmed",
               "basis_refs": ["source:SRC-TEST-0001"]}
        rec.update({k: x for k, x in (("value", v), ("value_text", vt), ("unit", u)) if x is not None})
        s.create("design_parameter", rec, R)


def test_design_basis_report(project):
    s = data(project)
    _params(s)
    s.update("document", "ALP-EPC-00000-PR-DBR-0001", {"produced_by": "design_basis", "title": "Plant Design Basis Report"}, R)
    m = run_engine(project, Store(project, who()), registry()["design_basis"], {"pdf": "no"})
    assert "ALP-EPC-00000-PR-DBR-0001.docx" in m["files"]
    t = _text(project.output / "design_basis" / "ALP-EPC-00000-PR-DBR-0001.docx")
    assert "Plant Design Basis Report" in t and "Rev A - IFR, draft (not issued)" in t
    assert "Summer design dry bulb (0.4 %)" in t and "32.8 degC" in t and "SRC-TEST-0001" in t


def test_specification_from_clauses_scope_and_vdrl(project):
    s = data(project)
    _params(s)
    dtype(s, "SPC", "review", prep=15)
    doc(s, "ALP-EPC-00PAC-ME-SPC-0001", "mechanical", "SPC", mr="MR-ME-003", produced_by="specifications",
        title="CW Pumps - Technical Specification")
    doc(s, "ALP-M03-00PAC-ME-GAD-0001", "mechanical", "DSH", mr="MR-ME-003", po_weeks_ifr=6, po_weeks_final=12,
        title="CW Pumps - Supplier GA Drawings")
    s.create("reference", {"id": "REF-TEST-0001", "code": "ISO 13709", "title": "Centrifugal pumps"}, R)
    s.update("equipment", "00PAC10AP001", {"service": "Condenser cooling seawater", "capacity": 19000, "capacity_unit": "m3/h",
                                           "head": 17, "rated_power": 1120, "voltage": 10000, "redundancy": "2 x 50 %"}, R)
    for i, (sec, txt, par, v, u) in enumerate([("technical", "The pumps shall be vertical wet-pit mixed-flow pumps.", None, None, None),
                                               ("technical", None, "Rated flow", 19000, "m3/h"),
                                               ("codes", "Pumps per ISO 13709 where applicable.", None, None, None)], 1):
        rec = {"document": "ALP-EPC-00PAC-ME-SPC-0001", "section": sec, "seq": i, "status": "preliminary",
               "basis_refs": ["reference:REF-TEST-0001"]}
        rec.update({k: x for k, x in (("text", txt), ("parameter", par), ("value", v), ("unit", u)) if x is not None})
        s.create("spec_clause", rec, R)
    m = run_engine(project, Store(project, who()), registry()["specifications"], {"pdf": "no"})
    assert list(m["files"]) == ["ALP-EPC-00PAC-ME-SPC-0001.docx"] and not m["warnings"], m["warnings"]
    t = _text(project.output / "specifications" / "ALP-EPC-00PAC-ME-SPC-0001.docx")
    assert "00PAC10AP001" in t and "19,000 m3/h, 17 m, 1,120 kW, 10 kV" in t          # scope from the equipment list
    assert "ISO 13709" in t and "vertical wet-pit" in t and "Rated flow" in t
    assert "ALP-M03-00PAC-ME-GAD-0001" in t and "PO + 6 wk" in t                       # VDRL from the MDL
    assert "32.8 degC" in t                                                             # site conditions


def test_site_survey_report_checks_transformation_and_levels(project):
    import math
    s = data(project)
    _params(s)
    s.create("document", {"id": "ALP-EPC-00000-CV-RPT-0001", "title": "Topographic Survey Report", "discipline": "civil",
                          "doc_type": "report", "type_code": "LST", "originator": "EPC", "produced_by": "site_survey"}, R)
    rot, k = math.radians(-0.3281), 0.999612
    pts = [("CP-01", "control", 15, 12, 15.05, 0.004), ("CP-02", "control", 392, 293, 15.07, 0.030),
           ("SP-0001", "spot", 100, 100, 15.02, 0), ("SP-0002", "spot", 125, 100, 14.80, 0)]
    for pid, kind, e, n, z, err in pts:
        s.create("survey_point", {"id": pid, "kind": kind, "e": e, "n": n, "z": z, "area": "plot", "source": "SRC-TEST-0001",
                                  "utm_e": 541250 + k * (e * math.cos(rot) - n * math.sin(rot)) + err,
                                  "utm_n": 4532480 + k * (e * math.sin(rot) + n * math.cos(rot))}, R)
    m = run_engine(project, Store(project, who()), registry()["site_survey"], {"pdf": "no"})
    w = " | ".join(m["warnings"])
    assert "CP-02 transformation residual 30 mm" in w and "CP-01" not in w           # 4 mm passes, 30 mm fails
    assert "1 platform points outside" in w                                           # SP-0002 at -0.20 m
    t = _text(project.output / "site_survey" / "ALP-EPC-00000-CV-RPT-0001.docx")
    assert "2 spot heights" in t or "(2 spot heights" in t
    assert "SP-0002" in t and "-200" in t


def test_consistency_checks(project):
    from engine.core import consistency
    s = data(project)
    s.update("equipment", "00PAC10AP001", {"service": "CW", "redundancy": "2 x 50 %", "capacity": 19000, "head": 17,
                                           "design_pressure": 6, "design_temperature": 50, "material": "duplex",
                                           "rated_power": 1900, "voltage": 400}, R)
    rows = consistency.check(Store(project, who()))
    txt = " | ".join(f"{a} {b} {c} {d}" for a, b, c, d in rows)
    assert "1900 kW is not an IEC standard motor size" in txt and "1900 kW motor not on MV" in txt
    assert "2 x 50 % but 1 pump unit(s)" in txt                            # redundancy needs 2 registered units
    s.update("equipment", "00PAC10AP001", {"rated_power": 2000, "voltage": 10000, "quantity": 2}, R)
    txt = " | ".join(f"{a} {b} {c} {d}" for a, b, c, d in consistency.check(Store(project, who())))
    assert "motor size" not in txt and "redundancy" not in txt
    doc(s, "ALP-EPC-00PAC-ME-DSH-0009", "mechanical", "DSH", aveva_class="Specification Documents")
    txt = " | ".join(f"{a} {b} {c} {d}" for a, b, c, d in consistency.check(Store(project, who())))
    assert "ALP-EPC-00PAC-ME-DSH-0009 document inputs WARN" in txt              # no inputs: missing links
