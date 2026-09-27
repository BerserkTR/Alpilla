"""Contract engines: document assembled from clause/appendix records; traceability matrix computed from basis refs."""
from docx import Document
from openpyxl import load_workbook

from engine.core import validate
from engine.core.runner import run_engine
from engine.core.store import Store, StoreError
from engine.engines import registry
from tests.conftest import who

R = "t"


def seed_contract(s):
    s.create("project", {"id": "ALP", "name": "Alpilla CCGT", "plant_type": "CCGT", "schedule_start": "2026-11-02"}, R)
    s.create("party", {"id": "OWNER", "name": "Owner Co", "role": "owner"}, R)
    s.create("party", {"id": "LEAD", "name": "Leader Co", "role": "consortium_leader"}, R)
    s.create("party", {"id": "MEMB", "name": "Member Co", "role": "consortium_member"}, R)
    s.create("contract", {"id": "K-1", "title": "Test EPC Contract", "owner": "OWNER", "contractor_parties": ["LEAD", "MEMB"],
                          "leader": "LEAD", "currency": "EUR", "price": 100000000, "status": "signed"}, R)
    s.create("contract_clause", {"contract": "K-1", "number": "A1", "part": "agreement", "title": "Parties", "text": "Between us."}, R)
    s.create("contract_clause", {"contract": "K-1", "number": "A7", "part": "agreement", "title": "Signatures", "text": "Signed."}, R)
    s.create("contract_clause", {"contract": "K-1", "number": "12", "part": "conditions", "title": "Guarantees"}, R)
    s.create("contract_clause", {"contract": "K-1", "number": "12.2", "part": "conditions", "title": "Performance LDs",
                                 "text": "The Contractor shall pay:\n\n- LD one\n- LD two"}, R)
    s.create("contract_clause", {"contract": "K-1", "number": "2.1", "part": "conditions", "title": "Access", "text": "Access."}, R)
    s.create("contract_clause", {"contract": "K-1", "number": "2", "part": "conditions", "title": "The Owner"}, R)
    s.create("guarantee", {"id": "PG-01", "contract": "K-1", "parameter": "Net output", "category": "performance",
                           "guaranteed_value": 600, "unit": "MW", "direction": "min", "conditions": "SRC",
                           "ld_rate": 2200, "ld_unit": "EUR per kW shortfall"}, R)
    s.create("requirement", {"id": "ER-04.01", "contract": "K-1", "section": "ER-04 Performance", "title": "Net output",
                             "text": "Not less than 600 MW.", "discipline": "process", "verification": "test", "guarantee": "PG-01"}, R)
    s.create("requirement", {"id": "ER-04.02", "contract": "K-1", "section": "ER-04 Performance", "title": "Heat rate",
                             "text": "Low.", "discipline": "process", "verification": "test"}, R)
    s.create("milestone", {"id": "MS-01", "contract": "K-1", "title": "Advance", "percent": 10, "trigger": "signature"}, R)
    s.create("milestone", {"id": "MS-02", "contract": "K-1", "title": "Taking-Over", "percent": 90, "trigger": "TOC"}, R)
    s.create("scope_item", {"id": "SC-001", "contract": "K-1", "area": "Power island", "item": "Gas turbine",
                            "design": "MEMB", "supply": "MEMB", "install": "LEAD", "terminal_point": "GT exhaust flange"}, R)
    s.create("design_parameter", {"category": "performance", "parameter": "Net output", "value": 600, "unit": "MW",
                                  "status": "confirmed", "basis_refs": ["requirement:ER-04.01"]}, R)
    s.create("wbs", {"id": "ALP", "title": "Alpilla"}, R)
    s.create("wbs", {"id": "ALP.KD", "title": "Key dates", "parent": "ALP"}, R)
    s.create("activity", {"id": "KD-040", "title": "Taking-Over", "wbs": "ALP.KD", "phase": "commissioning", "duration": 0,
                          "type": "finish_milestone", "constraint": "finish_no_later_than", "constraint_date": "2029-07-04"}, R)


def test_contract_document(project):
    seed_contract(Store(project, who()))
    assert validate.run(project).ok
    m = run_engine(project, Store(project, who()), registry()["contract_document"], {"pdf": "no"})
    doc = Document(project.output / "contract_document" / "K-1_contract.docx")
    text = [p.text for p in doc.paragraphs]
    order = [t for t in text if t.startswith(("2  The Owner", "2.1  Access", "12  Guarantees", "12.2  Performance"))]
    assert order == ["2  The Owner", "2.1  Access", "12  Guarantees", "12.2  Performance LDs"]   # numeric, not text order
    assert "LD one" in text and "Signed." in text
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "EUR 10,000,000" in cells and "EUR 90,000,000" in cells and "100 %" in cells     # amounts computed
    assert ">= 600 MW" in cells and "2,200 EUR per kW shortfall" in cells
    assert any(c.startswith("For the Contractor (Leader)") for c in cells) and "GT exhaust flange" in cells
    assert "2029-07-04" in cells and m["warnings"] == []


def test_payment_schedule_must_total_100(project):
    s = Store(project, who())
    seed_contract(s)
    s.update("milestone", "MS-02", {"percent": 80}, "t")
    m = run_engine(project, Store(project, who()), registry()["contract_document"], {"pdf": "no"})
    assert any("add up to 90" in w for w in m["warnings"])


def test_requirements_matrix(project):
    seed_contract(Store(project, who()))
    run_engine(project, Store(project, who()), registry()["requirements_matrix"], {})
    wb = load_workbook(project.output / "requirements_matrix" / "ALP_requirements_matrix.xlsx")
    rows = {r[0]: r for r in wb["RTM"].iter_rows(min_row=5, values_only=True)}
    assert rows["ER-04.01"][8] == ">= 600 MW" and rows["ER-04.01"][9].startswith("design_parameter/DP-AAA-")
    assert rows["ER-04.01"][10] == "yes" and rows["ER-04.02"][10] == "no"
    cov = list(wb["Coverage"].iter_rows(min_row=5, max_row=5, values_only=True))[0]
    assert cov[:5] == ("ER-04 Performance", 2, 1, 1, 50.0)


def test_requirement_refs_validated(project):
    s = Store(project, who())
    seed_contract(s)
    import pytest
    with pytest.raises(StoreError, match="does not exist"):
        s.create("design_parameter", {"category": "site", "parameter": "x", "value": 1, "basis_refs": ["requirement:ER-99.01"]}, R)
    with pytest.raises(StoreError, match="referenced by"):
        s.delete("requirement", "ER-04.01", "t")
