"""Codes/standards and permits register: ER citation coverage, authority, agreement and permit timing checks; engine outputs."""
from openpyxl import load_workbook

from engine.core.runner import run_engine
from engine.core.store import Store
from engine.engines import registry
from engine.engines.standards_register import check_registers, cited_codes, covered
from tests.conftest import who

PARTIES = {"OWNER": {"role": "owner", "name": "Owner"}, "IEPC": {"role": "consortium_leader", "name": "EPC"},
           "IEC": {"role": "consortium_member", "name": "Vendor"}, "MOEUCC": {"role": "authority", "name": "Ministry"},
           "KGM": {"role": "authority", "name": "Highways"}}


def test_citations_and_coverage():
    t = "per ASME Section I or EN 12952, Law (No. 2872), IEC 60034 and IEEE 62271, HEI Standards, the Grid Code"
    assert cited_codes(t) == ["ASME I", "EN 12952", "Grid Code", "HEI", "IEC 60034", "IEEE 62271", "Law 2872"]
    codes = ["ASME Section I", "EN 12952", "Law 2872", "IEC 60034-1", "IEC/IEEE 62271-37-013", "HEI 2629", "Grid Code"]
    assert all(covered(c, codes) for c in cited_codes(t))
    assert not covered("ASME IX", ["ASME Section I"])           # 'IX' is not a part of 'I'
    assert not covered("IEC 6003", ["IEC 60034-1"]) and not covered("NFPA 99", codes)


def test_register_checks():
    refs = [{"id": "R1", "code": "Law 2872", "kind": "law", "status": "agreed", "agreed_by": ["OWNER", "IEPC", "IEC"],
             "responsible": ["IEPC", "IEC"], "authority": "MOEUCC"},
            {"id": "R2", "code": "Law 2918", "kind": "law", "status": "agreed", "agreed_by": ["OWNER", "IEPC"],
             "responsible": ["IEPC", "IEC"]},
            {"id": "R3", "code": "ISO 9001", "kind": "standard", "status": "proposed", "responsible": ["IEPC"]}]
    reqs = [{"id": "ER-1", "text": "Comply with Law No. 2872 and NFPA 850."}]
    permits = [{"id": "PMT-001", "authority": "MOEUCC", "legal_basis": ["R1"], "applicant": "OWNER", "support_by": ["IEC"],
                "apply_by": "2027-01-01", "lead_time": 60, "needed_by": "2027-02-15", "status": "agreed",
                "agreed_by": ["OWNER", "IEPC"]},
               {"id": "PMT-002", "authority": "IEPC", "legal_basis": ["R3"], "applicant": "IEPC", "apply_by": "2027-01-01",
                "lead_time": 30, "needed_by": "2027-03-01", "status": "proposed"}]
    c = {(a, b): (st, t) for a, b, st, t in check_registers(refs, permits, reqs, PARTIES, "OWNER")}
    er = [(st, t) for a, b, st, t in check_registers(refs, permits, reqs, PARTIES, "OWNER") if b == "ER coverage"]
    assert sorted(er) == [("OK", "cites Law 2872: in the register"), ("WARN", "cites NFPA 850: not in the register")]
    assert c[("R2", "authority")][0] == "WARN"
    assert c[("R2", "agreement")][0] == "WARN" and "missing IEC" in c[("R2", "agreement")][1]
    assert c[("R1", "agreement")][0] == "OK" and c[("R3", "agreement")][0] == "WARN"
    assert c[("PMT-001", "timing")] == ("WARN", "apply 2027-01-01 + 60 d = 2027-03-02; needed 2027-02-15: float -15 d")
    assert c[("PMT-001", "agreement")][0] == "WARN" and "missing IEC" in c[("PMT-001", "agreement")][1]
    assert c[("PMT-002", "legal basis")][0] == "WARN" and c[("PMT-002", "authority")][0] == "WARN"
    assert c[("KGM", "authority use")][0] == "INFO"


def test_engine_outputs(project):
    s = Store(project, who())
    s.create("party", {"id": "OWNER", "name": "Owner", "role": "owner"}, "t")
    s.create("party", {"id": "MOEUCC", "name": "Ministry", "role": "authority"}, "t")
    r = s.create("reference", {"code": "Law 2872", "title": "Environmental Law", "kind": "law", "precedence": 1,
                               "authority": "MOEUCC", "status": "proposed"}, "t")
    s.create("permit", {"id": "PMT-001", "title": "EIA decision", "authority": "MOEUCC", "legal_basis": [r["id"]],
                        "phase": "development", "applicant": "OWNER", "scope": "plant", "permit_status": "granted",
                        "status": "proposed"}, "t")
    m = run_engine(project, Store(project, who()), registry()["standards_register"], {"pdf": "no"})
    assert sorted(f.split("/")[-1] for f in m["files"]) == ["PROJECT-TBD_codes_standards_permits_register.docx",
                                                             "PROJECT-TBD_codes_standards_permits_register.xlsx"]
    wb = load_workbook(project.output / "standards_register" / "PROJECT-TBD_codes_standards_permits_register.xlsx")
    assert wb.sheetnames == ["Codes and standards", "Permits", "Authorities", "Checks"]
    assert wb["Permits"]["A5"].value == "PMT-001"
    assert any("PMT-001 agreement" in w for w in m["warnings"])
