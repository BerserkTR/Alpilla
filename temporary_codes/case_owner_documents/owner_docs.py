"""Case authoring: produce the Owner's issued documents (docx + pdf, xlsx data) from the draft texts.
These are simulated *received* documents (sources), not project outputs."""
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path(__file__).parent / "owner_md"
OUT = Path(sys.argv[1])
OUT.mkdir(parents=True, exist_ok=True)
OWNER = "Alpilla Enerji Üretim A.Ş."
NAVY = RGBColor(0x0B, 0x4F, 0x6C)
REV0 = {"MET": "2026-09-10", "MAR": "2026-09-12", "FUL-001": "2026-09-08", "FUL-002": "2026-09-15", "GEN": "2026-09-14",
        "ELE": "2026-09-11", "ENV": "2026-09-09", "GEO": "2026-09-05"}
XLSX = {"MET", "MAR", "FUL-001", "FUL-002"}


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), fill)
    cell._element.get_or_add_tcPr().append(shd)


def runs(par, text, size=None, bold=False):
    """Minimal inline markdown: **bold**, *italic*."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*"):
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            r = par.add_run(part); r.bold = bold
        if size:
            r.font.size = Pt(size)


def page_field(par):
    run = par.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        el = OxmlElement("w:fldChar" if kind else "w:instrText")
        if kind:
            el.set(qn("w:fldCharType"), kind)
        else:
            el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)
    run.font.size = Pt(8)


def table(doc, rows, header=True, widths=None, size=8.5):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Table Grid"
    for i, r in enumerate(rows):
        row = t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))              # keep each row on one page
        if header and i == 0:
            trpr.append(OxmlElement("w:tblHeader"))          # repeat the header row on continuation pages
        cells = row.cells
        for c, v in zip(cells, r):
            c.text = ""
            runs(c.paragraphs[0], v, size, bold=header and i == 0)
            if header and i == 0:
                shade(c, "D6E6EE")
    if widths:
        t.autofit = False
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Cm(w)
    return t


def parse(md: str):
    lines = md.splitlines()
    title = lines[0].lstrip("# ").strip()
    m = re.match(r"(ALP-OWN-[A-Z]+-\d{3}) Rev (\d+) — (.+)", title)
    return m.group(1), m.group(2), m.group(3), lines[1:]


def build(md_path: Path):
    docno, rev, title, lines = parse(md_path.read_text(encoding="utf-8"))
    key = docno.split("ALP-OWN-")[1].replace("-001", "") if not docno.startswith("ALP-OWN-FUL") else docno.split("ALP-OWN-")[1]
    date = next((re.search(r"\*\*Date:\*\* (\S+)", l).group(1) for l in lines if "**Date:**" in l), "")
    status = next((re.search(r"\*\*Status:\*\* (.+)$", l).group(1).strip() for l in lines if "**Status:**" in l), "")
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(1.8)
    st = doc.styles
    st["Normal"].font.name, st["Normal"].font.size = "Arial", Pt(9.5)
    for name, size in (("Heading 1", 13), ("Heading 2", 11)):
        rpr = st[name].element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is not None:
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                fonts.attrib.pop(qn(a), None)
        st[name].font.name, st[name].font.size, st[name].font.color.rgb = "Arial", Pt(size), NAVY
    hp = sec.header.paragraphs[0]
    runs(hp, f"**{OWNER}**  |  Alpilla 600 MW CCGT  |  {docno} Rev {rev}", 8)
    fp = sec.footer.paragraphs[0]
    runs(fp, "Owner document - controlled copy when printed from the project repository.  CASE STUDY - fictional.   Page ", 7)
    page_field(fp)

    # ---- cover
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(OWNER); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = NAVY
    p = doc.add_paragraph("Alpilla 600 MW Combined Cycle Power Plant"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold = True; r.font.size = Pt(18)
    doc.add_paragraph()
    table(doc, [["Document number", docno], ["Revision", rev], ["Date", date], ["Status", status],
                ["Confidentiality", "Restricted - Contract ALP-EPC-001 parties only"]], header=False, widths=(4.5, 12.5), size=10)
    doc.add_paragraph()
    table(doc, [["", "Prepared", "Checked", "Approved"],
                ["Function", "Owner's Engineer (discipline lead)", "Owner's Engineer (project engineer)", "Owner's Project Director"],
                ["Signature / date", f"signed {date}", f"signed {date}", f"signed {date}"]], widths=(3.5, 4.5, 4.5, 4.5), size=9)
    doc.add_paragraph()
    hist = [["Rev", "Date", "Description"], ["0", REV0.get(key, date), "Issued for tender / contract"]]
    if rev != "0":
        hist.append([rev, date, status])
    table(doc, hist, widths=(1.5, 3.0, 12.5), size=9)
    doc.add_paragraph()
    n = doc.add_paragraph(); n.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(n, "*CASE STUDY - all organisations, data and signatures are fictional.*", 8)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- body
    body = [l for l in lines if not l.startswith("**Project:**") and not l.startswith("**Date:**")
            and not l.startswith("*CASE STUDY") and "**Status:**" not in l]
    i = 0
    tables_data = []
    current_heading = title
    while i < len(body):
        l = body[i].rstrip()
        if not l.strip():
            i += 1; continue
        if l.startswith("## "):
            current_heading = l[3:].strip()
            doc.add_heading(current_heading, level=1)
        elif l.startswith("|"):
            rows = []
            while i < len(body) and body[i].startswith("|"):
                cells = [c.strip() for c in body[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table(doc, rows)
            doc.add_paragraph()
            tables_data.append((current_heading, rows))
            continue
        elif l.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); runs(p, l[2:])
        elif l.startswith("  ") and doc.paragraphs:
            runs(doc.paragraphs[-1], " " + l.strip())
        else:
            # hard-wrapped text: join following plain lines into one paragraph
            text = l.strip()
            while i + 1 < len(body) and body[i + 1].strip() and not re.match(r"^(#|\||- |  )", body[i + 1]):
                i += 1
                text += " " + body[i].strip()
            p = doc.add_paragraph(); runs(p, text)
        i += 1
    stem = f"{docno}_Rev{rev}"
    docx_path = OUT / f"{stem}.docx"
    doc.save(docx_path)
    with tempfile.TemporaryDirectory() as prof:
        subprocess.run(["soffice", "--headless", f"-env:UserInstallation=file://{prof}", "--convert-to", "pdf", "--outdir",
                        str(OUT), str(docx_path)], capture_output=True, timeout=300)
    made = [docx_path, OUT / f"{stem}.pdf"]
    if key in XLSX:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook(); wb.remove(wb.active)
        for k, (heading, rows) in enumerate(tables_data, 1):
            ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "", f"T{k} {heading}")[:31])
            ws.append([f"{docno} Rev {rev} - {heading}"]); ws["A1"].font = Font(bold=True)
            ws.append([])
            for j, r in enumerate(rows):
                vals = []
                for v in r:
                    v = v.replace("**", "")
                    try:
                        vals.append(float(v.replace(",", "")) if re.fullmatch(r"-?[\d,]+(\.\d+)?", v) else v)
                    except ValueError:
                        vals.append(v)
                ws.append(vals)
                if j == 0:
                    for c in ws[ws.max_row]:
                        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="D6E6EE")
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = max(12, min(45, max(len(str(c.value or "")) for c in col) + 2))
        x = OUT / f"{stem}_data.xlsx"
        wb.save(x)
        made.append(x)
    return docno, rev, title, date, made


if __name__ == "__main__":
    for md in sorted(SRC.glob("*.md")):
        docno, rev, title, date, made = build(md)
        print(docno, rev, [m.name for m in made if m.exists()])
