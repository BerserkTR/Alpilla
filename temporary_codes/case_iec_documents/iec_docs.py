"""Case authoring: Imaginary Electric (IEC) reference documents for the EPC (simulated vendor documents, fictional).
`python iec_docs.py <iec_data.json> <out_dir>` writes Word + PDF (+ Excel for tabular documents). All numbers come from
iec_data.json (iec_model.py), so datasheets, heat balance and interface data are mutually consistent."""
from __future__ import annotations

import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

D = json.loads(Path(sys.argv[1]).read_text())
OUT = Path(sys.argv[2])
OUT.mkdir(parents=True, exist_ok=True)
VENDOR = "Imaginary Electric Company"
DIV = "Power Generation Division - Plant Integration"
RED = RGBColor(0x8B, 0x1A, 0x1A)
DATE = "2026-09-27"
P, PS, ST = D["perf"], {s["no"]: s for s in D["streams"]}, D["pressures"]
AMB = {r["T_amb"]: r for r in D["ambient"]}


def f(x, n=1):
    return f"{x:,.{n}f}"


# ------------------------------------------------------------------ rendering
def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), fill)
    cell._element.get_or_add_tcPr().append(shd)


def runs(par, text, size=None, bold=False):
    for part in re.split(r"(\*\*[^*]+\*\*)", str(text)):
        if not part:
            continue
        r = par.add_run(part[2:-2] if part.startswith("**") else part)
        r.bold = bold or part.startswith("**")
        if size:
            r.font.size = Pt(size)


def page_field(par):
    run = par.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        el = OxmlElement("w:fldChar" if kind else "w:instrText")
        if kind:
            el.set(qn("w:fldCharType"), kind)
        else:
            el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)
    run.font.size = Pt(7)


def table(doc, rows, widths=None, size=8, header=True):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Table Grid"
    for i, r in enumerate(rows):
        row = t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if header and i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for c, v in zip(row.cells, r):
            c.text = ""
            runs(c.paragraphs[0], v, size, bold=header and i == 0)
            if header and i == 0:
                shade(c, "F2DCDB")
    if widths:
        usable = 26.0 if LANDSCAPE[0] else 17.4
        if sum(widths) < usable - 1.0 and len(widths) > 3:           # use the full text width of the page
            widths = [w * usable / sum(widths) for w in widths]
        t.autofit = False
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):     # LibreOffice reads the grid widths
            gc.set(qn("w:w"), str(int(Cm(w).twips)))
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Cm(w)
    doc.add_paragraph()


LANDSCAPE = [False]


def render(docno, title, blocks, landscape=False, xlsx=None, status="Issued for EPC engineering (reference data)"):
    LANDSCAPE[0] = False                                             # cover tables are portrait-sized
    doc = Document()
    sec = doc.sections[0]
    if landscape:
        from docx.enum.section import WD_ORIENT
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_height, sec.page_width = Cm(21.0), Cm(29.7)
    else:
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.6)
    st = doc.styles
    st["Normal"].font.name, st["Normal"].font.size = "Arial", Pt(9)
    for name, size in (("Heading 1", 12), ("Heading 2", 10)):
        rpr = st[name].element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is not None:
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                fonts.attrib.pop(qn(a), None)
        st[name].font.name, st[name].font.size, st[name].font.color.rgb = "Arial", Pt(size), RED
    runs(sec.header.paragraphs[0], f"**{VENDOR}**  |  Alpilla 600 MW CCGT - power island  |  {docno} Rev 0", 7.5)
    fp = sec.footer.paragraphs[0]
    runs(fp, f"{VENDOR} proprietary - for use by the Alpilla consortium only.  CASE STUDY - fictional.   Page ", 7)
    page_field(fp)
    # cover
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(VENDOR.upper()); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = RED
    p = doc.add_paragraph(DIV); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Alpilla 600 MW Combined Cycle Power Plant - 1 x IE-9H.02 multi-shaft power island")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold = True; r.font.size = Pt(17)
    doc.add_paragraph()
    table(doc, [["Document number", docno], ["Revision", "0"], ["Date", DATE], ["Status", status],
                ["Addressee", "Istanbul EPC Muhendislik ve Taahhut A.S. (consortium leader) - EPC engineering"],
                ["Contract", "ALP-EPC-001 (Owner contract); consortium agreement IEPC-IEC"]], (4.5, 12.5), 9.5, header=False)
    table(doc, [["", "Prepared", "Checked", "Approved"],
                ["Function", "IEC plant integration engineer", "IEC lead engineer (discipline)", "IEC project engineering manager"],
                ["Signature / date", f"signed {DATE}", f"signed {DATE}", f"signed {DATE}"]], (3.4, 4.5, 4.5, 4.6), 8.5)
    table(doc, [["Rev", "Date", "Description"], ["0", DATE, status]], (1.5, 3.0, 12.5), 8.5)
    n = doc.add_paragraph(); n.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(n, "CASE STUDY - Imaginary Electric, its products and all data are fictional.", 8)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    LANDSCAPE[0] = landscape
    xl = []
    for b in blocks:
        kind = b[0]
        if kind == "h":
            doc.add_heading(b[1], level=1)
        elif kind == "h2":
            doc.add_heading(b[1], level=2)
        elif kind == "p":
            runs(doc.add_paragraph(), b[1])
        elif kind == "ul":
            for item in b[1]:
                runs(doc.add_paragraph(style="List Bullet"), item)
        elif kind == "t":
            table(doc, b[1], b[2] if len(b) > 2 else None, b[3] if len(b) > 3 else 8)
            if len(b) > 4 and b[4]:
                xl.append((b[4], b[1]))
        elif kind == "pb":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    while doc.paragraphs and not doc.paragraphs[-1].text.strip() and not doc.paragraphs[-1].runs:
        el = doc.paragraphs[-1]._element                              # no trailing spacer -> no blank last page
        el.getparent().remove(el)
    stem = f"{docno}_Rev0"
    path = OUT / f"{stem}.docx"
    doc.save(path)
    with tempfile.TemporaryDirectory() as prof:
        subprocess.run(["soffice", "--headless", f"-env:UserInstallation=file://{prof}", "--convert-to", "pdf", "--outdir",
                        str(OUT), str(path)], capture_output=True, timeout=300)
    made = [path.name, f"{stem}.pdf"]
    if xl:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook(); wb.remove(wb.active)
        for name, rows in xl:
            ws = wb.create_sheet(name[:31])
            ws.append([f"{docno} Rev 0 - {name}"]); ws["A1"].font = Font(bold=True)
            ws.append([])
            for j, row in enumerate(rows):
                vals = []
                for v in row:
                    v = str(v).replace("**", "")
                    vals.append(float(v.replace(",", "")) if re.fullmatch(r"-?[\d,]+(\.\d+)?", v) else v)
                ws.append(vals)
                if j == 0:
                    for c in ws[ws.max_row]:
                        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="F2DCDB")
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = max(10, min(60, max(len(str(c.value or "")) for c in col[2:]) + 2))
                for c in col:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        wb.save(OUT / f"{stem}_data.xlsx")
        made.append(f"{stem}_data.xlsx")
    print(docno, made)


# ------------------------------------------------------------------ documents
def per001():
    ex = P["exhaust_mol_pct"]
    s = PS
    stream_rows = [["No", "Description", "Fluid", "From", "To", "Flow kg/s", "p bar(a)", "T degC", "h kJ/kg"]]
    for n in sorted(s):
        x = s[n]
        stream_rows.append([str(n), x["description"], x["fluid"], x["from_node"], x["to_node"], f(x["mass_flow"], 2),
                            f(x["pressure"], 4 if x["pressure"] < 0.1 else 2), f(x["temperature"], 1),
                            f(x["enthalpy"], 1) if x["enthalpy"] is not None else "-"])
    sect = [["HRSG section (gas flow order)", "Gas in degC", "Gas out degC", "Duty MW"]] + [
        [k.replace("_", " / "), f(a), f(b), f(q, 2)] for k, (a, b, q) in P["sections"].items()]
    amb = [["Ambient degC", "RH %", "Seawater degC", "GT MW", "GT HR kJ/kWh", "Exhaust kg/s", "Exhaust degC", "ST MW",
            "Gross MW", "Gross HR kJ/kWh", "Condenser mbar", "Gas Sm3/h", "Remark"]]
    for r in D["ambient"]:
        amb.append([f(r["T_amb"]), f(r["RH"], 0), f(r["T_sw"]), f(r["gt_P"]), f(r["gt_HR"], 0), f(r["gt_exh_flow"]),
                    f(r["gt_exh_T"]), f(r["st_P"]), f(r["gross"]), f(r["gross_HR"], 0), f(r["p_cond"]), f(r["fuel_Sm3h"], 0),
                    "GT at output limit 470 MW" if r["gt_limited"] else ""])
    pl = [["GT load %", "GT MW", "GT HR kJ/kWh", "Exhaust kg/s", "Exhaust degC", "ST MW", "Gross MW", "Gross HR kJ/kWh"]]
    for r in D["part_load"]:
        pl.append([f(r["gt_load"], 0), f(r["gt_P"]), f(r["gt_HR"], 0), f(r["gt_exh_flow"]), f(r["gt_exh_T"]), f(r["st_P"]),
                   f(r["gross"]), f(r["gross_HR"], 0)])
    blocks = [
        ("h", "1. Purpose and basis"),
        ("p", "This document gives the expected thermal performance of the IEC power island (1 x IE-9H.02 gas turbine with "
              "generator, 1 x IE-HR3 triple-pressure reheat HRSG, 1 x IE-ST3R reheat steam turbine with generator, seawater "
              "condenser) for the EPC's plant heat and mass balance. Values are **expected values, new and clean**, without "
              "guarantee margins; IEC's guarantee values to the consortium are given in section 7."),
        ("ul", ["Case **SRC-NG-100**: Site Reference Conditions per ER-02.03 - ambient 15.0 degC, 70 % RH, 1011.5 mbar; "
                "seawater 16.0 degC; base load on natural gas; power factor 0.85 lagging; 50.0 Hz.",
                f"Fuel: Owner reference composition ALP-OWN-FUL-001 / ER-03.02. **LHV calculated per ISO 6976 from this "
                f"composition: {P['fuel_LHV']:.2f} MJ/kg ({P['fuel_LHV_vol']:.2f} MJ/Sm3, 15 degC/15 degC).** Heat input = "
                "fuel mass flow x this LHV. Fuel gas heated to 215 degC in the IEC performance gas heater with IP feedwater.",
                "GT inlet pressure loss 10 mbar, exhaust back-pressure 36 mbar (HRSG 30 mbar + stack 6 mbar, stack by EPC).",
                "Steam piping between HRSG and steam turbine is **by EPC**; IEC assumed pressure drops HP 7.0 bar, hot "
                "reheat 1.6 bar, cold reheat 0.8 bar, LP 0.45 bar, and temperature drops of 3 K (HP, HRH), 1 K (CRH), "
                "2 K (LP). The EPC shall confirm or IEC will correct the performance.",
                "Boiler feed pumps and condensate pumps are by EPC; IEC assumed BFP discharge 209 bar(a) (HP) / 46 bar(a) (IP), "
                "CEP discharge 16 bar(a).",
                "Steam turbine gland leak-offs, valve stem leakages and HRSG blowdown are not shown in this summary heat "
                "balance (blowdown zero at design; leak-offs are internal to the ST)."]),
        ("h", "2. Summary at SRC-NG-100"),
        ("t", [["Item", "Value", "Unit"],
               ["GT output at generator terminals", f(P["gt_P_gen"], 2), "MW"],
               ["GT heat rate (LHV)", f(P["gt_HR"], 1), "kJ/kWh"],
               ["GT heat input (LHV)", f(P["Q_fuel"], 2), "MW"],
               ["Fuel gas flow", f(P["fuel_flow"], 3) + " kg/s = " + f(P["fuel_Sm3h"], 0), "Sm3/h"],
               ["Compressor inlet air flow", f(P["air_flow"], 2), "kg/s"],
               ["GT exhaust flow / temperature", f(P["exh_flow"], 2) + " / " + f(PS[4]["temperature"], 1), "kg/s / degC"],
               ["GT generator efficiency", f(P["gt_eta_gen"] * 100, 2), "%"],
               ["GT losses other than generator (lube oil, radiation)", f(P["gt_Q_other"], 2), "MW"],
               ["Performance gas heater duty", f(P["Q_fgh"], 2), "MW"],
               ["HRSG heat transferred to water/steam", f(P["Q_hrsg"], 2), "MW"],
               ["HRSG heat loss (casing, % of gas-side duty)", f(P["hrsg_loss_pct"], 1), "%"],
               ["Stack temperature", f(P["T_stack"], 1), "degC"],
               ["ST output at generator terminals", f(P["st_P_gen"], 2), "MW"],
               ["ST mechanical / generator efficiency", f(P["st_eta_mech"] * 100, 2) + " / " + f(P["st_eta_gen"] * 100, 2), "%"],
               ["Condenser pressure", f(P["p_cond_mbar"], 2), "mbar(a)"],
               ["Condenser duty", f(P["Q_cond"], 2), "MW"],
               ["Circulating water flow (7 K rise)", f(P["cw_flow"], 0) + " kg/s = " + f(P["cw_m3h"], 0), "m3/h"],
               ["**Power island gross output**", "**" + f(P["gross"], 2) + "**", "MW"],
               ["**Power island gross heat rate (LHV)**", "**" + f(P["gross_HR"], 1) + "**", "kJ/kWh"]], (9.5, 5.0, 3.0), 8.5,
         "Summary"),
        ("p", "Net output and net heat rate depend on the EPC balance-of-plant auxiliary loads and transformer losses; IEC's "
              "auxiliary loads are listed in IEC-ALP-REQ-001 section 7."),
        ("h", "3. Stream data at SRC-NG-100"),
        ("p", "Node codes: AMB ambient, FGS EPC fuel gas system, FGH IEC performance gas heater, GT gas turbine, HRSG, STK "
              "stack (EPC), PIPE steam piping HRSG-ST (EPC), ST steam turbine, COND condenser, CEP condensate pumps (EPC), "
              "CMIX condensate/FGH-return junction (EPC), BFP boiler feed pumps (EPC), SEA intake, OUT outfall."),
        ("t", stream_rows, (1.0, 9.0, 1.8, 1.5, 1.5, 2.2, 2.2, 1.8, 1.8), 7.5, "Streams SRC-NG-100"),
        ("h", "4. GT exhaust gas composition at SRC-NG-100 (mol %, wet)"),
        ("t", [["N2", "O2", "Ar", "CO2", "H2O"], [f(ex["N2"], 2), f(ex["O2"], 2), f(ex["Ar"], 2), f(ex["CO2"], 2), f(ex["H2O"], 2)]],
         None, 8.5, "Exhaust composition"),
        ("h", "5. HRSG temperature profile at SRC-NG-100"),
        ("t", sect, (7.0, 3.0, 3.0, 3.0), 8.5, "HRSG profile"),
        ("p", f"Pinch points HP/IP/LP: {P['pinch'][0]:.0f} / {P['pinch'][1]:.0f} / {P['pinch'][2]:.0f} K; approach "
              f"temperatures at the economiser outlets: {P['approach'][0]:.0f} / {P['approach'][1]:.0f} / {P['approach'][2]:.0f} K. "
              f"Drum saturation temperatures HP/IP/LP: {P['Tsat']['HP_drum']} / {P['Tsat']['IP_drum']} / {P['Tsat']['LP_drum']} degC. "
              f"Steam turbine section outputs (shaft): HP {P['P_hp']:.2f} MW, IP {P['P_ip']:.2f} MW, LP {P['P_lp']:.2f} MW; "
              f"LP exhaust moisture {100 - P['x_ueep'] * 100:.2f} %."),
        ("pb",),
        ("h", "6. Performance at other ambient conditions and part load (estimates)"),
        ("p", "Base load on natural gas, seawater temperature as shown, other conditions as SRC. GT from the IE-9H.02 correction "
              "curves; the bottoming cycle is re-estimated with the design pinch/approach values. **These are preliminary "
              "estimates, not off-design ratings of the fixed HRSG and condenser surfaces**; rated off-design heat balances "
              "follow with the HRSG thermal design (IEC-ALP-DS-104 Rev A)."),
        ("t", amb, (1.4, 0.9, 1.3, 1.2, 1.4, 1.4, 1.3, 1.1, 1.2, 1.5, 1.3, 1.4, 2.2), 6.5, "Ambient cases"),
        ("p", "GT output is limited to **470 MW** at the generator terminals (generator and shaft limit) below about -1 degC; "
              "the GT then runs at part load with the IGV closing."),
        ("t", pl, (1.6, 1.6, 2.0, 2.0, 2.0, 1.6, 1.8, 2.2), 8, "Part load"),
        ("p", "Part load at SRC with the steam cycle at fixed pressure (sliding-pressure operation will be optimised in "
              "detailed design; values are conservative estimates)."),
        ("h", "7. IEC guarantees to the consortium (power island)"),
        ("t", [["Guarantee (SRC-NG-100, new and clean, per ASME PTC 46 / ISO 2314)", "Value"],
               ["Power island gross output at generator terminals", "not less than 614.0 MW"],
               ["Power island gross heat rate (LHV)", "not more than 5,745 kJ/kWh"],
               ["GT output at generator terminals", "not less than 418.0 MW"],
               ["Condenser pressure at 16.0 degC seawater, 36,050 m3/h", "not more than 35.5 mbar(a)"],
               ["IEC auxiliary loads (IEC-ALP-REQ-001 table 7.1, excl. GSU losses)", "not more than 2,700 kW"],
               ["NOx / CO at HRSG outlet (15 % O2, dry), 40-100 % GT load, natural gas", "30 / 30 mg/Nm3"]], (12.0, 5.0), 8.5),
        ("p", "Degradation is excluded from the values above; recoverable and non-recoverable degradation curves follow in "
              "IEC-ALP-PER-002."),
    ]
    render("IEC-ALP-PER-001", "Power Island Thermal Performance Data (Heat Balance)", blocks, landscape=True)


def ds101():
    gt = PS
    blocks = [
        ("h", "1. General"),
        ("t", [["Item", "Data"],
               ["Model", "IE-9H.02 heavy-duty single-shaft gas turbine, 50 Hz, 3,000 rpm, cold-end drive, axial exhaust"],
               ["Fleet experience", "11 units in commercial operation, lead unit 38,000 fired hours (ER-01.08: complies)"],
               ["Compressor", "14 stages, pressure ratio 23.8:1, variable IGV + 3 rows of variable stator vanes"],
               ["Combustion system", "IE DLN-H2, 16 can-annular combustors, dual fuel (natural gas / LDO), water injection on LDO, "
                                     "up to 30 vol % H2 capable hardware (5 % without modification)"],
               ["Turbine", "4 stages, air-cooled, single-crystal stage 1 blades, thermal barrier coatings"],
               ["Starting", "Static frequency converter (SFC) through the generator, 8.5 MW input from 10.5 kV"],
               ["Enclosure", "Acoustic enclosure 85 dB(A) at 1 m, ventilation 2 x 100 % fans, gas detection, CO2 fire suppression"],
               ["Inlet air system", "Filter house 3 stages: G4 weather louvre/coalescer, F9, EPA E12; anti-icing by compressor "
                                    "bleed; silencer. Evaporative cooler: **space provision only, not included**"]], (5.0, 12.4), 8.5),
        ("h", "2. Ratings"),
        ("t", [["Condition", "Output MW", "Heat rate kJ/kWh", "Efficiency %", "Exhaust kg/s", "Exhaust degC"],
               ["ISO 2314 (15 degC, 60 %, 1013.25 mbar, no inlet/outlet losses), gas", "432.0", "8,240", "43.69", "801.0", "638.0"],
               ["SRC-NG-100 (10/36 mbar losses, fuel 215 degC)", f(P["gt_P_gen"]), f(P["gt_HR"], 0), f(P["gt_eta"] * 100, 2),
                f(gt[4]["mass_flow"]), f(gt[4]["temperature"])],
               ["SRC, LDO with water injection (base load)", "395.0", "8,650", "41.62", "812.0", "622.0"],
               ["Maximum output (cold ambient, generator/shaft limit)", "470.0", "-", "-", "-", "-"]],
         (6.6, 1.8, 2.2, 1.8, 1.8, 1.8), 8),
        ("h", "3. Fuel requirements (at IEC terminal points, see IEC-ALP-IF-001)"),
        ("t", [["Parameter", "Natural gas (IF-01, performance gas heater inlet)", "LDO (IF-02, liquid fuel skid inlet)"],
               ["Pressure", "36.5 barg minimum at base load; 41.0 barg maximum; max. rate of change 0.7 bar/s",
                "3.0 - 6.0 barg"],
               ["Temperature", "5 - 40 degC and at least 28 K above the hydrocarbon and water dew points",
                "10 - 45 degC"],
               ["Flow at base load (SRC)", f(P["fuel_flow"], 2) + " kg/s", "22.1 kg/s (+ demin water 22.5 kg/s at IF-03)"],
               ["Quality", "ALP-OWN-FUL-001 range; Wobbe index change <= 1 %/min; trace metals Na+K <= 0.02 mg/Sm3 equivalent; "
                           "particulates filtered to 3 um absolute (EPC final filter)",
                "ER-03.07 limits at the GT inlet; filtration 5 um (IEC skid)"],
               ["Fuel heating", "IEC performance gas heater to 215 degC with IP feedwater; start-up heating by EPC gas "
                                "station (electric or water-bath heater, gas >= 28 K superheat)", "-"]], (3.2, 8.2, 6.0), 8),
        ("h", "4. Emissions and noise"),
        ("t", [["Parameter", "Natural gas", "LDO"],
               ["NOx at GT outlet (15 % O2, dry), 40-100 % load", "<= 50 mg/Nm3", "<= 90 mg/Nm3 (water/fuel ratio 1.02)"],
               ["CO at GT outlet, 40-100 % load", "<= 20 mg/Nm3", "<= 25 mg/Nm3"],
               ["Minimum emissions-compliant GT load (MECL)", "40 % GT load (with SCR)", "50 % GT load"],
               ["Near-field noise, enclosure and inlet", "85 dB(A) at 1 m", "85 dB(A) at 1 m"]], (8.0, 4.7, 4.7), 8.5),
        ("h", "5. Operation and maintenance"),
        ("t", [["Item", "Data"],
               ["Start to base load (GT)", "hot start 25 min, warm 35 min, cold 45 min (GT only; CC limited by HRSG/ST)"],
               ["Load ramp (GT)", "40 MW/min normal, 65 MW/min fast"],
               ["Combustion inspection / hot gas path / major inspection", "12,500 / 25,000 / 50,000 EOH, or 900 / 900 / 1,800 "
                                                                           "starts, whichever first"],
               ["Equivalent operating hours (EOH)", "EOH = fired hours + 20 x starts + 60 x trips from base load + 2 x h on LDO"],
               ["Lube oil", "ISO VG 32 mineral, tank 28 m3, 2 x 100 % AC pumps + DC emergency pump; cooling duty 1.6 MW"],
               ["Compressor washing", "on-line and off-line; detergent skid by IEC; demin water 5 m3 per off-line wash"]], (6.0, 11.4), 8.5),
        ("h", "6. Dimensions, weights and transport"),
        ("t", [["Item", "Dimensions L x W x H (m)", "Weight (t)"],
               ["GT core engine (shipping, with transport frame)", "13.2 x 5.9 x 5.6", "405"],
               ["Exhaust diffuser (2 pieces)", "8.5 x 7.2 x 7.0", "2 x 62"],
               ["Inlet filter house (modules)", "18 modules", "max. 28 per module"],
               ["GT rotor (maintenance lift)", "11.0 x 3.2 x 3.2", "112"]], (8.0, 5.0, 4.4), 8.5),
        ("p", "The GT core engine exceeds the 120 t road limit (ALP-OWN-GEN-001): delivery by barge to the landing TP-M1; the EPC "
              "shall design the barge landing and the haul road for a 480 t gross transport (SPMT)."),
        ("h", "7. Auxiliary power (base load, SRC)"),
        ("p", "GT package auxiliaries 1,180 kW (lube oil, hydraulic, enclosure ventilation, fuel gas skid, filter house); "
              "SFC 8.5 MW during start-up only. See IEC-ALP-REQ-001 table 7.1."),
    ]
    render("IEC-ALP-DS-101", "Gas Turbine IE-9H.02 - Datasheet", blocks)


def ds102():
    blocks = [
        ("h", "1. Generator data"),
        ("t", [["Item", "GT generator IE-GH560", "ST generator IE-GA270"],
               ["Rated apparent power", "560 MVA", "270 MVA"],
               ["Rated power factor / active power", "0.85 lagging / 476 MW", "0.85 lagging / 229.5 MW"],
               ["Reactive capability at rated active power", "0.85 lagging to 0.95 leading (ER-04.08)", "0.85 lagging to 0.95 leading"],
               ["Rated voltage / frequency / speed", "21.0 kV +/- 5 % / 50 Hz / 3,000 rpm", "15.75 kV +/- 5 % / 50 Hz / 3,000 rpm"],
               ["Cooling", "hydrogen-cooled rotor and stator core, 4.0 barg H2; water-cooled stator winding", "TEWAC (air, water coolers)"],
               ["Efficiency at SRC load (incl. excitation)", f(P["gt_eta_gen"] * 100, 2) + " %", f(P["st_eta_gen"] * 100, 2) + " %"],
               ["Losses to cooling water at SRC", "4.3 MW (H2 coolers + stator water)", "2.2 MW (air coolers)"],
               ["Excitation", "static, from generator terminals via excitation transformer", "static"],
               ["Short-circuit ratio / Xd'' (unsaturated)", "0.50 / 0.21 pu", "0.52 / 0.18 pu"],
               ["Insulation / temperature rise", "class F / class B", "class F / class B"],
               ["Generator circuit breaker (GCB)", "included, IEC 62271-37-013, 21 kV, 20 kA rated current, 130 kA", "included, 15.75 kV, 12.5 kA, 100 kA"],
               ["Isolated phase busduct (IPB)", "generator - GCB - GSU LV, with tap-off to unit auxiliary transformer (UAT by EPC)", "as GT"],
               ["Stator transport weight / dimensions", "340 t / 10.8 x 4.6 x 4.9 m", "185 t / 8.6 x 4.1 x 4.2 m"],
               ["Rotor weight (maintenance lift)", "72 t", "46 t"],
               ["H2 / CO2 consumption", "H2 make-up 12 Nm3/day; CO2 for purging 150 Nm3 per purge", "-"],
               ["Protection", "IEC generator protection panels (IEEE C37.102 functions, duplicated main 1 / main 2)", "as GT"]],
         (5.5, 6.2, 5.7), 8),
        ("h", "2. Interfaces"),
        ("ul", ["H2 and CO2 bottle racks and supply piping to the IEC gas control panel: **by EPC** (IEC-ALP-DOR-001 item 5.6).",
                "Closed cooling water for H2 coolers, stator water coolers and TEWAC coolers: by EPC (IEC-ALP-REQ-001 table 5.1).",
                "UAT tap-off from the IPB between GCB and GSU: IPB tap-off flange by IEC, UAT and MV busduct by EPC.",
                "Neutral earthing: high-resistance earthing transformer in the generator neutral cubicle by IEC."]),
    ]
    render("IEC-ALP-DS-102", "Generators, GCB and Isolated Phase Busduct - Datasheet", blocks)


def ds103():
    blocks = [
        ("h", "1. General"),
        ("t", [["Item", "Data"],
               ["Model", "IE-ST3R reheat condensing steam turbine, 3,000 rpm, two casings: HP barrel turbine and combined IP/LP "
                         "turbine with double-flow LP section, **side exhaust** (both sides) to the condenser"],
               ["Last stage blade / exhaust annulus", "1,070 mm, 12 % Cr steel with erosion shields / 2 x 10.8 m2"],
               ["Output at SRC-NG-100", f(P["st_P_gen"], 2) + " MW at the generator terminals"],
               ["Maximum continuous output", "235 MW (cold ambient and valves wide open)"],
               ["Section efficiencies (SRC, isentropic, shaft)", "HP 89.0 %, IP 92.5 %, LP 88.5 % (to the ELEP, incl. moisture)"],
               ["Exhaust loss at SRC", f(P["st_exh_loss"], 1) + " kJ/kg"],
               ["Governing", "throttle (sliding pressure above 50 % load), HP/IP combined stop-control valves, LP admission valve"],
               ["Bypass valves (supplied by IEC, installed in EPC piping)", "HP bypass 100 % to cold reheat; IP bypass 100 % and LP "
                                                                             "bypass 100 % to condenser dump devices"],
               ["Gland steam", "self-sealing above 40 % load; start-up sealing steam 4.5 t/h at 10 bar(a), 250 - 300 degC from the "
                               "EPC auxiliary boiler; gland steam condenser with 2 x 100 % exhausters (IEC)"],
               ["Lube / control oil", "ISO VG 46 lube oil, tank 32 m3, 2 x 100 % AC + DC pump; fire-resistant EHC fluid (ER-07.06)"],
               ["Turning gear", "hydraulic, 5 rpm"]], (5.0, 12.4), 8.5),
        ("h", "2. Steam conditions (at the ST terminal points)"),
        ("t", [["Admission", "Flow kg/s", "Pressure bar(a)", "Temperature degC", "Design pressure / temperature"],
               ["HP main stop valve inlet (IF-06)", f(PS[7]["mass_flow"], 2), f(PS[7]["pressure"], 1), f(PS[7]["temperature"], 1),
                "200 bar(a) / 605 degC"],
               ["HP exhaust / cold reheat (IF-07)", f(PS[8]["mass_flow"], 2), f(PS[8]["pressure"], 1), f(PS[8]["temperature"], 1),
                "52 bar(a) / 420 degC"],
               ["IP stop-control valve inlet (IF-10)", f(PS[12]["mass_flow"], 2), f(PS[12]["pressure"], 1), f(PS[12]["temperature"], 1),
                "48 bar(a) / 605 degC"],
               ["LP admission valve inlet (IF-12)", f(PS[14]["mass_flow"], 2), f(PS[14]["pressure"], 2), f(PS[14]["temperature"], 1),
                "8 bar(a) / 330 degC"],
               ["LP exhaust (to condenser)", f(PS[16]["mass_flow"], 2), f(PS[16]["pressure"], 4), f(PS[16]["temperature"], 1),
                "full vacuum / 1.5 bar(a)"]], (5.5, 2.0, 2.4, 2.6, 4.9), 8),
        ("h", "3. Operating limits"),
        ("t", [["Parameter", "Value"],
               ["Exhaust pressure: continuous / alarm / trip", "120 / 150 / 250 mbar(a)"],
               ["Main and reheat steam temperature: continuous maximum / 400 h per year", "605 degC / 615 degC"],
               ["Start-up (ST roll to full load) after 8 h / 48 h / cold", "25 / 60 / 150 min"],
               ["Allowable nozzle loads", "per IEC-ALP-IF-001 attachment A (NEMA SM-23 x 3.0 for HP/IP, x 1.5 for LP admission)"]],
         (9.0, 8.4), 8.5),
        ("h", "4. Weights"),
        ("t", [["Component", "Weight (t)", "Remark"],
               ["HP turbine (shipped assembled)", "92", "heaviest lift in the turbine hall for maintenance: IP/LP rotor 86 t"],
               ["IP/LP outer casing lower half (2 pieces)", "2 x 78", ""],
               ["IP/LP rotor", "86", "turbine hall crane: 100 t main hook minimum"],
               ["ST generator stator", "185", "erection by hydraulic jacking/skidding, not by crane"]], (6.0, 2.5, 8.9), 8.5),
    ]
    render("IEC-ALP-DS-103", "Steam Turbine IE-ST3R - Datasheet", blocks)


def ds104():
    s = P["sections"]
    blocks = [
        ("h", "1. General"),
        ("t", [["Item", "Data"],
               ["Type", "IE-HR3 horizontal gas flow, triple-pressure reheat, natural circulation (HP drum thin-wall design with "
                        "external cyclone separators), modular (22 modules), self-supporting casing, internally insulated"],
               ["Design code", "ASME Section I, PED 2014/68/EU module G (notified body)"],
               ["Cycling", "250 starts/year, 30-year fatigue life per EN 12952-3 (hot 180, warm 60, cold 10 per year)"],
               ["Gas side pressure loss at SRC", "30 mbar (incl. SCR and CO catalyst, clean)"],
               ["Outlet to stack", "HRSG outlet transition duct flange, 9.2 m diameter (stack by EPC, IF-26)"],
               ["SCR", "V2O5/TiO2 honeycomb at 360 - 375 degC (between HP evaporator and HP economiser 2), aqueous ammonia "
                       "24.5 % injection grid (IEC); design NOx at HRSG outlet 22 mg/Nm3 (15 % O2); NH3 slip <= 3 mg/Nm3; "
                       "catalyst life 24,000 h"],
               ["CO catalyst", "Pt/Pd oxidation catalyst at 450 degC (after HP superheater/reheater)"],
               ["Aqueous ammonia consumption", "125 kg/h at base load on gas (24.5 % solution); 190 kg/h on LDO"],
               ["Stack damper", "not included (stack by EPC)"]], (4.5, 12.9), 8.5),
        ("h", "2. Thermal data at SRC-NG-100"),
        ("t", [["Circuit", "Flow kg/s", "Outlet pressure bar(a)", "Outlet temperature degC", "Drum pressure bar(a)",
                "Design pressure bar(a)"],
               ["HP superheater", f(PS[6]["mass_flow"], 2), f(PS[6]["pressure"], 1), f(PS[6]["temperature"], 1), f(ST["HP_drum"], 1), "210"],
               ["Reheater", f(PS[11]["mass_flow"], 2), f(PS[11]["pressure"], 1), f(PS[11]["temperature"], 1), "-", "50"],
               ["IP superheater", f(PS[10]["mass_flow"], 2), f(PS[10]["pressure"], 1), f(PS[10]["temperature"], 1), f(ST["IP_drum"], 1), "50"],
               ["LP superheater", f(PS[13]["mass_flow"], 2), f(PS[13]["pressure"], 2), f(PS[13]["temperature"], 1), f(ST["LP_drum"], 1), "10"],
               ["Condensate preheater", f(PS[20]["mass_flow"], 2), f(ST["LP_drum"] + 1, 1), f(P["Tsat"]["LP_drum"] - P["approach"][2], 1),
                "-", "25"]], (3.8, 2.0, 2.8, 3.0, 2.8, 3.0), 8),
        ("t", [["Section", "Gas in degC", "Gas out degC", "Duty MW"]] + [[k.replace("_", " / "), f(a), f(b), f(q, 2)]
                                                                         for k, (a, b, q) in s.items()], (7.0, 3.0, 3.0, 3.0), 8.5),
        ("p", f"Stack temperature {P['T_stack']:.1f} degC at SRC. Condensate preheater inlet temperature is kept at not less than "
              "50 degC (acid dew point margin, LDO) by a 2 x 100 % CPH recirculation pump (IEC); at SRC the mixed inlet from the "
              "condensate and the fuel gas heater return is 30.9 degC before recirculation."),
        ("h", "3. Water/steam quality required"),
        ("p", "Per IEC-ALP-REQ-001 section 4 (IAPWS TGD and VGB-S-010): all-volatile treatment (AVT(O)) with trisodium "
              "phosphate provision for the HP drum; continuous blowdown 0.5 % maximum at start-up, 0 % in normal operation."),
        ("h", "4. Weights and transport"),
        ("t", [["Item", "Data"],
               ["Heaviest pressure-part module (HP evaporator)", "185 t, 24.0 x 3.6 x 26.5 m (shipped horizontal)"],
               ["HP drum", "96 t, 14.0 m long, OD 1.9 m"],
               ["Total HRSG weight (pressure parts, casing, steel, catalysts)", "6,400 t"]], (7.5, 9.9), 8.5),
    ]
    render("IEC-ALP-DS-104", "Heat Recovery Steam Generator IE-HR3 - Datasheet", blocks)


def ds105():
    blocks = [
        ("h", "1. Thermal design (HEI 11th edition)"),
        ("t", [["Item", "Value"],
               ["Type", "single-shell, two-pass surface condenser, divided water boxes (two halves), side-mounted to the ST"],
               ["Design duty at SRC-NG-100", f(P["Q_cond"], 2) + " MW"],
               ["Condenser pressure at SRC", f(P["p_cond_mbar"], 2) + " mbar(a) (guarantee 35.5 mbar(a); ER-07.08 limit 36 mbar(a))"],
               ["Seawater inlet / outlet", "16.0 / 23.0 degC"],
               ["Circulating water flow", f(P["cw_m3h"], 0) + " m3/h (" + f(P["cw_flow"], 0) + " kg/s)"],
               ["Seawater properties", "density 1,016.5 kg/m3, cp 4.07 kJ/kg K (salinity 22 PSU, ALP-OWN-MAR-001)"],
               ["Surface area (incl. 5 % plugging margin)", "15,200 m2"],
               ["Tubes", "titanium Grade 2 welded, OD 25.4 mm, 0.5 mm wall (0.7 mm in the impingement zone), effective "
                         "length 9.35 m, 20,400 tubes"],
               ["Tube velocity", "2.1 m/s (ER-07.07: 1.8 - 2.2 m/s)"],
               ["Cleanliness factor", "0.85"],
               ["Seawater pressure loss", "0.62 bar (water boxes and tubes, clean)"],
               ["Tube sheets / water boxes", "titanium-clad carbon steel / rubber-lined carbon steel with impressed-current cathodic "
                                             "protection"],
               ["Hotwell", "30 m3 working volume (3.9 min at SRC condensate flow)"],
               ["Dissolved oxygen at hotwell outlet", "<= 20 ug/l above 50 % load"],
               ["Bypass operation", "100 % HP/IP/LP bypass at SRC: 365 MW; maximum (cold ambient, LDO): 395 MW; seawater rise "
                                    "then 8.8 - 9.5 K (transient)"],
               ["Vacuum system", "2 x 100 % liquid-ring vacuum pumps 75 kW (holding) + hogging by both pumps in parallel "
                                 "(30 min to 150 mbar)"],
               ["Tube cleaning", "**tube cleaning ball system and debris filters: not included (by EPC with the CW system)**; "
                                 "condenser inlet water boxes prepared for ball injection and strainer connection"]], (5.0, 12.4), 8.5),
        ("h", "2. Performance at other seawater temperatures (base load, SRC steam flow)"),
        ("t", [["Seawater inlet degC", "6.9", "11.0", "16.0", "22.0", "24.0", "27.0"],
               ["Condenser pressure mbar(a)", "19.8", "25.5", "34.4", "48.7", "54.4", "64.2"]], None, 8.5),
        ("p", "At 27 degC seawater the condenser pressure is 64 mbar(a), well inside the ST continuous limit of 120 mbar(a) "
              "(ER-07.08)."),
        ("h", "3. Weights"),
        ("t", [["Item", "Weight (t)"], ["Condenser shell with tube bundles (shipped in 4 modules)", "4 x 95"],
               ["Operating weight (water-filled)", "1,650"], ["Flooded weight (hydrotest)", "2,300"]], (11.0, 6.4), 8.5),
    ]
    render("IEC-ALP-DS-105", "Seawater-Cooled Surface Condenser - Datasheet", blocks)


def ds106():
    blocks = [
        ("h", "1. Generator step-up transformers"),
        ("t", [["Item", "GT GSU", "ST GSU"],
               ["Rating", "560 MVA ONAF (336 MVA ONAN)", "270 MVA ONAF (162 MVA ONAN)"],
               ["Voltage ratio / vector group", "400 +/- 8 x 1.25 % / 21 kV, YNd11, OLTC", "400 +/- 8 x 1.25 % / 15.75 kV, YNd11, OLTC"],
               ["Impedance", "14.5 %", "13.5 %"],
               ["No-load losses / load losses at rating", "190 kW / 1,150 kW", "110 kW / 620 kW"],
               ["Losses at SRC-NG-100 load (P, pf 0.85)", "1,110 kW", "580 kW"],
               ["Insulation level HV / neutral", "LI 1425 kV / SI 1050 kV; neutral solidly earthed", "same"],
               ["Oil", "natural ester (K-class) - fire risk reduction; oil volume 98 m3", "natural ester; 62 m3"],
               ["Transport weight (without oil, nitrogen-filled) / dimensions", "268 t / 11.2 x 4.1 x 4.6 m", "168 t / 9.4 x 3.8 x 4.3 m"],
               ["Total weight", "395 t", "245 t"],
               ["HV connection", "HV bushings with air terminal; 380 kV connection to the GIS by EPC (IF-34/IF-35)", "same"],
               ["Monitoring", "on-line DGA, bushing monitoring, fibre-optic hot-spot; IEC 61850 to the EPC substation system", "same"]],
         (5.0, 6.2, 6.2), 8),
        ("h", "2. Interfaces"),
        ("ul", ["Transformer foundations, oil pits (110 % oil volume), fire walls and deluge system: by EPC (NFPA 850, IEC 61936-1).",
                "380 kV HV connection from the bushings to the GIS, surge arresters at the HV bushings: by EPC.",
                "IPB connection flanges at the LV bushings: by IEC."]),
    ]
    render("IEC-ALP-DS-106", "Generator Step-up Transformers - Datasheet", blocks)


IF_ROWS = [
    ["Tag", "Service", "IEC side", "EPC side", "Medium", "Size / rating", "Design p / T", "Operating (SRC)"],
    ["IF-01", "Natural gas to GT", "performance gas heater inlet flange", "gas from EPC gas receiving/metering station",
     "natural gas", "12\" ASME CL600 RF", "63 barg / 80 degC", f"{P['fuel_flow']:.2f} kg/s, 37.0 barg, 25 degC"],
    ["IF-02", "LDO to GT", "liquid fuel skid inlet flange", "LDO forwarding/treatment system", "LDO", "6\" CL150", "16 barg / 60 degC",
     "22.1 kg/s (LDO operation)"],
    ["IF-03", "Demin water for water injection", "water injection skid inlet", "demin water storage/pumps", "demin water",
     "4\" CL150", "16 barg / 50 degC", "22.5 kg/s (LDO operation)"],
    ["IF-04", "Compressor wash water", "wash skid inlet", "demin water", "demin water", "2\" CL150", "10 barg / 50 degC", "5 m3 per wash"],
    ["IF-05", "HP steam", "HRSG HP SH outlet stop valve, weld end", "main steam piping (EPC)", "steam", "12\" (P91) weld end",
     "210 bar(a) / 610 degC", f"{PS[6]['mass_flow']:.2f} kg/s, {PS[6]['pressure']:.1f} bar(a), {PS[6]['temperature']:.0f} degC"],
    ["IF-06", "HP steam", "ST main stop valve inlet, weld end", "main steam piping (EPC)", "steam", "12\" (P91) weld end",
     "200 bar(a) / 605 degC", f"{PS[7]['pressure']:.1f} bar(a), {PS[7]['temperature']:.0f} degC"],
    ["IF-07", "Cold reheat", "HP turbine exhaust, weld end", "cold reheat piping (EPC)", "steam", "24\" (P22/A106) weld end",
     "52 bar(a) / 420 degC", f"{PS[8]['pressure']:.1f} bar(a), {PS[8]['temperature']:.0f} degC"],
    ["IF-08", "Cold reheat", "HRSG reheater inlet, weld end", "cold reheat piping (EPC)", "steam", "24\" weld end",
     "52 bar(a) / 420 degC", f"{PS[9]['pressure']:.1f} bar(a), {PS[9]['temperature']:.0f} degC"],
    ["IF-09", "Hot reheat", "HRSG reheater outlet stop valve, weld end", "hot reheat piping (EPC)", "steam", "28\" (P91) weld end",
     "50 bar(a) / 610 degC", f"{PS[11]['mass_flow']:.2f} kg/s, {PS[11]['pressure']:.1f} bar(a)"],
    ["IF-10", "Hot reheat", "ST IP stop-control valve inlet", "hot reheat piping (EPC)", "steam", "28\" (P91) weld end",
     "48 bar(a) / 605 degC", f"{PS[12]['pressure']:.1f} bar(a), {PS[12]['temperature']:.0f} degC"],
    ["IF-11", "LP steam", "HRSG LP SH outlet stop valve", "LP steam piping (EPC)", "steam", "20\" weld end", "10 bar(a) / 330 degC",
     f"{PS[13]['mass_flow']:.2f} kg/s, {PS[13]['pressure']:.2f} bar(a)"],
    ["IF-12", "LP steam", "ST LP admission valve inlet", "LP steam piping (EPC)", "steam", "20\" weld end", "8 bar(a) / 330 degC",
     f"{PS[14]['pressure']:.2f} bar(a)"],
    ["IF-13", "HP / IP / LP bypass valves", "valves supplied loose by IEC (incl. desuperheating)", "installation in EPC piping, "
     "spray water from BFP/CEP (EPC)", "steam / water", "per valve datasheets", "-", "start-up / trip"],
    ["IF-14", "Bypass steam to condenser", "condenser neck dump devices (IEC)", "IP/LP bypass downstream piping (EPC)", "steam",
     "2 x 36\" + 1 x 28\"", "16 bar(a) / 250 degC", "start-up / trip"],
    ["IF-15", "Condensate from hotwell", "hotwell outlet nozzle", "CEP suction piping (EPC)", "condensate", "20\" CL150",
     "full vacuum - 3.5 bar(a) / 80 degC", f"{PS[17]['mass_flow']:.2f} kg/s, {PS[17]['temperature']:.1f} degC"],
    ["IF-16", "Condensate to CPH", "HRSG CPH inlet", "condensate line (EPC)", "condensate", "12\" CL300", "25 bar(a) / 180 degC",
     f"{PS[20]['mass_flow']:.2f} kg/s"],
    ["IF-17", "BFP suction", "HRSG LP drum downcomer outlet nozzle", "BFP suction piping (EPC)", "feedwater", "16\" CL150",
     "10 bar(a) / 185 degC", f"{PS[21]['mass_flow']:.2f} kg/s, {PS[21]['temperature']:.1f} degC"],
    ["IF-18", "HP feedwater", "HRSG HP ECO1 inlet", "HP feedwater piping (EPC)", "feedwater", "10\" CL2500", "235 bar(a) / 200 degC",
     f"{PS[22]['mass_flow']:.2f} kg/s, {PS[22]['pressure']:.0f} bar(a)"],
    ["IF-19", "IP feedwater", "HRSG IP ECO inlet", "IP feedwater piping (EPC)", "feedwater", "6\" CL600", "60 bar(a) / 200 degC",
     f"{PS[23]['mass_flow']:.2f} kg/s, {PS[23]['pressure']:.0f} bar(a)"],
    ["IF-20", "FGH water return", "performance gas heater water outlet", "return line to condensate (EPC)", "feedwater", "4\" CL600",
     "60 bar(a) / 260 degC", f"{PS[19]['mass_flow']:.2f} kg/s, {PS[19]['temperature']:.0f} degC"],
    ["IF-21", "HRSG blowdown", "continuous / intermittent blowdown valves outlet", "blowdown tanks (EPC)", "water", "2\" / 3\"",
     "210 bar(a) / 370 degC", "0 in normal operation"],
    ["IF-22", "HRSG drains and vents", "drain/vent valves outlet (IEC)", "drain collection, flash tanks, silencers (EPC)",
     "water/steam", "various", "-", "start-up"],
    ["IF-23", "Chemical dosing", "dosing nozzles on drums/condensate", "dosing skids and lines (EPC)", "chemicals", "1\"", "-", "-"],
    ["IF-24", "Sampling", "sample nozzles with root valves (IEC)", "sample lines, coolers, SWAS (EPC)", "water/steam", "1/2\"", "-", "-"],
    ["IF-25", "Aqueous ammonia", "ammonia injection skid inlet (IEC)", "storage, unloading, forwarding pumps (EPC)", "NH4OH 24.5 %",
     "1\"", "10 barg / 40 degC", "125 kg/h"],
    ["IF-26", "Flue gas to stack", "HRSG outlet transition flange", "stack (EPC), CEMS (EPC)", "flue gas", "D 9.2 m",
     "+/- 50 mbar / 200 degC", f"{PS[5]['mass_flow']:.1f} kg/s, {PS[5]['temperature']:.1f} degC"],
    ["IF-27", "Circulating water", "condenser water box flanges (2 inlets, 2 outlets)", "CW pipes (EPC)", "seawater", "4 x DN1800",
     "4 barg / 40 degC", f"{P['cw_m3h']:,.0f} m3/h"],
    ["IF-28", "Closed cooling water", "IEC coolers inlet/outlet flanges", "CCW system (EPC)", "demin water + inhibitor", "various",
     "10 barg / 60 degC", "see IEC-ALP-REQ-001 table 5.1"],
    ["IF-29", "Instrument and service air", "IEC skids", "compressed air system (EPC)", "air", "1\" - 2\"", "10 barg", "see REQ-001 6.1"],
    ["IF-30", "Nitrogen", "HRSG lay-up / generator purge connections", "N2 supply (EPC)", "N2", "1\"", "10 barg", "lay-up"],
    ["IF-31", "Hydrogen / CO2", "GT generator gas control panel inlet", "bottle racks and piping (EPC)", "H2 / CO2", "1/2\"", "16 barg", "12 Nm3/day H2"],
    ["IF-32", "Auxiliary steam", "gland steam header inlet (IEC)", "auxiliary boiler steam (EPC)", "steam", "4\"", "16 bar(a) / 320 degC",
     "4.5 t/h at start-up"],
    ["IF-33", "Fire protection", "GT enclosure CO2 system (IEC)", "fire detection/alarm, transformer deluge, hydrants (EPC)", "-", "-",
     "-", "signals to EPC fire alarm panel"],
    ["IF-34", "GT GSU HV", "GT GSU HV bushings", "380 kV connection to GIS (EPC)", "electrical", "400 kV", "-", "560 MVA"],
    ["IF-35", "ST GSU HV", "ST GSU HV bushings", "380 kV connection to GIS (EPC)", "electrical", "400 kV", "-", "270 MVA"],
    ["IF-36", "UAT supply", "IPB tap-off flange (GT and ST IPB)", "UAT and MV busduct (EPC)", "electrical", "21 kV / 15.75 kV", "-", "-"],
    ["IF-37", "MV supplies", "SFC isolation transformer, IEC MV consumers", "10.5 kV switchgear feeders (EPC)", "electrical", "10.5 kV",
     "-", "SFC 8.5 MW at start"],
    ["IF-38", "LV supplies", "IEC MCCs incoming", "0.4 kV switchgear feeders (EPC)", "electrical", "0.4 kV", "-", "see REQ-001 7.1"],
    ["IF-39", "DC and UPS", "TCS, protection, emergency oil pumps", "220 V DC and 230 V AC UPS (EPC)", "electrical", "-", "-", "see REQ-001 7.2"],
    ["IF-40", "Control system interface", "IEC turbine control system (TCS) gateway", "plant DCS (EPC)", "data", "OPC UA + Modbus TCP, "
     "redundant; hardwired trips", "-", "about 2,500 signals"],
    ["IF-41", "Protection interface", "generator/GSU protection panels", "switchyard protection and TSO SCADA (EPC)", "signals", "IEC 61850 + hardwired", "-", "-"],
    ["IF-42", "Earthing", "equipment earthing bosses", "earthing grid (EPC)", "electrical", "-", "-", "-"],
    ["IF-43", "Foundations", "anchor bolts, sole plates, load data (IEC)", "foundation design and construction (EPC)", "civil", "-", "-",
     "see REQ-001 section 9"],
]


def if001():
    blocks = [
        ("h", "1. Terminal points between IEC and the EPC"),
        ("p", "Each terminal point is the physical boundary between the IEC supply (power island) and the EPC's balance of plant. "
              "Design conditions are the IEC design values of the IEC component; the EPC side shall be designed for at least the "
              "same conditions. Nozzle load limits: attachment A (to follow with IEC-ALP-IF-001 Rev A, 8 weeks after NTP)."),
        ("t", IF_ROWS, (1.1, 2.6, 3.8, 3.6, 1.9, 2.4, 2.6, 3.4), 6.5, "Terminal points"),
    ]
    render("IEC-ALP-IF-001", "Terminal Points and Interface Data - Power Island", blocks, landscape=True)


DOR_ROWS = [
    ["No", "Item", "Design", "Supply", "Erection", "Supervision", "Commissioning", "Remarks"],
    ["1.1", "Gas turbine package incl. enclosure, lube oil, hydraulic and fuel skids", "IEC", "IEC", "EPC", "IEC", "IEC with EPC", ""],
    ["1.2", "GT inlet air filter house, anti-icing, silencer, inlet duct", "IEC", "IEC", "EPC", "IEC", "IEC", "evaporative cooler: space provision only (not included)"],
    ["1.3", "GT exhaust diffuser and expansion joint to HRSG", "IEC", "IEC", "EPC", "IEC", "IEC", ""],
    ["1.4", "Performance fuel gas heater and GT gas skid (from IF-01)", "IEC", "IEC", "EPC", "IEC", "IEC", ""],
    ["1.5", "Gas receiving and metering station, start-up gas heating, final filtration 3 um", "EPC", "EPC", "EPC", "-", "EPC", "IEC fuel gas requirements per DS-101"],
    ["1.6", "Liquid fuel skid, water injection skid, purge and recirculation (from IF-02/03)", "IEC", "IEC", "EPC", "IEC", "IEC", ""],
    ["1.7", "LDO unloading, storage, forwarding and treatment", "EPC", "EPC", "EPC", "-", "EPC", ""],
    ["1.8", "Compressor washing skid", "IEC", "IEC", "EPC", "IEC", "IEC", "demin water by EPC"],
    ["1.9", "GT enclosure fire suppression (CO2) and gas detection", "IEC", "IEC", "EPC", "IEC", "IEC", "fire alarm panel by EPC"],
    ["2.1", "HRSG pressure parts, casing, drums, platforms, CPH recirculation pumps", "IEC", "IEC", "EPC", "IEC", "EPC with IEC", ""],
    ["2.2", "SCR catalyst, ammonia injection grid and skid, CO catalyst", "IEC", "IEC", "EPC", "IEC", "EPC with IEC", ""],
    ["2.3", "Aqueous ammonia unloading, storage and forwarding", "EPC", "EPC", "EPC", "-", "EPC", "not in the Owner contract scope list"],
    ["2.4", "HRSG blowdown tanks, drain tanks, atmospheric flash tanks", "EPC", "EPC", "EPC", "-", "EPC", "IEC gives flows and conditions"],
    ["2.5", "Chemical dosing and steam/water sampling (SWAS)", "EPC", "EPC", "EPC", "-", "EPC", "IEC sample/dosing nozzles"],
    ["2.6", "Stack incl. silencer, CEMS, aviation lighting", "EPC", "EPC", "EPC", "-", "EPC", "65 m (ER-06.06)"],
    ["3.1", "Steam turbine with valves, lube/EHC oil, gland steam system, turning gear", "IEC", "IEC", "EPC", "IEC", "IEC", ""],
    ["3.2", "HP / IP / LP bypass valves with desuperheaters", "IEC", "IEC", "EPC", "IEC", "IEC with EPC", "loose supply"],
    ["3.3", "HP, cold reheat, hot reheat and LP steam piping HRSG - ST incl. supports", "EPC", "EPC", "EPC", "-", "EPC",
     "IEC: nozzle loads, pressure drop assumptions (PER-001 1)"],
    ["3.4", "Bypass piping and spray water piping", "EPC", "EPC", "EPC", "-", "EPC", ""],
    ["3.5", "ST drains, drain collection and flash tank", "EPC", "EPC", "EPC", "-", "EPC", "IEC drain list"],
    ["4.1", "Condenser incl. dump devices, hotwell, vacuum pumps", "IEC", "IEC", "EPC", "IEC", "EPC with IEC", ""],
    ["4.2", "Tube cleaning ball system and debris filters", "EPC", "EPC", "EPC", "-", "EPC", "required by ER-07.03"],
    ["4.3", "Circulating water system, intake, pumps, outfall, electro-chlorination", "EPC", "EPC", "EPC", "-", "EPC", ""],
    ["4.4", "Condensate pumps, boiler feed pumps, condensate polishing", "EPC", "EPC", "EPC", "-", "EPC", "IEC pressure assumptions PER-001"],
    ["4.5", "Closed cooling water system", "EPC", "EPC", "EPC", "-", "EPC", "IEC cooler demands REQ-001 5.1"],
    ["5.1", "GT and ST generators incl. excitation, neutral earthing", "IEC", "IEC", "EPC", "IEC", "IEC", ""],
    ["5.2", "Generator circuit breakers, isolated phase busducts", "IEC", "IEC", "EPC", "IEC", "IEC", ""],
    ["5.3", "GT and ST generator step-up transformers", "IEC", "IEC", "EPC", "IEC", "IEC with EPC", ""],
    ["5.4", "SFC and isolation transformer", "IEC", "IEC", "EPC", "IEC", "IEC", ""],
    ["5.5", "Generator and GSU protection panels", "IEC", "IEC", "EPC", "IEC", "IEC with EPC", ""],
    ["5.6", "H2 and CO2 bottle racks and supply piping", "EPC", "EPC", "EPC", "-", "EPC", "IEC gas control panel"],
    ["5.7", "Unit auxiliary transformers, MV/LV switchgear, UPS/DC, cabling", "EPC", "EPC", "EPC", "-", "EPC", ""],
    ["5.8", "380 kV connection GSU - GIS, GIS switchyard", "EPC", "EPC", "EPC", "-", "EPC", ""],
    ["6.1", "Turbine control system (TCS) for GT, ST, generators, HRSG protection (SIS)", "IEC", "IEC", "EPC", "IEC", "IEC",
     "stand-alone TCS with OPC UA gateway to the DCS"],
    ["6.2", "Plant DCS incl. HRSG and BOP control, operator stations", "EPC", "EPC", "EPC", "-", "EPC", "TCS HMI screens via gateway"],
    ["6.3", "CEMS", "EPC", "EPC", "EPC", "-", "EPC", ""],
    ["6.4", "Vibration and performance monitoring, remote diagnostics", "IEC", "IEC", "EPC", "IEC", "IEC", "secure link by EPC (IEC 62443)"],
    ["7.1", "Foundations of IEC equipment (GT, generators, ST, HRSG, condenser, GSUs)", "EPC", "EPC", "EPC", "-", "-",
     "IEC load data, anchor bolts and sole plates"],
    ["7.2", "Turbine hall(s), cranes, HRSG access structures beyond IEC platforms", "EPC", "EPC", "EPC", "-", "-", ""],
    ["8.1", "Transport of IEC equipment to site (incl. heavy transport, barge, SPMT)", "-", "IEC to port of discharge; EPC from "
     "port to foundation", "-", "-", "-", "barge landing TP-M1 by EPC"],
    ["8.2", "Performance test", "EPC with IEC", "EPC (test instruments)", "-", "IEC", "EPC", "ASME PTC 46"],
    ["8.3", "Training of Owner staff on IEC equipment", "IEC", "IEC", "-", "-", "-", ""],
]


def dor001():
    blocks = [
        ("h", "1. Division of responsibility between IEC and the EPC"),
        ("p", "Design / Supply / Erection / Supervision (of erection) / Commissioning. 'IEC with EPC' = IEC leads, EPC provides "
              "labour, utilities and plant systems. Items not listed follow the consortium agreement: power island equipment "
              "by IEC (Owner contract items SC-006 to SC-010, SC-012 to SC-014, SC-029, SC-037, SC-054), everything else by EPC."),
        ("t", DOR_ROWS, (0.9, 7.3, 1.6, 2.4, 1.6, 1.8, 2.3, 6.5), 7, "Division of responsibility"),
    ]
    render("IEC-ALP-DOR-001", "Division of Responsibility - Power Island (IEC) and Balance of Plant (EPC)", blocks, landscape=True)


AUX = [["Consumer (IEC scope)", "kW at base load, SRC"],
       ["GT package auxiliaries (lube oil, hydraulic, enclosure ventilation, fuel skid, filter house)", "1,180"],
       ["GT generator auxiliaries (seal oil, stator water, excitation)", "420"],
       ["ST auxiliaries (lube oil, EHC, gland steam exhausters)", "380"],
       ["ST generator auxiliaries (excitation, TEWAC)", "260"],
       ["HRSG auxiliaries (CPH recirculation pump, ammonia skid, drains)", "210"],
       ["Condenser vacuum pump (1 running)", "75"],
       ["Turbine control system and UPS load", "60"],
       ["**Total IEC auxiliaries**", "**2,585**"],
       ["GT GSU losses at SRC load", "1,110"],
       ["ST GSU losses at SRC load", "580"]]


def req001():
    blocks = [
        ("h", "1. Purpose"),
        ("p", "Requirements of IEC to the EPC for the design of the balance of plant around the power island. They complement the "
              "Owner's Employer's Requirements (Appendix A of ALP-EPC-001); where the Owner's requirement is stricter it prevails. "
              "Section 12 lists IEC's exceptions and clarifications to the Employer's Requirements."),
        ("h", "2. Fuel gas at IF-01"),
        ("ul", ["Pressure 36.5 barg minimum at base load, 41.0 barg maximum, rate of change <= 0.7 bar/s; at 36.5 barg IEC "
                "requires 45 barg at TP-G1 less the EPC station losses (EPC to confirm about 4 bar).",
                "Temperature 5 - 40 degC and >= 28 K superheat above the hydrocarbon and water dew points at all loads.",
                "Final filtration 3 um absolute (beta 200) with coalescing filter-separator (EPC), no free liquids.",
                "Trace metals: Na + K <= 0.02 mg/Sm3 equivalent, V <= 0.01, Pb <= 0.02, Ca <= 0.02 mg/Sm3 at IF-01; H2S + COS <= 5 "
                "mg/Sm3; total sulphur <= 30 mg S/Sm3.",
                "Gas chromatograph signal (LHV, Wobbe index, H2) to the TCS for combustion tuning (update every <= 4 min)."]),
        ("h", "3. LDO and water injection"),
        ("ul", ["LDO at IF-02 per ER-03.07 (GT inlet limits) after EPC treatment; 3.0 - 6.0 barg; 10 - 45 degC.",
                "Demin water for water injection at IF-03: conductivity <= 0.2 uS/cm, Na+K <= 0.02 mg/kg, silica <= 0.02 mg/kg; "
                "22.5 kg/s at base load on LDO."]),
        ("h", "4. Water/steam chemistry (IAPWS TGD, VGB-S-010)"),
        ("t", [["Parameter", "Feedwater", "HP steam", "Unit"],
               ["Cation conductivity (25 degC)", "<= 0.2", "<= 0.2", "uS/cm"],
               ["Sodium", "<= 3", "<= 2", "ug/kg"],
               ["Silica", "<= 10", "<= 10", "ug/kg"],
               ["Dissolved oxygen", "<= 20 (AVT(O))", "-", "ug/kg"],
               ["pH (ammonia)", "9.4 - 9.8", "-", "-"],
               ["Iron / copper", "<= 5 / <= 2", "-", "ug/kg"]], (6.0, 3.0, 3.0, 2.0), 8.5),
        ("h", "5. Closed cooling water demand"),
        ("t", [["Consumer", "Duty kW", "Flow m3/h", "Max. supply temperature degC"],
               ["GT lube oil coolers", "1,600", "140", "38"],
               ["GT generator H2 coolers and stator water coolers", "4,300", "520", "38"],
               ["ST lube oil coolers", "850", "75", "38"],
               ["ST generator TEWAC coolers", "2,200", "260", "36"],
               ["SFC, excitation cubicles, vacuum pump seal water", "450", "60", "38"],
               ["**Total**", "**9,400**", "**1,055**", ""]], (8.0, 2.5, 2.5, 4.4), 8.5),
        ("h", "6. Utilities"),
        ("t", [["Utility", "Requirement"],
               ["Instrument air", "dew point -40 degC at 7 barg, 350 Nm3/h continuous (valves, purge, pulse cleaning of filter house)"],
               ["Service air", "150 Nm3/h intermittent"],
               ["Nitrogen", "HRSG and generator purge/lay-up: 3,000 Nm3 per event, 99.9 %"],
               ["Hydrogen / CO2", "H2 99.9 %, 12 Nm3/day make-up; CO2 150 Nm3 per purge"],
               ["Auxiliary steam", "4.5 t/h at 10 bar(a), 250 - 300 degC for gland sealing at start-up; HRSG sparging 3 t/h optional"]],
         (4.5, 12.9), 8.5),
        ("h", "7. Electrical"),
        ("t", AUX, (13.0, 4.4), 8.5, "Auxiliary loads"),
        ("p", "Largest IEC MV consumer: SFC 8.5 MW at 10.5 kV during start-up (30 min). DC: 220 V DC emergency oil pumps "
              "(GT 37 kW, ST 45 kW, 3 h autonomy). UPS 230 V AC: 45 kVA for the TCS."),
        ("h", "8. Instrumentation and control"),
        ("ul", ["IEC TCS (IE-Control T3) controls and protects GT, ST, generators and the HRSG safety functions; it is a stand-alone "
                "system connected to the plant DCS through a redundant OPC UA gateway (Modbus TCP back-up) with hardwired trips "
                "and permissives.",
                "Operator control of the power island from the DCS operator stations through the gateway (TCS HMI screens are "
                "also available on dedicated IEC stations in the central control room).",
                "Time synchronisation GPS/PTP by EPC; cybersecurity zone and conduit design per IEC 62443-3-2 by EPC with IEC input."]),
        ("h", "9. Foundations and loads"),
        ("t", [["Equipment", "Static load (t)", "Dynamic / design data"],
               ["GT + generator train (common table-top or block foundation)", "1,050", "unbalance and short-circuit torque (4.8 x "
                                                                                       "rated), natural frequencies +/- 20 % from 50 Hz"],
               ["ST + generator (spring-mounted table-top)", "980", "same criteria; spring elements by EPC"],
               ["HRSG (steel structure base plates)", "7,800 operating", "wind 30 m/s, seismic PGA 0.45 g (TBDY 2018)"],
               ["Condenser", "1,650 operating / 2,300 flooded", "side-exhaust; expansion joint to ST by IEC"],
               ["GT GSU / ST GSU", "395 / 245", "oil pits and fire walls by EPC"]], (6.0, 3.2, 8.2), 8),
        ("h", "10. Transport and erection"),
        ("t", [["Piece", "Weight (t)", "Dimensions (m)", "Route"],
               ["GT core engine", "405", "13.2 x 5.9 x 5.6", "barge to TP-M1, SPMT to foundation"],
               ["GT generator stator", "340", "10.8 x 4.6 x 4.9", "barge, SPMT, jacking/skidding"],
               ["GT GSU", "268", "11.2 x 4.1 x 4.6", "barge, SPMT"],
               ["HRSG HP evaporator module", "185", "24.0 x 3.6 x 26.5", "barge, SPMT, heavy crane (by EPC)"],
               ["ST generator stator", "185", "8.6 x 4.1 x 4.2", "barge, SPMT, jacking"],
               ["ST GSU", "168", "9.4 x 3.8 x 4.3", "barge, SPMT"],
               ["HRSG HP drum", "96", "14.0 x 1.9", "barge or road (under 120 t)"]], (4.5, 1.8, 3.4, 7.7), 8.5),
        ("p", "All pieces above 120 t exceed the road limit (ALP-OWN-GEN-001): the barge landing TP-M1 and the haul road shall "
              "take 480 t gross (SPMT) with ground bearing to be verified by the EPC."),
        ("h", "11. Maintenance"),
        ("ul", ["GT hall crane: 120 t main hook (GT rotor 112 t), 20 t auxiliary; ST hall crane: 100 t (IP/LP rotor 86 t).",
                "Laydown for GT major inspection: 900 m2 next to the GT, rated 10 t/m2."]),
        ("h", "12. IEC exceptions and clarifications to the Employer's Requirements"),
        ("t", [["No", "Requirement", "IEC position"],
               ["E-01", "ER-05.02 evaporative cooler provision", "space provision only; cooler, water supply and drain not included"],
               ["E-02", "SC-013 TCS 'integrated in' the plant DCS", "stand-alone IEC TCS with OPC UA gateway to the EPC DCS (IEC "
                                                                    "standard, proven on the fleet)"],
               ["E-03", "ER-07.03 automatic ball tube cleaning", "not in IEC scope (EPC with the CW system); water boxes prepared"],
               ["E-04", "ER-05.07 HGP >= 25,000 EOH", "complies with IEC's EOH formula (starts count 20 EOH each, see DS-101 5)"],
               ["E-05", "Main steam, reheat and LP steam piping", "by EPC (not an IEC item in SC-006 to SC-014)"],
               ["E-06", "H2/CO2 supply for the GT generator", "by EPC"],
               ["E-07", "ER-04.09 house load operation 60 min", "complies for the GT; ST follows on bypass operation"],
               ["E-08", "Guarantees", "IEC guarantees are gross values of the power island (PER-001 section 7); net plant "
                                      "guarantees to the Owner remain with the consortium"]], (1.2, 5.5, 10.7), 8, "Exceptions"),
    ]
    render("IEC-ALP-REQ-001", "Power Island Interface Requirements to the EPC (utilities, chemistry, electrical, civil, transport)",
           blocks)


DOCS = {"PER": per001, "DS101": ds101, "DS102": ds102, "DS103": ds103, "DS104": ds104, "DS105": ds105, "DS106": ds106,
        "IF": if001, "DOR": dor001, "REQ": req001}

if __name__ == "__main__":
    for k in (sys.argv[3:] or DOCS):
        DOCS[k]()
